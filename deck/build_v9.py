# -*- coding: utf-8 -*-
"""build_v9.py: v9 연구 결과 보고 덱(국문, 16:9). 명세 deck/deck_spec_v9.json. 지침 docs/DESIGN_BRIEF_v9_2026-10-15.md.
문법 V9: 주황 제목(명사구) + 회색 부제 한 줄 + 그림이 본문 폭 12 in 을 채움 + 의미 줄(잉크 14 pt, 개조식) + 필요할 때만 회색 설명 줄(12 pt).
강조는 주황 외곽선 상자(목표 · 결론 · 의의 쪽만, 소제목 없이 문장 1–3개). 한글은 어절 단위로 직접 줄을 나눈다(LibreOffice 의 낱말 중간 끊김 방지).
표지는 사용자가 v8 에서 직접 고친 설정(흰 제목 32 pt, '예측과' 뒤 줄바꿈, 줄 간격 112 %)을 따른다.
수치는 docs/numbers_v6.json + docs/numbers_v9.json 의 {nMacro} 토큰({nMacro|pt} = 구간을 뺀 점추정).
출력 deck/render/v9_report.{pptx,pdf}, deck/render/v9_report_pages/page-XX.png. 실행: OMP_NUM_THREADS=2 nice -n 10 python3 deck/build_v9.py [--no-pdf]
PDF 변환은 paper_report_lib.export_pdf(GPU 차단 환경변수 적용).
"""
from __future__ import annotations
import argparse, json, os, re, subprocess, sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.util import Inches, Pt  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
import final_lib as fl  # noqa: E402
import paper_report_lib as L  # noqa: E402
from paper_report_lib import ACC, F_M, F_S, F_X, GRAY, GRAY2, INK, run, text  # noqa: E402

ROOT = L.ROOT; DECK = L.DECK; RENDER = DECK / "render"; OUTSTEM = os.environ.get("DECK_OUT", "v9_report"); PAGES = RENDER / f"{OUTSTEM}_pages"
PPTX = RENDER / f"{OUTSTEM}.pptx"; PDF = RENDER / f"{OUTSTEM}.pdf"
assert OUTSTEM != "v8_report", "v8_report.pptx 는 사용자가 직접 고친 파일이라 덮어쓰지 않는다"
FIG = ROOT / "outputs" / "figures" / "paper" / "v7" / "slide"
DECKNATIVE = ROOT / "outputs" / "figures" / "paper" / "v7" / "deck"
DECKFIG = DECK / "assets" / "v7"; DECKFIG8 = DECK / "assets" / "v8"; DECKFIG9 = DECK / "assets" / "v9"; ASSET = DECK / "assets" / "paper_report"
SCRATCH = Path(os.environ.get("PR_TMP", tempfile.gettempdir()))
NUM = {}
for f in ("numbers_v6.json", "numbers_v9.json"):
    p = ROOT / "docs" / f
    if p.exists():
        NUM.update(json.load(open(p, encoding="utf-8")))
Y0, YB = 1.30, 6.40          # 본문 위·아래
CAP_Y = 6.62                 # 설명 줄
W_BODY = 12.0
MEAN_PT, CAP_PT, BOX_PT = 14, 12, 15
LINE_IN = {14: 0.285, 12: 0.25, 15: 0.31}
MISSING = []


def subst(obj):
    if isinstance(obj, str):
        def rep(m):
            v = NUM.get(m.group(1)[1:])
            if v is None:
                MISSING.append(m.group(1)); return m.group(0)
            return v.split(" [")[0] if m.group(2) else v
        return re.sub(r"\{(n[A-Za-z]+)(\|pt)?\}", rep, obj)
    if isinstance(obj, list):
        return [subst(x) for x in obj]
    if isinstance(obj, dict):
        return {k: subst(v) for k, v in obj.items()}
    return obj


SPEC_NAME = os.environ.get("DECK_SPEC", "deck_spec_v9.json"); SPEC = subst(json.load(open(DECK / SPEC_NAME, encoding="utf-8")))
CACHE = DECK / "assets" / "v7" / "_rgb"


def asset(name):
    """그림 파일을 찾고, RGBA 면 흰 바탕에 평탄화한 RGB 사본(캐시)을 돌려준다(PDF 변환 시 smask 제거)."""
    for p in (DECKFIG9 / name, DECKNATIVE / name, FIG / name, DECKFIG8 / name, DECKFIG / name, ASSET / name):
        if p.exists():
            from PIL import Image
            with Image.open(p) as im:
                if im.mode != "RGBA":
                    return p
                CACHE.mkdir(parents=True, exist_ok=True); q = CACHE / p.name
                if not q.exists() or q.stat().st_mtime < p.stat().st_mtime:
                    bg = Image.new("RGB", im.size, (255, 255, 255)); bg.paste(im, mask=im.split()[3]); bg.save(q, dpi=im.info.get("dpi", (300, 300)))
                return q
    return None


def wrapped(t, width_in, pt, weight="Medium"):
    """어절 단위 줄 목록(공백에서만 나눔, 숫자 + 단위 묶음)."""
    out = []
    for part in t.split("\n"):
        out += L.wrap_text(part, width_in, pt, weight)
    return out


def ktext(sl, x, y, w, h, t, pt, color, font, weight="Medium", **kw):
    """어절 단위로 미리 나눈 문단 하나(a:br 줄바꿈)."""
    lines = wrapped(t, w - 0.2, pt, weight)
    tb = text(sl, x, y, w, h, [run("\n".join(lines), pt, color, font)], **kw); L.break_runs(tb)
    return tb, len(lines)


def meaning_h(meaning):
    return 0.0 if not meaning else LINE_IN[MEAN_PT] * len(wrapped(meaning, W_BODY - 0.2, MEAN_PT)) + 0.10


def cap_h(t):
    return 0.0 if not t else 0.26 * len(wrapped(t, W_BODY - 0.2, CAP_PT)) + 0.02


def caption(sl, t, meaning=None, top=None):
    """의미 줄(잉크 14 pt)과 회색 설명 줄(12 pt). 기본은 맨 아래(설명 줄 y = CAP_Y)에 붙이고,
    그림이 높아 겹치면 그림 바로 아래(top)부터 차례로 놓는다(설명 줄 끝 ≤ 7.10 in)."""
    mh, ch = meaning_h(meaning), cap_h(t)
    y_m = CAP_Y - mh - 0.02 if meaning else CAP_Y
    if top is not None and top > y_m:
        y_m = top
    if meaning:
        ktext(sl, L.ML, y_m, W_BODY, mh, meaning, MEAN_PT, INK, F_M, align=PP_ALIGN.LEFT, line_spacing=1.08)
    if t:
        y_c = max(CAP_Y, y_m + mh + 0.02) if meaning else max(CAP_Y, y_m)
        ktext(sl, L.ML, y_c, W_BODY, ch, t, CAP_PT, GRAY, F_M, align=PP_ALIGN.LEFT)


def box_h(lines, w=W_BODY):
    n = sum(len(wrapped(ln, w - 0.36 - 0.2, BOX_PT)) for ln in lines)
    return LINE_IN[BOX_PT] * n + 0.07 * (len(lines) - 1) + 0.26


def accent_box(sl, x, y, w, h, lines):
    """주황 외곽선 상자(채움 없음). 소제목 없이 문장마다 문단 하나, 어절 단위 줄바꿈, 잉크 Medium 15 pt."""
    shp = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.background(); shp.line.color.rgb = ACC; shp.line.width = Pt(1.5); shp.shadow.inherit = False; shp.name = "accent box"
    tf = shp.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0.18); tf.margin_top = tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = "\n".join(wrapped(ln, w - 0.36 - 0.2, BOX_PT)); r.font.size = Pt(BOX_PT); r.font.color.rgb = INK; r.font.name = F_M
        fl._set_fonts(r, F_M); p.space_after = Pt(5)
    from lxml import etree  # noqa: F401
    L.break_runs(type("T", (), {"text_frame": tf})())
    return shp


def gap(sl, x, y, w, h, name):
    L.gap_marker(sl, x, y, w, h, ["그림 대기", name])


def fig_slide(sl, s):
    paths = [asset(a) for a in s["assets"]]
    box = s.get("box"); bh = box_h(box) if box else 0.0
    mh, ch = meaning_h(s.get("meaning")), cap_h(s.get("caption", ""))
    body_h = (YB - Y0) - (bh + 0.16 if box else 0.0) - (mh + 0.02 if mh else 0.0)
    top = None
    if all(paths) and len(paths) == 1:
        p = paths[0]; iw, ih = L.img_size(p); h_full = W_BODY * ih / iw
        room = 7.10 - Y0 - (bh + 0.16 if box else 0.0) - mh - ch - 0.06          # 의미·설명 줄을 그림 바로 아래로 내릴 때의 최대 높이
        h_use = h_full if h_full <= max(body_h, room) else max(body_h, room)
        w, h = L.fit(p, W_BODY, h_use); L.place(sl, p, L.ML + (W_BODY - w) / 2, Y0, w=w, h=h)
        s["_fig_fill"] = round(w / W_BODY, 3)
        if h > body_h + 1e-6 and not box:
            top = Y0 + h + 0.04
    elif not all(paths):
        gap(sl, L.ML, Y0, W_BODY, body_h, ", ".join(a for a, p in zip(s["assets"], paths) if p is None))
    else:
        nom = [L.img_size(p) for p in paths]; hh = body_h; ws = [hh * a / b for a, b in nom]; tot = sum(ws) + 0.3 * (len(ws) - 1)
        if tot > W_BODY:
            f = (W_BODY - 0.3 * (len(ws) - 1)) / (tot - 0.3 * (len(ws) - 1)); hh *= f; ws = [x * f for x in ws]; tot = W_BODY
        x = L.ML + (W_BODY - tot) / 2
        for p, w in zip(paths, ws):
            L.place(sl, p, x, Y0, w=w, h=hh); x += w + 0.3
    if box:
        accent_box(sl, L.ML, YB - bh, W_BODY, bh, box)
    caption(sl, s.get("caption", ""), s.get("meaning"), top=top)


def goal_slide(sl, s):
    """목표 쪽(2026-10-16 사용자 지시 1): 그림을 본문 위쪽에 크게(12 in 폭), 아래에 목표 한 문장을 주황 굵은 글씨(18 pt)로."""
    p = asset(s["assets"][0]); t = s["goal"]
    lines = wrapped(t, W_BODY - 0.2, 18, "Bold"); gh = 0.36 * len(lines) + 0.10
    y_goal = 7.02 - gh
    if p is not None:
        w, h = L.fit(p, W_BODY, y_goal - 0.14 - Y0); L.place(sl, p, L.ML + (W_BODY - w) / 2, Y0, w=w, h=h)
    tb = text(sl, L.ML, y_goal, W_BODY, gh, [run("\n".join(lines), 18, ACC, F_X)], align=PP_ALIGN.LEFT, line_spacing=1.10)
    L.break_runs(tb)


def claims_slide(sl, s):
    """결론 · 의의 쪽(2026-10-16 사용자 지시 12): 근거 그림 한 장 + 주장 카드(주황 세로 막대 + 굵은 머리 줄 + 근거 문장). 한 상자에 몰지 않는다.
    s["figure_side"] = "right"|"left", s["fig_w"] 그림 폭(in), s["cards"] = [{"head","body","y","h"}] (y·h 없으면 고르게 배분)."""
    p = asset(s["assets"][0]); fw = s.get("fig_w", 5.9); side = s.get("figure_side", "right"); top, bot = Y0, 6.95
    if p is not None:
        w, h = L.fit(p, fw, bot - top)
        fx = L.ML + W_BODY - w if side == "right" else L.ML
        L.place(sl, p, fx, top, w=w, h=h)
    else:
        fx = L.ML + W_BODY - fw if side == "right" else L.ML
        gap(sl, fx, top, fw, bot - top, s["assets"][0])
    cx = L.ML if side == "right" else L.ML + fw + 0.35
    cw = W_BODY - fw - 0.35
    cards = s["cards"]; n = len(cards); gap_c = 0.16
    ch_def = (bot - top - gap_c * (n - 1)) / n
    for i, c in enumerate(cards):
        y = c.get("y", top + i * (ch_def + gap_c)); hgt = c.get("h", ch_def)
        bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(y + 0.06), Inches(0.07), Inches(hgt - 0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background(); bar.shadow.inherit = False; bar.name = "claim bar"
        hl = wrapped(c["head"], cw - 0.45, 16, "Bold"); bl = wrapped(c["body"], cw - 0.45, 14)
        hh = 0.30 * len(hl) + 0.06
        tb = text(sl, cx + 0.22, y + 0.04, cw - 0.25, hh, [run("\n".join(hl), 16, INK, F_X)], line_spacing=1.05); L.break_runs(tb)
        tb = text(sl, cx + 0.22, y + 0.04 + hh + 0.02, cw - 0.25, hgt - hh - 0.08, [run("\n".join(bl), 14, INK, F_M)], line_spacing=1.10); L.break_runs(tb)


def cover(sl, s):
    """사용자 v8 수정본(2026-10-15 16:19)의 표지: 어두운 배경 + 흰 제목(32 pt, 112 %), 부제·이름 흰색, 소속 #C8D0DA."""
    bg = DECKFIG8 / "cover_bg_v8.png"; dark = bg.exists()
    if not dark:
        bg = ASSET / "cover_bg_v3.png"
    pic = sl.shapes.add_picture(str(bg), 0, 0, Inches(13.333), Inches(7.5)); pic.shadow.inherit = False
    white = RGBColor(0xFF, 0xFF, 0xFF); soft = RGBColor(0xC8, 0xD0, 0xDA)
    tb = text(sl, L.ML, 1.15, 6.9, 1.77, [run(s["title"], 32, white if dark else ACC, F_X)], anchor=MSO_ANCHOR.TOP, line_spacing=1.12)
    L.break_runs(tb)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(4)
    text(sl, L.ML, 3.75, 6.9, 0.45, [run(s["subtitle"], 18, white if dark else GRAY2, F_M)])
    text(sl, L.ML, 4.55, 6.9, 0.40, [run(s["presenter"], 18, white if dark else INK, F_S)])
    text(sl, L.ML, 4.98, 6.9, 0.34, [run(s["affiliation"], 14, soft if dark else GRAY2, F_M)])
    L.logos_row(sl, x_left=L.ML, y_bottom=6.90, h=L.LOGO_H_COVER, gap=L.LOGO_GAP_COVER)


def refs(sl, s):
    half = (len(s["refs"]) + 1) // 2
    for x, chunk in zip((L.COL[0], L.COL[3]), (s["refs"][:half], s["refs"][half:])):
        text(sl, x, L.Y0, 5.85, 5.2, [[run(r, 12, INK, F_M)] for r in chunk], line_spacing=1.15, space_after=6)


def build(out=PPTX):
    prs = fl.new_deck(); warn = []
    for i, s in enumerate(SPEC["slides"]):
        sl = fl.blank(prs)
        if s["type"] == "cover":
            cover(sl, s)
        elif s["type"] == "references":
            L.header(sl, s["title"], None, i); refs(sl, s)
        else:
            sub = s.get("subtitle")
            if sub and L.text_w_in(sub, 14) > W_BODY - 0.25:
                warn.append(f"p{i} 부제 한 줄 초과({L.text_w_in(sub, 14):.1f} in): {sub[:40]}…")
            L.header(sl, s["title"], sub, i)
            {"goal": goal_slide, "claims": claims_slide}.get(s["type"], fig_slide)(sl, s)
            if s.get("meaning") and len(wrapped(s["meaning"], W_BODY - 0.2, MEAN_PT)) > 2:
                warn.append(f"p{i} 의미 줄 3줄 이상")
            if s.get("_fig_fill") and s["_fig_fill"] < 0.97:
                warn.append(f"p{i} 그림 폭 {s['_fig_fill'] * 100:.0f}% (좌우 빈 띠)")
        L.notes(sl, s.get("notes"), None, None)
    RENDER.mkdir(parents=True, exist_ok=True); prs.save(out)
    print(f"[pptx] {out.relative_to(ROOT)} {len(prs.slides)} 쪽")
    for w in warn:
        print("  [layout]", w)
    if MISSING:
        print("  [macro 없음]", sorted(set(MISSING)))
    return prs


def audit():
    rep = L.pptx_audit(str(PPTX)); ext = dict(L.deck_extra_audit(str(PPTX))); hits = 0
    for i, v in rep:
        e = ext[i]; msgs = []
        for k in ("off_font", "off_size", "big_num"):
            if v.get(k):
                msgs.append(f"{k} {v[k]}")
        bad = [x for x in (v.get("pic_scale_bad") or []) if not (0.85 <= x <= 1.65)]
        if bad:
            msgs.append(f"pic_scale {bad}")
        if e.get("dash"):
            msgs.append(f"dash {e['dash']}")
        if msgs:
            hits += 1; print(f"  [audit] p{i}: {'; '.join(msgs)}")
    print(f"[audit] 지적 쪽 {hits}/{len(rep)}")


def render():
    L.export_pdf(PPTX, PDF, SCRATCH / "lo_v9")
    PAGES.mkdir(parents=True, exist_ok=True)
    for old in PAGES.glob("page-*.png"):
        old.unlink()
    subprocess.run(["nice", "-n", "10", "pdftoppm", "-r", "80", "-png", str(PDF), str(PAGES / "page")], check=True)
    print(f"[pdf] {PDF.relative_to(ROOT)} 쪽 PNG {len(list(PAGES.glob('page-*.png')))}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--no-pdf", action="store_true"); a = ap.parse_args()
    build(); audit()
    if not a.no_pdf:
        render()
