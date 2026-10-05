# -*- coding: utf-8 -*-
"""paper_report_lib.py: 논문 작성 전 보고 덱(문법 V)용 final_lib.py 확장.

final_lib.py 는 고치지 않고 감싼다. 바뀐 값(지침 5.0 '변경(보고)'):
여백 0.52 → 0.667 in(6열 격자), 결론 줄 헤어라인 y 6.82, 본문·표 18 pt.

담는 것
- 머리(제목, 부제, 쪽 번호), 결론 줄, 굵은 리드 줄, 패널 머리, 그림 대기 표지
- 논문 그림 v3 패널 크롭(사전 정의 상자, crops/ 에 저장)
- S13–S15 순차 공개 차트(Fig 2a·b 원천 값으로 단계별 다시 그림, 경로 B)
- 헤어라인 표(python-pptx 원어 표 객체, 채움 0, 세로선 0)
- 발표자 노트, pptx_audit(지침 부록 A.3 + 덱 금지 항목 확장), PDF 내보내기(자동 간격 끔)
"""
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import final_lib as fl

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "deck"
ASSET = DECK / "assets" / "paper_report"
CROP_DIR = ASSET / "crops"
PANEL_DIR = ASSET / "slide_panels"            # 슬라이드판 패널(글자 12 pt 이상, README.md 의 쪽별 대체표)
SRC_DIR = ASSET / "source_data"
V3 = ROOT / "outputs" / "figures" / "paper" / "v3_restructure"
FONT_DIR = Path.home() / ".fonts"

# ---------------------------------------------------------------- 격자(지침 5.2, 원형 1.1)
ML = 0.667
BODY_W = 12.0
Y0, H = 1.30, 5.20
COL = [0.667, 2.717, 4.767, 6.817, 8.867, 10.917]
COLR = [2.417, 4.467, 6.517, 8.567, 10.617, 12.667]
HAIR_Y, TAKE_Y = 6.82, 6.88

ACC, INK, GRAY, GRAY2, HAIR = fl.ACC, fl.INK, fl.GRAY, fl.GRAY2, fl.HAIR
F_X, F_S, F_M = fl.F_X, fl.F_S, fl.F_M
C_TOP = RGBColor(0x4A, 0x4A, 0x4A)


def span(c0, c1):
    """열 c0–c1(1부터)의 (x, w)."""
    return COL[c0 - 1], COLR[c1 - 1] - COL[c0 - 1]


# ---------------------------------------------------------------- 글자 폭 측정(줄 수 점검용)
_FONTS = {}


def text_w_in(s, pt=18, weight="Medium"):
    key = (weight, pt)
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(str(FONT_DIR / f"Pretendard-{weight}.otf"), size=pt * 10)
    return _FONTS[key].getlength(s) / 10.0 / 72.0


def run(t, size=18, color=INK, font=F_M):
    return {"t": t, "size": size, "color": color, "font": font}


def text(sl, x, y, w, h, paras, **kw):
    return fl.text(sl, x, y, w, h, paras, **kw)


# ---------------------------------------------------------------- 머리와 결론 줄(원형 1.2)
# ---------------------------------------------------------------- 로고(2026-10-05 조정 지시 11–12)
LOGO_DIR = DECK / "assets" / "final" / "logos"
LOGOS = ["sail.png", "snu.png", "kopri.png"]          # SAI'L, 서울대, 극지연구소(RGBA)
LOGO_H_PAGE, LOGO_H_COVER = 0.28, 0.45
LOGO_GAP_PAGE, LOGO_GAP_COVER = 0.18, 0.35
TITLE_BASELINE_Y = 0.24 + 0.93 * 26 / 72              # 제목 26 pt 의 기준선(약 0.58 in)


def _logo_w(name, h):
    iw, ih = img_size(LOGO_DIR / name)
    return h * iw / ih


def logos_row(sl, x_right=None, x_left=None, y_bottom=None, h=LOGO_H_PAGE, gap=LOGO_GAP_PAGE):
    """로고 3개를 한 줄로. x_right 를 주면 오른쪽 끝 정렬, x_left 를 주면 왼쪽 끝 정렬. 아래끝을 y_bottom 에 맞춘다."""
    ws = [_logo_w(n, h) for n in LOGOS]
    total = sum(ws) + gap * (len(ws) - 1)
    x = x_left if x_left is not None else x_right - total
    for n, w in zip(LOGOS, ws):
        pic = sl.shapes.add_picture(str(LOGO_DIR / n), Inches(x), Inches(y_bottom - h), Inches(w), Inches(h))
        pic.shadow.inherit = False
        pic.name = f"logo {n}"
        x += w + gap
    return total


def logos_topright(sl):
    """내용 쪽 공통: 오른쪽 위, 높이 0.28 in, 제목 기준선에 아래끝 정렬."""
    return logos_row(sl, x_right=COLR[5], y_bottom=TITLE_BASELINE_Y)


def pagenum_bottomright(sl, page):
    if page is None:
        return
    text(sl, COLR[5] - 1.0, 7.12, 1.0, 0.25, [run(str(page), 12, GRAY, F_M)], align=PP_ALIGN.RIGHT)


def header(sl, title, subtitle=None, page=None, logos=True):
    """제목(폭 9.3 in, 로고 영역 앞에서 끝남) + 부제 + 오른쪽 위 로고 3개 + 오른쪽 아래 쪽 번호."""
    text(sl, ML, 0.24, 9.30, 0.55, [run(title, 26, ACC, F_X)])
    if subtitle:
        text(sl, ML, 0.80, BODY_W, 0.34, [run(subtitle, 14, GRAY, F_M)])
    if logos:
        logos_topright(sl)
    pagenum_bottomright(sl, page)


def takeaway(sl, msg):
    fl.hline(sl, ML, HAIR_Y, BODY_W, color=HAIR, lw=0.9)
    text(sl, ML, TAKE_Y, BODY_W, 0.36, [run(msg, 15, INK, F_X)], align=PP_ALIGN.CENTER)


def break_runs(tb):
    """런 글자의 '\n' 을 줄바꿈(a:br)으로 바꾼다. 구 경계에서 직접 나누는 용도(지침 5.9)."""
    for p in tb.text_frame.paragraphs:
        for r in list(p.runs):
            if "\n" not in r.text:
                continue
            parts = r.text.split("\n")
            r.text = parts[0]
            anchor = r._r
            for part in parts[1:]:
                br = etree.Element(qn("a:br"))
                rpr = r._r.find(qn("a:rPr"))
                if rpr is not None:
                    br.append(etree.fromstring(etree.tostring(rpr)))
                anchor.addnext(br)
                nr = etree.fromstring(etree.tostring(r._r))
                for t in nr.iter(qn("a:t")):
                    t.text = part
                br.addnext(nr)
                anchor = nr
    return tb


def lead_wrap(lead, body, width_in, pt=18, safety=0.93):
    """리드 줄 본문을 공백(구 경계)에서만 나눈 줄바꿈 문자열로 바꾼다. 첫 줄은 리드 폭만큼 좁다.
    LibreOffice 는 한글 낱말 가운데서도 줄을 바꾸므로 렌더 전에 직접 나눈다(지침 5.9)."""
    if "\n" in body:
        return body
    maxw = width_in * safety
    lead_w = text_w_in(lead + "  ", pt, "SemiBold") if lead else 0.0
    t = nbsp(body)
    for k in KEEP:
        if k in t:
            t = t.replace(k, k.replace(" ", "\u00a0"))
    words = t.split(" ")
    lines, cur, avail = [], "", maxw - lead_w
    for w in words:
        cand = w if not cur else cur + " " + w
        if text_w_in(cand, pt) <= avail or not cur:
            cur = cand
        else:
            lines.append(cur)
            cur, avail = w, maxw
    if cur:
        lines.append(cur)
    return "\n".join(lines)


def lead_lines(sl, x, y, w, items, gap_pt=12, line_spacing=1.12):
    """굵은 리드 줄(지침 5.5): 줄 머리 SemiBold + 한 칸 + Medium 본문, 18 pt, 콜론 없음, 3줄 이하."""
    assert len(items) <= 3
    paras = [[run(lead + "  ", 18, INK, F_S), run(lead_wrap(lead, body, w), 18, INK, F_M)] for lead, body in items]
    return break_runs(text(sl, x, y, w, 2.6, paras, line_spacing=line_spacing, space_after=gap_pt))


def support_line(sl, x, y, w, body, lead=None, h=0.8):
    runs = ([run(lead + "  ", 18, INK, F_S)] if lead else []) + [run(body, 18, INK, F_M)]
    return break_runs(text(sl, x, y, w, h, [runs], line_spacing=1.1, space_after=0))


def panel_head(sl, x, y, w, t, align=PP_ALIGN.LEFT):
    # wrap=False 이면 LibreOffice 가 글을 상자 가운데에 놓으므로 줄바꿈을 켠 상자를 쓴다
    return text(sl, x, y, w, 0.34, [run(t, 18, INK, F_S)], align=align, wrap=True)


def gap_marker(sl, x, y, w, h, lines):
    """그림 대기 표지: 테두리·채움 없는 회색 글(16 pt), 예약 영역 가운데."""
    paras = [[run(t, 16, GRAY, F_M)] for t in lines]
    return text(sl, x, y, w, h, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                line_spacing=1.15, space_after=4)


def notes(sl, script=None, ref=None, extra=None):
    parts = []
    if script:
        parts.append(script)
    if ref:
        parts.append("[참고] " + ref)
    if extra:
        parts.extend(extra)
    sl.notes_slide.notes_text_frame.text = "\n\n".join(parts)


# ---------------------------------------------------------------- 그림 배치
def img_size(path):
    with Image.open(path) as im:
        return im.size


def place(sl, path, x, y, w=None, h=None):
    """정확한 크기로 넣는다. w 나 h 하나만 주면 종횡비를 지킨다. (w, h) 반환."""
    iw, ih = img_size(path)
    if w is None and h is None:
        with Image.open(path) as im:
            dpi = (im.info.get("dpi") or (300, 300))[0]
        w, h = iw / dpi, ih / dpi
    elif w is None:
        w = h * iw / ih
    elif h is None:
        h = w * ih / iw
    pic = sl.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.shadow.inherit = False
    return w, h


def fit(path, max_w, max_h):
    iw, ih = img_size(path)
    s = min(max_w / iw, max_h / ih)
    return iw * s, ih * s


# ---------------------------------------------------------------- 논문 그림 v3 패널 크롭
# 좌표는 2000 px 폭 기준(원본 4015 px 의 1/2.0075). pieces = [(원본 상자, 캔버스 위치)].
# 다른 패널의 글자 조각과 패널 문자(a, b, …)가 들어오지 않게 상자를 나누어 붙인다(자료 변경 없음).
CROPS = {
    "f1a": dict(src="Fig1", canvas=(1130, 1170), pieces=[
        ((30, 18, 1132, 1088), (0, 0)),           # 범북극 지도(패널 b 글자 'C' 제외)
        ((1132, 18, 1160, 225), (1102, 0)),       # 티베트 삽도 오른쪽 띠
        ((640, 1118, 1132, 1212), (610, 1076)),   # 크기 열쇠와 영구동토 열쇠(확대도 지시선 제외)
    ]),
    "f1b": dict(src="Fig1", canvas=(868, 1114), pieces=[
        ((1162, 34, 2000, 1148), (30, 0)),
        ((1132, 225, 1162, 1148), (0, 191)),      # 'Central Russia' 머리 글자(삽도 테두리 제외)
    ]),
    "f1c": dict(src="Fig1", box=(30, 1242, 558, 1772)),
    "f1d": dict(src="Fig1", box=(632, 1242, 1158, 1772)),
    "f3b": dict(src="Fig3", box=(795, 0, 1995, 606)),
    "f3c": dict(src="Fig3", canvas=(917, 606), pieces=[
        ((28, 786, 945, 1340), (0, 0)),
        ((890, 1345, 1110, 1392), (397, 557)),    # 가로축 이름(패널 c 축 가운데로)
    ]),
    "f4b": dict(src="Fig4", box=(1040, 28, 1895, 832)),
    "f4c": dict(src="Fig4", canvas=(925, 533), pieces=[    # 패널 c: 행 이름은 x 0 부터, 패널 문자(y 883–902)는 흰 조각으로 덮지 않고 잘라냄
        ((0, 903, 925, 1436), (0, 0)),                        # 열쇠 줄(y 903 부터)과 본문
    ]),
    "f4d": dict(src="Fig4", box=(1040, 902, 2000, 1436)),   # 재보정 몫 분리(2026-10-05 갱신판 Fig 4 패널 d)
    "f5a": dict(src="Fig5", box=(0, 50, 478, 526)),      # 재채색판(2026-10-05 17:29) 패널 a
    "f5c": dict(src="Fig5", box=(1010, 50, 1488, 526)),  # 패널 c
    "f5e": dict(src="Fig5", box=(0, 662, 1242, 1398)),
    "f7a_ak": dict(src="Fig7", box=(0, 40, 698, 455)),
    "f7a_ca": dict(src="Fig7", box=(1380, 40, 2000, 455)),
    "f7b": dict(src="Fig7", box=(0, 568, 1082, 1428)),
    "f7c": dict(src="Fig7", box=(1094, 572, 1962, 1418)),
}
# Fig 6 은 다른 작업에서 나온다(2026-10-05 02:19 생성 확인). 파일이 없으면 빌드가 그림 대기 표지를 둔다.
FIG6_CROPS = {
    "f6a": dict(src="Fig6", box=(0, 40, 965, 927)),        # 직접 ML 지도(영구동토 열쇠, 티베트·알래스카 삽도 포함)
    "f6b": dict(src="Fig6", box=(1025, 40, 1995, 927)),    # 물리 유사라벨 지도(같은 높이 = 같은 배율)
    "f6_cbar": dict(src="Fig6", box=(105, 930, 1672, 1054)),  # 공유 컬러바와 방향 표지
    "f6c": dict(src="Fig6", box=(0, 1108, 1105, 1939)),    # 학습기 포레스트(기존 ALT 지도 묶음 포함, 2026-10-05 갱신판)
}


def _scale(box, f):
    return tuple(int(round(v * f)) for v in box)


def make_crop(name, spec=None):
    """크롭 PNG 를 crops/ 에 만들고 경로를 돌려준다. dpi 정보는 원본 값(600)을 그대로 둔다."""
    spec = spec or CROPS.get(name) or FIG6_CROPS.get(name)
    src = V3 / f"{spec['src']}.png"
    CROP_DIR.mkdir(parents=True, exist_ok=True)
    out = CROP_DIR / f"{name}.png"
    with Image.open(src) as im0:
        im = im0.convert("RGB")
        dpi = im0.info.get("dpi", (600, 600))
    f = im.size[0] / 2000.0
    if "box" in spec:
        b = _scale(spec["box"], f)
        b = (max(0, b[0]), max(0, b[1]), min(im.size[0], b[2]), min(im.size[1], b[3]))
        res = im.crop(b)
    else:
        cw, ch = _scale(spec["canvas"], f)
        res = Image.new("RGB", (cw, ch), "white")
        for box, (dx, dy) in spec["pieces"]:
            b = _scale(box, f)
            b = (max(0, b[0]), max(0, b[1]), min(im.size[0], b[2]), min(im.size[1], b[3]))
            res.paste(im.crop(b), _scale((dx, dy), f))
    res.save(out, dpi=(round(dpi[0]), round(dpi[1])))
    return out


# ---------------------------------------------------------------- S13–S15 순차 공개 차트(경로 B)
# v4 토큰(design/style_tokens_v4.json): 방법 색, Pretendard Regular, 슬라이드 선 굵기. 차트 양식 부분만 바꿨다(배치 코드는 그대로)
import json as _json
_TOK = _json.loads((Path(__file__).resolve().parent.parent / "design" / "style_tokens_v4.json").read_text(encoding="utf-8"))
_TMc, _TLs, _TFs = _TOK["color"]["methods"], _TOK["lines"]["slide"], _TOK["font"]["slide"]
METHOD = {"source": _TMc["P0"]["hex"], "recal": _TMc["P1"]["hex"], "resid": _TMc["R1"]["hex"], "direct": _TMc["D0"]["hex"],
          "ym": _TMc["Pstar"]["hex"]}
_INK = _TOK["color"]["neutral"]["text_slide"]
EMPH = RGBColor.from_string(_TMc["R1"]["hex"].lstrip("#"))     # 제안 방법 글자 강조(주홍 #D55E00), 쪽당 한 곳 이하


def _setup_mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager as fm
    fm.fontManager.addfont(str(FONT_DIR / "Pretendard-Regular.otf"))
    plt.rcParams.update({
        "font.family": "Pretendard", "font.weight": "normal", "font.size": _TFs["direct_label_pt"],
        "text.color": _INK, "axes.labelcolor": _INK, "axes.edgecolor": _INK,
        "axes.labelsize": _TFs["axis_label_pt"], "axes.labelweight": "normal", "axes.linewidth": _TLs["axis_pt"],
        "axes.spines.top": False, "axes.spines.right": False, "axes.grid": False,
        "axes.unicode_minus": True, "axes.facecolor": "white", "figure.facecolor": "white",
        "xtick.labelsize": _TFs["tick_pt"], "ytick.labelsize": _TFs["tick_pt"], "xtick.color": _INK, "ytick.color": _INK,
        "xtick.major.width": _TLs["tick_pt"], "ytick.major.width": _TLs["tick_pt"], "xtick.major.size": _TLs["tick_len_pt"],
        "ytick.major.size": _TLs["tick_len_pt"],
        "xtick.major.pad": 6, "ytick.major.pad": 8, "xtick.minor.visible": False, "ytick.minor.visible": False,
        "savefig.dpi": 300, "figure.dpi": 300, "lines.solid_capstyle": "butt",
    })
    return plt


def make_label_curve_stages(out_dir=ASSET):
    """Fig 2a·b 원천 값(Fig2_source_data.csv)으로 순차 공개 3단계를 같은 축·같은 위치로 그린다.
    1단계 0선(원천 계수 Stefan), 연도 정합 Stefan, 직접 ML. 2단계 + 재보정 Stefan. 3단계 + 물리 잔차 결합."""
    import numpy as np
    import pandas as pd
    from matplotlib.lines import Line2D
    from matplotlib.ticker import FixedLocator, NullLocator
    plt = _setup_mpl()
    d = pd.read_csv(V3 / "Fig2_source_data.csv")
    d = d[d.panel.isin(["a", "b"])].copy()
    SRC_DIR.mkdir(parents=True, exist_ok=True)
    d.to_csv(SRC_DIR / "S13_S15_label_curve.csv", index=False, encoding="utf-8")
    ser = {"Year-matched Stefan": "ym", "Recalibrated Stefan": "recal", "Anchor + residual ML": "resid",
           "Direct ML": "direct"}
    W, Hh = 12.0, 5.2
    YL = (-4.0, 4.0)
    yb, yh = 1.00, 3.55               # 자료 영역 y 1.00–4.55 in
    lay = {"a": dict(n0=(1.45, 0.40), log=(2.15, 2.30), lim=(2.45, 12.5), ticks=[3, 10], all=None),
           "b": dict(n0=(5.25, 0.40), log=(5.95, 3.40), lim=(2.45, 420), ticks=[3, 10, 40, 160, 320],
                     all=(9.65, 0.42))}
    head = {"a": "주 4지역 평균", "b": "레나델타·캐나다 평균"}
    outs = []
    for stage in (1, 2, 3):
        fig = plt.figure(figsize=(W, Hh), dpi=300)
        tr = fig.dpi_scale_trans
        axes = {}
        for p, L in lay.items():
            def ax_at(x, w):
                return fig.add_axes([x / W, yb / Hh, w / W, yh / Hh])
            a0 = ax_at(*L["n0"])
            am = ax_at(*L["log"])
            aa = ax_at(*L["all"]) if L["all"] else None
            for a in (a0, am, aa):
                if a is None:
                    continue
                a.set_ylim(*YL)
                a.axhline(0, color=METHOD["source"], lw=_TLs["data_pt"], zorder=1)
            for a in (am, aa):
                if a is None:
                    continue
                a.spines["left"].set_visible(False)
                a.tick_params(axis="y", left=False, labelleft=False)
            a0.set_yticks([-4, -2, 0, 2, 4])
            if p == "a":
                a0.set_yticklabels(["−4", "−2", "0", "2", "4"])
            else:
                a0.set_yticklabels([])
            a0.set_xlim(0, 1)
            a0.set_xticks([0.5])
            a0.set_xticklabels(["0"])
            am.set_xscale("log")
            am.set_xlim(*L["lim"])
            am.xaxis.set_major_locator(FixedLocator(L["ticks"]))
            am.xaxis.set_minor_locator(NullLocator())
            am.set_xticklabels([str(t) for t in L["ticks"]])
            if aa is not None:
                aa.set_xlim(0, 1)
                aa.set_xticks([0.5])
                aa.set_xticklabels(["전량"])
            axes[p] = (a0, am, aa)
            fig.text(((L["n0"][0]) + (L["all"][0] + L["all"][1] if L["all"] else L["log"][0] + L["log"][1])) / 2 / W,
                     (yb + yh + 0.38) / Hh, head[p], ha="center", va="center", fontsize=_TFs["direct_label_pt"])
        fig.text(0.42 / W, (yb + yh / 2) / Hh, "원천 계수 Stefan 대비\n오차 변화 (cm)", rotation=90,
                 ha="center", va="center", fontsize=_TFs["axis_label_pt"], linespacing=1.25)
        fig.text(5.55 / W, 0.22 / Hh, "대상 라벨 수, n", ha="center", va="center", fontsize=_TFs["axis_label_pt"])

        def rows(p, s):
            r = d[(d.panel == p) & (d.series.map(ser) == s)]
            return r

        def draw(p, s, color, ls, band):
            a0, am, aa = axes[p]
            r = rows(p, s)
            r0 = r[r.x_axis == "n = 0 axis"]
            rl = r[r.x_axis == "log axis"]
            ra = r[r.x_axis == "All axis"]
            off = {"recal": -0.10, "resid": 0.10, "direct": 0.0, "ym": 0.0}[s]
            for rr, ax in ((r0, a0), (ra, aa)):
                if ax is None or rr.empty:
                    continue
                v = float(rr.error_change_cm.iloc[0])
                lo, hi = float(rr.ci95_low_cm.iloc[0]), float(rr.ci95_high_cm.iloc[0])
                if band and str(rr.ci_drawn.iloc[0]) == "yes":
                    ax.fill_between([0.5 + off - 0.09, 0.5 + off + 0.09], [lo, lo], [hi, hi], color=color,
                                    alpha=0.20, lw=0, zorder=2)
                ax.plot([0.5 + off], [v], marker="o", ms=_TLs["marker_pt"], mfc=color, mec=color, ls="none", zorder=4)
            if not rl.empty:
                x = rl.target_labels_n.astype(float).values
                y = rl.error_change_cm.values
                if band and (rl.ci_drawn == "yes").all():
                    am.fill_between(x, rl.ci95_low_cm.values, rl.ci95_high_cm.values, color=color, alpha=0.20,
                                    lw=0, zorder=2)
                am.plot(x, y, color=color, lw=_TLs["data_pt"], ls=ls, zorder=3, dash_capstyle="butt")
                am.plot(x, y, marker="o", ms=_TLs["marker_pt"], mfc=color, mec=color, ls="none", zorder=4)

        def ym(p, label):
            a0 = axes[p][0]
            r = rows(p, "ym")
            v = float(r.error_change_cm.iloc[0])
            a0.plot([0.12, 0.88], [v, v], color=METHOD["ym"], lw=_TLs["data_pt"], ls=(0, (1.0, 1.4)), zorder=3)
            lo, hi = r.ci95_low_cm.iloc[0], r.ci95_high_cm.iloc[0]
            if str(r.ci_drawn.iloc[0]) == "yes" and np.isfinite(lo):
                a0.plot([0.5, 0.5], [lo, hi], color=METHOD["ym"], lw=_TLs["data_aux_pt"], zorder=3)
            if label:
                yy = a0.transData.transform((0, v))[1] / fig.dpi
                xx = a0.get_window_extent(fig.canvas.get_renderer()).x1 / fig.dpi + 0.12
                fig.text(xx / W, yy / Hh, "연도 정합 Stefan", ha="left", va="center", fontsize=_TFs["direct_label_pt"],
                         color=METHOD["ym"])

        fig.canvas.draw()
        for p in ("a", "b"):
            draw(p, "direct", METHOD["direct"], (0, (2, 1.5)), band=False)
            ym(p, label=(p == "b"))
            if stage >= 2:
                draw(p, "recal", METHOD["recal"], "-", band=True)
            if stage >= 3:
                draw(p, "resid", METHOD["resid"], "-", band=True)
        # 직접 라벨: 오른쪽 전량 축 옆, 각 계열 1회(16 pt SemiBold)
        aa = axes["b"][2]
        xl = (lay["b"]["all"][0] + lay["b"]["all"][1] + 0.18)
        labs = [("direct", "직접 ML")]
        if stage >= 2:
            labs.append(("recal", "재보정 Stefan"))
        if stage >= 3:
            labs.append(("resid", "물리 잔차 결합"))
        for s, name in labs:
            r = rows("b", s)
            v = float(r[r.x_axis == "All axis"].error_change_cm.iloc[0])
            yy = aa.transData.transform((0, v))[1] / fig.dpi
            fig.text(xl / W, yy / Hh, name, ha="left", va="center", fontsize=_TFs["direct_label_pt"], color=METHOD[s])
        out = out_dir / f"S{12 + stage}_label_curve_stage{stage}.png"
        fig.savefig(out, dpi=300, facecolor="white")
        plt.close(fig)
        outs.append(out)
    return outs


# ---------------------------------------------------------------- 헤어라인 표(원어 표 객체)
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid


def _ln(tcPr, tag, w_pt=None, color=None):
    ln = etree.SubElement(tcPr, qn(tag))
    if w_pt is None:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(int(round(w_pt * 12700))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", color)
        dsh = etree.SubElement(ln, qn("a:prstDash"))
        dsh.set("val", "solid")
    return ln


def _cell_props(cell, top=None, bottom=None, marL=0.0, marR=0.0, marT=0.03, marB=0.03):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for ch in list(tcPr):
        tcPr.remove(ch)
    tcPr.set("marL", str(int(Inches(marL))))
    tcPr.set("marR", str(int(Inches(marR))))
    tcPr.set("marT", str(int(Inches(marT))))
    tcPr.set("marB", str(int(Inches(marB))))
    tcPr.set("anchor", "ctr")
    _ln(tcPr, "a:lnL")
    _ln(tcPr, "a:lnR")
    _ln(tcPr, "a:lnT", *(top or (None, None)))
    _ln(tcPr, "a:lnB", *(bottom or (None, None)))
    etree.SubElement(tcPr, qn("a:noFill"))


def _split_emph(line, words):
    """줄을 강조 낱말(줄바꿈 없는 공백 판 포함)과 나머지 조각으로 나눈다."""
    segs = [(line, False)]
    for w in words:
        for form in (w, w.replace(" ", "\u00a0")):
            out = []
            for seg, e in segs:
                if e or form not in seg:
                    out.append((seg, e))
                    continue
                parts = seg.split(form)
                for k, part in enumerate(parts):
                    if part:
                        out.append((part, False))
                    if k < len(parts) - 1:
                        out.append((form, True))
            segs = out
    return segs


def equation(sl, x, y, w, parts, size=18, h=0.40, align=PP_ALIGN.LEFT):
    """수식 한 줄: parts = [(글자, 'sub'|'lead'|'')]. 아래 첨자는 baseline −25%, 리드는 SemiBold."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for t, kind in parts:
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.color.rgb = INK
        fl._set_fonts(r, F_S if kind == "lead" else F_M)
        if kind == "sub":
            r._r.get_or_add_rPr().set("baseline", "-25000")
    return tb


def hairline_table(sl, t, x=None, y=None, row_h=None, header_h=None):
    """원어 표: 위·아래 1.25 pt #4A4A4A, 머리 아래 0.75 pt #DDDDDD, 머리 18 pt SemiBold #4A4A4A,
    셀 18 pt Medium #111111, 세로 가운데. 줄은 표 JSON 의 낱말 단위 줄(wrapped_*)을 문단으로 넣는다.
    돌려주는 값: 표 아래끝 y(in)."""
    g = t["geometry_in"]
    x = g["x"] if x is None else x
    y = g["y"] if y is None else y
    rh = g["row_h"] if row_h is None else row_h
    hh = (g.get("header_h") or t["style"]["header_h"]) if header_h is None else header_h
    cw = g["col_w"]
    nr, nc = len(t["wrapped_rows"]) + 1, len(cw)
    rhs = list(rh) if isinstance(rh, (list, tuple)) else [rh] * (nr - 1)
    total_h = hh + sum(rhs)
    gf = sl.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(cw)), Inches(total_h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE
    for j, w in enumerate(cw):
        tbl.columns[j].width = Inches(w)
    tbl.rows[0].height = Inches(hh)
    for i in range(1, nr):
        tbl.rows[i].height = Inches(rhs[i - 1])
    gap = t["style"].get("col_gap", 0.18)
    acc = {tuple(a) for a in t.get("accent_cells", [])}
    top_line = (1.25, "4A4A4A")
    mid_line = (0.75, "DDDDDD")
    for i in range(nr):
        lines_row = t["wrapped_header"] if i == 0 else t["wrapped_rows"][i - 1]
        for j in range(nc):
            cell = tbl.cell(i, j)
            al = t["align"][j]
            top = top_line if i == 0 else (mid_line if i == 1 else None)
            bottom = mid_line if i == 0 else (top_line if i == nr - 1 else None)
            _cell_props(cell, top=top, bottom=bottom,
                        marL=(gap if al == "r" else 0.0), marR=(0.0 if al == "r" else gap))
            tf = cell.text_frame
            tf.word_wrap = True
            if i == 0:
                font, colr = F_S, C_TOP
            else:
                font, colr = F_M, (ACC if (i - 1, j) in acc else INK)
            lines = lines_row[j]
            emph = [] if i == 0 else t.get("emph", [])
            for k, ln in enumerate(lines):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.RIGHT if al == "r" else PP_ALIGN.LEFT
                p.line_spacing = 1.0
                p.space_after = Pt(0)
                p.space_before = Pt(0)
                for seg, is_emph in _split_emph(ln, emph):
                    r = p.add_run()
                    r.text = seg
                    r.font.size = Pt(18)
                    r.font.color.rgb = EMPH if is_emph else colr
                    fl._set_fonts(r, font)
    return y + total_h


def table_support(sl, t, x, y, w):
    s = t.get("support_line")
    if not s:
        return y
    support_line(sl, x, y, w, s["text"], lead=s.get("lead") or None)
    return y


# ---------------------------------------------------------------- 스펙 표 → 표 정의(낱말 단위 줄 나눔, 지침 5.9)
TBL_STYLE = {"header_h": 0.5, "col_gap": 0.18, "safety": 0.95, "line_h": 0.3}
NBSP_UNITS = r"(cm|단어|in|pt|km|m|%|쪽|편|개|종|곳)"


def nbsp(s):
    return re.sub(r"(\d)\s" + NBSP_UNITS + r"(?![가-힣A-Za-z])", lambda m: m.group(1) + "\u00a0" + m.group(2), s)


KEEP = ["원천 계수 Stefan", "연도 정합 Stefan", "재보정 Stefan", "물리 잔차 결합", "잔차 ML", "직접 ML", "CCI v5", "사전 지정",
        "학습형 정책", "두 가설", "한 가설", "두 가중", "두 방법", "마감 10-08", "마감 10-11", "10-05 오전", "오차 감소",
        "오차 증가"]
DETERMINERS = {"두", "세", "네", "각", "첫", "새", "큰", "그", "이", "전", "한", "온", "총"}


def wrap_text(s, width_in, pt=18, weight="Medium"):
    """공백에서만 나누는 낱말 단위 줄바꿈(지침 5.9). 용어(KEEP)와 숫자·단위는 묶는다.
    두 줄이면 쉼표 뒤 구 경계를 우선하는 균형 분할, 관형어(두, 각 …)로 줄을 끝내지 않음.
    세 줄 이상이면 두 글자 이하 마지막 줄을 앞 줄 끝 낱말과 합친다."""
    if "\n" in s:                                   # 칸 문구의 줄바꿈 문자는 직접 나눈 구 경계로 따른다
        return [ln for part in s.split("\n") for ln in wrap_text(part, width_in, pt, weight)]
    maxw = width_in * TBL_STYLE["safety"]
    W = lambda x: text_w_in(x, pt, weight)  # noqa: E731
    t = nbsp(s)
    for k in KEEP:
        if k in t and W(k) <= maxw:
            t = t.replace(k, k.replace(" ", "\u00a0"))
    words = t.split(" ")
    lines, cur = [], ""
    for w in words:
        cand = w if not cur else cur + " " + w
        if W(cand) <= maxw or not cur:
            cur = cand
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    if len(lines) == 2:
        best = None
        for k in range(1, len(words)):
            a, b = " ".join(words[:k]), " ".join(words[k:])
            if W(a) > maxw or W(b) > maxw:
                continue
            score = max(W(a), W(b)) - (0.15 * maxw if words[k - 1].endswith(",") else 0)
            if len(b.replace("\u00a0", "")) <= 2:
                score += maxw
            if words[k - 1] in DETERMINERS:
                score += maxw
            if best is None or score < best[0]:
                best = (score, [a, b])
        if best:
            lines = best[1]
    elif len(lines) >= 3:
        last = lines[-1]
        if len(last.replace("\u00a0", "").strip()) <= 2 and " " in lines[-2]:
            head, tail = lines[-2].rsplit(" ", 1)
            if W(tail + " " + last) <= maxw:
                lines[-2], lines[-1] = head, tail + " " + last
    return lines


def row_heights(t, min_h=0.70, line_h=0.30, pad=0.12):
    """행마다 가장 긴 칸의 줄 수로 높이를 정한다(18 pt 한 줄 0.30 in)."""
    return [max(min_h, max(len(c) for c in r) * line_h + pad) for r in t["wrapped_rows"]]


def table_def(columns, rows, geo, align=None, accent=(), support=None, emph=None):
    """스펙 표(columns, rows, geometry_in)로 hairline_table 이 받는 정의를 만든다."""
    cw = geo["col_w"]
    gap = TBL_STYLE["col_gap"]
    align = align or ["l"] * len(columns)
    return dict(geometry_in=dict(geo), columns=columns, rows=rows, align=align, emph=list(emph or []),
                accent_cells=[list(a) for a in accent], style=dict(TBL_STYLE), support_line=support,
                wrapped_header=[wrap_text(c, cw[j] - gap, 18, "SemiBold") for j, c in enumerate(columns)],
                wrapped_rows=[[wrap_text(c, cw[j] - gap) for j, c in enumerate(r)] for r in rows])


# ---------------------------------------------------------------- 점검(지침 부록 A.3 그대로 + 확장)
EMU = 914400
COL_X = COL
SIZES = {"V": {26, 18, 16, 15, 14, 12}, "N": {22, 18, 16, 12}}
COVER = {32, 28, 18, 16, 14}        # 28: 표지 제목 두 줄(사용자 요청, 2026-10-05 보고 항목)
TITLE_PT = {"V": 26, "N": 18}
ACCENT = "EA851B"
NUM = re.compile(r"^[\s\d.,+\-−%×~]+(cm|°C|%|m)?$")
CODE_RE = re.compile(r"\b(L\d{1,2}[a-z]?|AB\d{1,2}|LG[A-Z]?(-[A-Z]\d)?|WF\d{1,2}|SC\d\w*|H\d{1,2}\w?|X[A-J]|P4|PE[12]"
                     r"|P[0-2]\*?|R[0-3]|D[01]|F1[kn]?|S[1-8]|AL-\d|CA-\d|LE-\d|MEAN\d?|(?<!CCI )(?<!CCI\u00a0)v[1-5])\b")
DASH_RE = re.compile(r"(?<![0-9])\s?[—–]\s?(?![0-9])|—")


def has_fill(sh):
    el = sh._element
    spPr = el.find(qn("p:spPr"))
    if spPr is None:
        return False
    if spPr.find(qn("a:noFill")) is not None:
        return False
    if any(spPr.find(qn(t)) is not None for t in ("a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill")):
        return True
    st = el.find(qn("p:style"))
    if st is not None:
        fr = st.find(qn("a:fillRef"))
        if fr is not None and fr.get("idx", "0") != "0":
            return True
    return False


def _runs(sh):
    if not sh.has_text_frame:
        return []
    return [r for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]


def pptx_audit(path, grammar="V", tol=0.03, cover_index=1):
    """지침 부록 A.3 pptx_audit 전문(동작 동일)."""
    from pptx import Presentation
    prs = Presentation(path)
    rep = []
    for i, s in enumerate(prs.slides, 1):
        allowed = COVER if i == cover_index else SIZES[grammar]
        v = dict(fill_shapes=0, off_font=set(), off_size=set(), accent_outside_title=0, big_num=0,
                 off_grid=0, pic_scale=[], para_end_da=0)
        for sh in s.shapes:
            x = sh.left / EMU if sh.left is not None else None
            if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER) \
                    and has_fill(sh):
                v["fill_shapes"] += 1
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                if sh.height is not None and sh.height / EMU <= 0.5 and str(sh.name).startswith("logo"):
                    continue                                   # 로고는 배율 점검 대상이 아님(조정 지시 11–12)
                try:
                    im = Image.open(io.BytesIO(sh.image.blob))
                    dpi = (im.info.get("dpi") or (300, 300))[0]
                    v["pic_scale"].append(round((sh.width / EMU) / (im.size[0] / dpi), 2))
                except Exception:
                    pass
            if x is not None and sh.width is not None and sh.width / EMU > 1.0 and sh.width / EMU < 13.0 and \
                    not any(abs(x - c) < tol for c in COL_X):
                v["off_grid"] += 1                          # 전면 바탕 그림(폭 13.3 in)은 격자 점검 대상이 아님
            for r in _runs(sh):
                f = r.font
                if f.name and not f.name.startswith("Pretendard"):
                    v["off_font"].add(f.name)
                if f.size is not None and round(f.size.pt) not in allowed:
                    v["off_size"].add(round(f.size.pt, 1))
                try:
                    col = str(f.color.rgb) if f.color and f.color.type is not None else ""
                except Exception:
                    col = ""
                is_title = f.size is not None and round(f.size.pt) in (TITLE_PT[grammar], 32, 28)   # 28: 표지 제목 두 줄
                if col.upper() == ACCENT and not is_title:
                    v["accent_outside_title"] += 1
                if f.size is not None and f.size.pt >= 24 and NUM.match(r.text.strip()):
                    v["big_num"] += 1
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = p.text.strip()
                    if t and re.search(r"(다|니다)\.?$", t):
                        v["para_end_da"] += 1
        v["pic_scale_bad"] = [k for k in v["pic_scale"] if abs(k - 1.0) > 0.02]
        v["accent_fail"] = v["accent_outside_title"] > 1
        rep.append((i, v))
    return rep


def deck_extra_audit(path, cover_index=1):
    """A.3 이 보지 않는 항목: 표 셀 글자(글꼴·크기·주황), 글씨 든 도형, 내부 코드, 연결어 대시,
    글머리 목록 수, 화살표·곡선 연결선, 쪽당 주황 1곳(표 포함)."""
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, s in enumerate(prs.slides, 1):
        allowed = COVER if i == cover_index else SIZES["V"]
        v = dict(table_off=set(), table_accent=0, boxed_text=0, codes=[], dash=[], bullets=0, connectors_bad=0,
                 texts=[])
        for sh in s.shapes:
            texts = []
            if sh.has_text_frame:
                texts = [p.text for p in sh.text_frame.paragraphs]
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and any(t.strip() for t in texts):
                    v["boxed_text"] += 1
                for p in sh.text_frame.paragraphs:
                    if p.text.strip().startswith(("●", "•", "-", "·", "▪")):
                        v["bullets"] += 1
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        for p in c.text_frame.paragraphs:
                            texts.append(p.text)
                            for r in p.runs:
                                if not r.text.strip():
                                    continue
                                f = r.font
                                if (f.name and not f.name.startswith("Pretendard")) or \
                                        (f.size is not None and round(f.size.pt) not in allowed):
                                    v["table_off"].add((f.name, f.size.pt if f.size else None))
                                try:
                                    if str(f.color.rgb).upper() == ACCENT:
                                        v["table_accent"] += 1
                                except Exception:
                                    pass
            if sh._element.tag == qn("p:cxnSp"):
                ln = sh._element.find(".//" + qn("a:ln"))
                if ln is not None and (ln.find(qn("a:headEnd")) is not None or ln.find(qn("a:tailEnd")) is not None):
                    v["connectors_bad"] += 1
                prst = sh._element.find(".//" + qn("a:prstGeom"))
                if prst is not None and "curved" in prst.get("prst", "").lower():
                    v["connectors_bad"] += 1
            for t in texts:
                if not t.strip():
                    continue
                v["texts"].append(t)
                for m in CODE_RE.finditer(t):
                    v["codes"].append(m.group(0))
                if DASH_RE.search(re.sub(r"수[십백천]–수[십백천]", "", t)):   # 수 범위(수백–수천)는 허용
                    v["dash"].append(t[:40])
        out.append((i, v))
    return out


# ---------------------------------------------------------------- PDF 내보내기(지침 5.6, 자동 간격 끔)
def export_pdf(pptx_path, pdf_path, work_dir):
    """PPTX → ODP → 모든 문단 속성에 style:text-autospace='none' → PDF(soffice --headless --convert-to)."""
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="pr_", dir=str(work_dir)))
    prof = f"file://{tmp}/lo_profile"
    soff = shutil.which("soffice")
    env = dict(os.environ, OMP_NUM_THREADS="2")
    pre = ["nice", "-n", "10"]
    subprocess.run(pre + [soff, f"-env:UserInstallation={prof}", "--headless", "--convert-to", "odp", "--outdir",
                          str(tmp), str(pptx_path)], check=True, capture_output=True, timeout=900, env=env)
    odp = tmp / (Path(pptx_path).stem + ".odp")
    patched = tmp / (Path(pptx_path).stem + ".odp.tmp")
    with zipfile.ZipFile(odp) as zin, zipfile.ZipFile(patched, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ("styles.xml", "content.xml"):
                s = data.decode("utf-8")
                s = re.sub(r'style:text-autospace="[^"]*"', "", s)
                s = s.replace("<style:paragraph-properties",
                              '<style:paragraph-properties style:text-autospace="none"')
                data = s.encode("utf-8")
            ct = zipfile.ZIP_STORED if item.filename == "mimetype" else zipfile.ZIP_DEFLATED
            zout.writestr(item, data, compress_type=ct)
    final_odp = tmp / "na" / (Path(pptx_path).stem + ".odp")
    final_odp.parent.mkdir()
    shutil.move(str(patched), final_odp)
    # 그림을 무손실로 넣는다(기본 JPEG 압축은 흰 배경을 254 로 바꿔 그림 둘레에 옅은 회색 상자가 보인다).
    # LibreOffice 6.0 은 명령줄 JSON 옵션이 없어 임시 프로필의 PDF 내보내기 설정을 바꾼다.
    reg = tmp / "lo_profile" / "user" / "registrymodifications.xcu"
    if reg.exists():
        items = "".join(
            f'<item oor:path="/org.openoffice.Office.Common/Filter/PDF/Export"><prop oor:name="{k}" oor:op="fuse">'
            f"<value>{v}</value></prop></item>"
            for k, v in (("UseLosslessCompression", "true"), ("ReduceImageResolution", "false")))
        x = reg.read_text(encoding="utf-8")
        reg.write_text(x.replace("</oor:items>", items + "</oor:items>"), encoding="utf-8")
    subprocess.run(pre + [soff, f"-env:UserInstallation={prof}", "--headless", "--convert-to", "pdf", "--outdir",
                          str(tmp), str(final_odp)], check=True, capture_output=True, timeout=900, env=env)
    shutil.move(str(tmp / (Path(pptx_path).stem + ".pdf")), pdf_path)
    shutil.rmtree(tmp, ignore_errors=True)
    return pdf_path
