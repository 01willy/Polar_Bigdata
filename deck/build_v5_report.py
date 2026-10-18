# -*- coding: utf-8 -*-
"""build_v5_report.py: v5 연구 결과 보고 덱(국문, 16:9, 문법 V). 명세 deck/deck_spec_v5.json, 그림 outputs/figures/paper/v5/, 자산 deck/assets/paper_report/.
출력 deck/render/v5_report.pptx, .pdf, deck/render/v5_report_pages/page-XX.png. 실행: cd deck && OMP_NUM_THREADS=2 nice -n 10 python3 build_v5_report.py [--no-pdf]
"""
from __future__ import annotations
import argparse, json, os, subprocess, sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.util import Inches  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
import final_lib as fl  # noqa: E402
import paper_report_lib as L  # noqa: E402
from paper_report_lib import ACC, F_M, F_S, F_X, GRAY, GRAY2, INK, run, text  # noqa: E402

ROOT = L.ROOT; DECK = L.DECK; RENDER = DECK / "render"; PAGES = RENDER / "v5_report_pages"
PPTX = RENDER / "v5_report.pptx"; PDF = RENDER / "v5_report.pdf"
FIG = ROOT / "outputs" / "figures" / "paper" / "v5"
ASSET = DECK / "assets" / "paper_report"
SPEC = json.load(open(DECK / "deck_spec_v5.json", encoding="utf-8"))
SCRATCH = Path(os.environ.get("PR_TMP", tempfile.gettempdir()))


def asset(name):
    for p in (FIG / "slide" / name, FIG / name, ASSET / name):
        if p.exists():
            return p
    raise FileNotFoundError(name)


def place_fig(sl, path, max_w=12.0, max_h=L.H, y=L.Y0):
    """A2: 폭 12.0 안, 높이 H 안에 맞춰 왼쪽 열(x 0.667)에 둔다."""
    w, h = L.fit(path, max_w, max_h)
    L.place(sl, path, L.ML, y, w=w, h=h)


def table_slide(sl, s):
    t = s["table"]
    geo = dict(x=L.ML, y=L.Y0 + 0.15, col_w=t["col_w"], row_h=t["row_h"])
    td = L.table_def(t["columns"], t["rows"], geo)
    L.hairline_table(sl, td, y=geo["y"], row_h=t["row_h"])


def cover(sl, s):
    bg = ASSET / "cover_bg_v3.png"
    pic = sl.shapes.add_picture(str(bg), 0, 0, Inches(13.333), Inches(7.5)); pic.shadow.inherit = False
    title = s["title"].replace(" 본 ", " 본\n", 1)
    L.break_runs(text(sl, L.ML, 0.95, 7.9, 1.30, [run(title, 28, ACC, F_X)], anchor=MSO_ANCHOR.TOP, line_spacing=1.1))
    text(sl, L.ML, 2.55, 7.0, 0.70, [run(s["subtitle"], 18, GRAY2, F_M)])
    text(sl, L.ML, 3.60, 5.0, 0.36, [run(s["presenter"], 16, INK, F_S)])
    text(sl, L.ML, 4.00, 5.0, 0.30, [run(s["affiliation"], 14, GRAY2, F_M)])
    L.logos_row(sl, x_left=L.ML, y_bottom=6.90, h=L.LOGO_H_COVER, gap=L.LOGO_GAP_COVER)


def refs(sl, s):
    half = (len(s["refs"]) + 1) // 2
    for x, chunk in zip((L.COL[0], L.COL[3]), (s["refs"][:half], s["refs"][half:])):
        paras = [[run(r, 12, INK, F_M)] for r in chunk]
        text(sl, x, L.Y0, 5.85, 5.2, paras, line_spacing=1.15, space_after=6)


def build(out=PPTX):
    prs = fl.new_deck()
    for i, s in enumerate(SPEC["slides"]):
        sl = fl.blank(prs)
        if s["type"] == "cover":
            cover(sl, s)
        elif s["type"] == "references":
            L.header(sl, s["title"], None, i)
            refs(sl, s)
        else:
            L.header(sl, s["title"], s.get("subtitle"), i)
            if "table" in s:
                table_slide(sl, s)
            else:
                place_fig(sl, asset(s["assets"][0]))
            if s.get("takeaway"):
                L.takeaway(sl, s["takeaway"])
        L.notes(sl, s.get("notes"), None, None)
    RENDER.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"[pptx] {out.relative_to(ROOT)} {len(prs.slides)} 쪽")
    return prs


def audit():
    rep = L.pptx_audit(str(PPTX))
    ext = dict(L.deck_extra_audit(str(PPTX)))
    hits = 0
    for i, v in rep:
        e = ext[i]; msgs = []
        for k in ("fill_shapes", "off_font", "off_size", "big_num", "off_grid", "pic_scale_bad"):
            if v.get(k):
                msgs.append(f"{k} {v[k]}")
        if v.get("accent_fail"):
            msgs.append(f"accent {v['accent_outside_title']}")
        if e.get("boxed_text"):
            msgs.append(f"boxed_text {e['boxed_text']}")
        if e.get("dash"):
            msgs.append(f"dash {e['dash']}")
        if e.get("bullets", 0) > 4:
            msgs.append(f"bullets {e['bullets']}")
        if msgs:
            hits += 1; print(f"  [audit] p{i}: {'; '.join(msgs)}")
    print(f"[audit] 지적 쪽 {hits}/{len(rep)}")


def render():
    L.export_pdf(PPTX, PDF, SCRATCH / "lo_v5")
    PAGES.mkdir(parents=True, exist_ok=True)
    for old in PAGES.glob("page-*.png"):
        old.unlink()
    subprocess.run(["nice", "-n", "10", "pdftoppm", "-r", "60", "-png", str(PDF), str(PAGES / "page")], check=True)
    print(f"[pdf] {PDF.relative_to(ROOT)} 쪽 PNG {len(list(PAGES.glob('page-*.png')))}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--no-pdf", action="store_true"); a = ap.parse_args()
    build(); audit()
    if not a.no_pdf:
        render()
