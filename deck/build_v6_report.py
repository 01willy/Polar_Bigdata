# -*- coding: utf-8 -*-
"""build_v6_report.py: v6 연구 결과 보고 덱(국문, 16:9, 문법 V). 명세 deck/deck_spec_v6.json.
그림: outputs/figures/paper/v6/slide/(패널 크롭), deck/assets/v6/method/(방법 도식), deck/assets/paper_report/(표지 배경).
쪽 구성: 그림 왼쪽 열(폭 7.9 in) + 설명 줄 오른쪽 열(폭 3.8 in). wide=true 는 그림을 위에 12 in 폭으로 두고 설명을 아래에 둔다.
출력 deck/render/v6_report.pptx, .pdf, deck/render/v6_report_pages/page-XX.png. 실행: OMP_NUM_THREADS=2 nice -n 10 python3 deck/build_v6_report.py [--no-pdf]
"""
from __future__ import annotations
import argparse, json, os, re, subprocess, sys, tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx.util import Inches  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
import final_lib as fl  # noqa: E402
import paper_report_lib as L  # noqa: E402
from paper_report_lib import ACC, F_M, F_S, F_X, GRAY, GRAY2, INK, run, text  # noqa: E402

ROOT = L.ROOT; DECK = L.DECK; RENDER = DECK / "render"; PAGES = RENDER / "v6_report_pages"
PPTX = RENDER / "v6_report.pptx"; PDF = RENDER / "v6_report.pdf"
FIG = ROOT / "outputs" / "figures" / "paper" / "v6" / "slide"
METHOD = DECK / "assets" / "v6"
ASSET = DECK / "assets" / "paper_report"
SPEC = json.load(open(DECK / "deck_spec_v6.json", encoding="utf-8"))
NUM = json.load(open(ROOT / "docs" / "numbers_v6.json", encoding="utf-8")) if (ROOT / "docs" / "numbers_v6.json").exists() else {}


def subst(obj):
    """명세 문자열의 {nXxx} 를 수치 매크로(docs/numbers_v6.json, v6_numbers.py 생성)로 바꾼다. 수치를 손으로 적지 않기 위한 장치."""
    if isinstance(obj, str):
        def rep(m):
            k = m.group(1)[1:]
            return NUM.get(k, m.group(0))
        return re.sub(r"\{(n[A-Za-z]+)\}", rep, obj)
    if isinstance(obj, list):
        return [subst(x) for x in obj]
    if isinstance(obj, dict):
        return {k: subst(v) for k, v in obj.items()}
    return obj


SPEC = subst(SPEC)
SCRATCH = Path(os.environ.get("PR_TMP", tempfile.gettempdir()))
PART_PREFIX = {"배경": "1", "목적": "1", "자료": "2", "방법": "2", "결과": "3", "결론": "4"}


def asset(name):
    for p in (FIG / name, METHOD / name, ASSET / name):
        if p.exists():
            return p
    return None


def cover(sl, s):
    bg = ASSET / "cover_bg_v3.png"
    pic = sl.shapes.add_picture(str(bg), 0, 0, Inches(13.333), Inches(7.5)); pic.shadow.inherit = False
    L.break_runs(text(sl, L.ML, 0.95, 8.4, 1.30, [run(s["title"], 28, ACC, F_X)], anchor=MSO_ANCHOR.TOP, line_spacing=1.1))
    text(sl, L.ML, 2.55, 7.4, 0.95, [run(s["subtitle"], 18, GRAY2, F_M)], line_spacing=1.15)
    text(sl, L.ML, 3.80, 5.0, 0.36, [run(s["presenter"], 16, INK, F_S)])
    text(sl, L.ML, 4.20, 5.0, 0.30, [run(s["affiliation"], 14, GRAY2, F_M)])
    L.logos_row(sl, x_left=L.ML, y_bottom=6.90, h=L.LOGO_H_COVER, gap=L.LOGO_GAP_COVER)


def refs(sl, s):
    half = (len(s["refs"]) + 1) // 2
    for x, chunk in zip((L.COL[0], L.COL[3]), (s["refs"][:half], s["refs"][half:])):
        paras = [[run(r, 12, INK, F_M)] for r in chunk]
        text(sl, x, L.Y0, 5.85, 5.2, paras, line_spacing=1.15, space_after=6)


def leads(sl, x, y, w, items, pt=18, gap_pt=10):
    """설명 줄: 줄 머리 SemiBold + 본문 Medium. 본문은 폭 w 에 맞춰 줄바꿈(lead_wrap). 최대 4개."""
    assert len(items) <= 4
    paras = [[run(a + "  ", pt, INK, F_S), run(L.lead_wrap(a, b, w, pt=pt), pt, INK, F_M)] for a, b in items]
    return L.break_runs(text(sl, x, y, w, L.H, paras, line_spacing=1.12, space_after=gap_pt))


def fit_cap(p, max_w, max_h, cap=1.6):
    """배치 크기: 폭·높이 안에 맞추되 300 dpi 명목 크기의 cap 배를 넘지 않는다(크롭 글자 12 → 최대 19 pt)."""
    w, h = L.fit(p, max_w, max_h)
    wi, hi = L.img_size(p)
    nw, nh = wi / 300.0, hi / 300.0
    if w / nw > cap:
        w, h = nw * cap, nh * cap
    return w, h


def row_place(sl, paths, x, y, max_w, max_h, gap_in=0.25, cap=1.6):
    """여러 그림을 한 줄에 같은 높이로. 폭 합 ≤ max_w, 높이 ≤ max_h, 배율 ≤ cap. 반환 (총폭, 높이)."""
    sizes = [L.img_size(p) for p in paths]
    nw = [w / 300.0 for w, _ in sizes]; nh = [h / 300.0 for _, h in sizes]
    h = min(max_h, cap * min(nh))
    tot = sum(w * h / hh for w, hh in zip(nw, nh)) + gap_in * (len(paths) - 1)
    if tot > max_w:
        h *= (max_w - gap_in * (len(paths) - 1)) / (tot - gap_in * (len(paths) - 1))
    cx = x
    for p, w, hh in zip(paths, nw, nh):
        ww = w * h / hh; L.place(sl, p, cx, y, w=ww, h=h); cx += ww + gap_in
    return cx - gap_in - x, h


def gap(sl, x, y, w, h, name):
    L.gap_marker(sl, x, y, w, h, ["그림 대기", name, "Rescale core_v6 결과 도착 뒤 생성"])


def is_method(p):
    return p is not None and "assets/v6/method" in str(p)


def fig_slide(sl, s):
    p = asset(s["assets"][0]); wide = s.get("wide", False); items = [tuple(i) for i in s.get("leads", [])]
    if is_method(p):
        # 방법 도식은 12 × 5.2 in 배치 크기 그대로(배율 1.0). 설명 줄은 발표자 노트로
        L.place(sl, p, L.ML, L.Y0, w=12.0, h=5.2)
        s["notes"] = (s.get("notes") or "") + " " + " ".join(f"{a}: {b}." for a, b in items)
        return
    paths = [asset(a) for a in s["assets"]]
    if s.get("layout") == "left-tall" and all(paths):
        # 첫 그림은 왼쪽에 높이 H 로, 둘째 그림은 오른쪽 위, 설명 줄은 그 아래
        w0, h0 = fit_cap(paths[0], 6.0, L.H); L.place(sl, paths[0], L.ML, L.Y0, w=w0, h=h0)
        xr = L.ML + w0 + 0.30; wr = L.ML + 12.0 - xr
        w1, h1 = fit_cap(paths[1], wr, 2.9); L.place(sl, paths[1], xr, L.Y0, w=w1, h=h1)
        if items:
            leads(sl, xr, L.Y0 + h1 + 0.2, wr, items, pt=16, gap_pt=6)
        return
    if wide:
        if not all(paths):
            gap(sl, L.ML, L.Y0, 12.0, 3.3, s["assets"][0])
            yt = L.Y0 + 3.45
        else:
            w, h = row_place(sl, paths, L.ML, L.Y0, 12.0, 3.55); yt = L.Y0 + h + 0.18
        if items:
            # 아래 설명: 두 열(각 5.85 in)
            n = len(items); half = (n + 1) // 2
            leads(sl, L.COL[0], yt, 5.85, items[:half], pt=16, gap_pt=6)
            if items[half:]:
                leads(sl, L.COL[3], yt, 5.85, items[half:], pt=16, gap_pt=6)
    else:
        if not all(paths):
            gap(sl, L.ML, L.Y0, 7.9, L.H, s["assets"][0])
        else:
            w, h = row_place(sl, paths, L.ML, L.Y0, 7.9, L.H)
        if items:
            leads(sl, L.COL[4], L.Y0 + 0.05, 3.80, items, pt=16, gap_pt=8)


def text_slide(sl, s):
    items = [tuple(i) for i in s.get("leads", [])]
    leads(sl, L.ML, L.Y0 + 0.15, 12.0, items, pt=18, gap_pt=14)


def table_slide(sl, s):
    t = s["table"]
    geo = dict(x=L.ML, y=L.Y0 + 0.15, col_w=t["col_w"], row_h=t["row_h"])
    td = L.table_def(t["columns"], t["rows"], geo)
    L.hairline_table(sl, td, y=geo["y"], row_h=t["row_h"])


def build(out=PPTX):
    prs = fl.new_deck()
    for i, s in enumerate(SPEC["slides"]):
        sl = fl.blank(prs)
        if s["type"] == "cover":
            cover(sl, s)
        elif s["type"] == "references":
            L.header(sl, s["title"], None, i); refs(sl, s)
        else:
            L.header(sl, s["title"], s.get("subtitle"), i)
            {"fig": fig_slide, "text": text_slide, "table": table_slide}[s["type"]](sl, s)
            if s.get("takeaway"):
                L.takeaway(sl, s["takeaway"])
        L.notes(sl, s.get("notes"), None, None)
    RENDER.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"[pptx] {out.relative_to(ROOT)} {len(prs.slides)} 쪽")
    return prs


def audit():
    rep = L.pptx_audit(str(PPTX)); ext = dict(L.deck_extra_audit(str(PPTX))); hits = 0
    for i, v in rep:
        e = ext[i]; msgs = []
        for k in ("fill_shapes", "off_font", "off_size", "big_num", "off_grid"):
            if v.get(k):
                msgs.append(f"{k} {v[k]}")
        bad = [x for x in (v.get("pic_scale_bad") or []) if not (0.88 <= x <= 1.62)]
        if bad:
            msgs.append(f"pic_scale {bad}")
        if v.get("accent_fail"):
            msgs.append(f"accent {v['accent_outside_title']}")
        for k in ("boxed_text", "dash"):
            if e.get(k):
                msgs.append(f"{k} {e[k]}")
        if e.get("bullets", 0) > 4:
            msgs.append(f"bullets {e['bullets']}")
        if msgs:
            hits += 1; print(f"  [audit] p{i}: {'; '.join(msgs)}")
    print(f"[audit] 지적 쪽 {hits}/{len(rep)}")


def render():
    L.export_pdf(PPTX, PDF, SCRATCH / "lo_v6")
    PAGES.mkdir(parents=True, exist_ok=True)
    for old in PAGES.glob("page-*.png"):
        old.unlink()
    subprocess.run(["nice", "-n", "10", "pdftoppm", "-r", "60", "-png", str(PDF), str(PAGES / "page")], check=True)
    print(f"[pdf] {PDF.relative_to(ROOT)} 쪽 PNG {len(list(PAGES.glob('page-*.png')))}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--no-pdf", action="store_true"); a = ap.parse_args()
    build(); audit()
    if not a.no_pdf:
        render()
