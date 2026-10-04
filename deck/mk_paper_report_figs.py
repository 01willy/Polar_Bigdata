#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""보고 덱(paper_report)의 덱 전용 근거 차트와 헤어라인 표를 만든다.

대상: deck/deck_spec_paper_report.json 에서 v3 논문 그림(또는 그 슬라이드판)이 아닌 근거.
  차트 14개: S03, S06, S08, S09(개념 좌표도), S20, S24(오른쪽 포레스트), S25, S26(워크플로 경로), S27(경로 축소판, 열 1),
            S29, S32, AP3, AP5, AP8
  표 15개 : S02, S04, S06, S10, S29, S30, S31, S33, S34, AP2, AP4, AP6, AP7, AP9, AP10
  수치 대조: values_check.txt(스펙 numbers ↔ 원천 CSV 재계산)

형식 선택(지침 5.6, 3.2)
  - 차트: 경로 B. Pretendard 로 배치 크기 그대로 그려 300 dpi PNG 로 저장(PPTX 배율 1.00).
    한글이 든 차트라 EMF 벡터 경로(A)는 쓰지 않는다(한글이 Noto Sans CJK 로 대체됨, 지침 R-16).
  - 표: python-pptx 원어 헤어라인 표(텍스트 상자 + 규칙선, final_lib.mini_table 과 같은 구성).
    편집 가능해야 하므로 그림으로 만들지 않는다. 표 정의는 tables/<쪽>_table.json 이고,
    build_hairline_table() 를 덱 빌더가 그대로 불러 쓴다.

입력(읽기 전용): paper/registry/experiments.csv, paper/claims/*/tables/*.csv,
  data/processed/fidelity_base_v{2,3,4}_meta.json, data/processed/paper_figs/fig1_meta.csv,
  docs/QA_FINAL_REVIEW_2026-10-02.md Q8·Q15 표와 docs/QA_FOLLOWUP_2026-10-04.md 3.4(S06 값),
  docs/research/2026-10-04/scirep_format_and_drafts.md 3.2, paper/manuscript/MANUSCRIPT_SPEC.md 1.1(S29 값),
  docs/EXPERIMENT_PLAN_FINAL_BATCH_2026-10-04.md 0.1, 6(S32 값)
출력: deck/assets/paper_report/
  <쪽>_<이름>.png, source_data/<쪽>_<이름>.csv, tables/<쪽>_table.json,
  preview_evidence.pptx(근거 쪽 미리보기: 머리, 근거, 결론 줄), preview/(PDF, 쪽 PNG), MANIFEST.json

실행: python deck/mk_paper_report_figs.py [--only S03,S08] [--no-preview]
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager as fm  # noqa: E402
from matplotlib.collections import Collection  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch, Rectangle  # noqa: E402
from matplotlib.text import Text  # noqa: E402
from matplotlib.ticker import FixedLocator, NullLocator  # noqa: E402
from PIL import ImageFont  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "deck"
SPEC = DECK / "deck_spec_paper_report.json"
OUT = DECK / "assets" / "paper_report"
SRC_DIR = OUT / "source_data"
TAB_DIR = OUT / "tables"
PREV_DIR = OUT / "preview"
CLAIMS = ROOT / "paper" / "claims"
FONT_DIR = Path.home() / ".fonts"

# ---------------------------------------------------------------- 양식 토큰(지침 2.4, 5.3)
INK = "#111111"
AUX = "#4d4d4d"          # 0선, 기준선, 방향 표지, 보조 표지
BAND = "#ededed"         # ±0.5 cm 동등 띠
GRAY_ST = "#b0b0b0"      # 대체·폐기 상태 원, 오차 하한 띠
METHOD = {
    "source": "#4d4d4d",   # 원천 계수 Stefan(Stefan 최소제곱 포함)
    "recal": "#2b5c8f",    # 재보정 Stefan
    "resid": "#9a7bc9",    # 물리 잔차 결합
    "aug": "#568f72",      # 물리 유사라벨 증강
    "direct": "#6b7280",   # 직접 ML
    "physin": "#ad921a",   # 물리 입력 ML
    "cci": "#84480c",      # Stefan·CCI 평균 앵커
}
AX_PT, TK_PT = 18, 16    # 축 라벨 18 pt, 눈금·직접 라벨 16 pt
LW_MAIN, LW_AUX, LW_AXIS, LW_ZERO = 3.0, 2.0, 1.5, 2.0
MS = 8                   # 측정점 8 pt
MS_SMALL = 6             # 지역별 개별 점(alpha 0.5)
DPI = 300
W_MED, W_SEMI = 500, 600

# ---------------------------------------------------------------- 공통
CODE_RE = re.compile(r"\b(L\d{1,2}|AB\d{1,2}|LG[A-Z]?(-[A-Z]\d)?|WF\d{1,2}|SC\d\w*|H\d{1,2}\w?|X[A-J]|P4|PE[12]"
                     r"|L\+C|P[0-2]\*?|R[0-2]|D[01]|v[1-5])\b|\((x|i)\)")
DASH_RE = re.compile(r"(?<![0-9])\s?[—–]\s?(?![0-9])|—")


def load_spec():
    with open(SPEC, encoding="utf-8") as f:
        d = json.load(f)
    return {s["id"]: s for s in d["slides"] + d["appendix"]}


def rel(p):
    return str(Path(p).resolve().relative_to(ROOT))


def setup_mpl():
    for w in ("Medium", "SemiBold"):
        fm.fontManager.addfont(str(FONT_DIR / f"Pretendard-{w}.otf"))
    plt.rcParams.update({
        "font.family": "Pretendard", "font.weight": W_MED, "font.size": TK_PT,
        "text.color": INK, "axes.labelcolor": INK, "axes.edgecolor": INK,
        "axes.labelsize": AX_PT, "axes.labelweight": W_MED, "axes.linewidth": LW_AXIS,
        "axes.spines.top": False, "axes.spines.right": False, "axes.grid": False,
        "axes.unicode_minus": True, "axes.facecolor": "white", "figure.facecolor": "white",
        "xtick.labelsize": TK_PT, "ytick.labelsize": TK_PT, "xtick.color": INK, "ytick.color": INK,
        "xtick.major.width": LW_AXIS, "ytick.major.width": LW_AXIS,
        "xtick.major.size": 5, "ytick.major.size": 5, "xtick.major.pad": 6, "ytick.major.pad": 8,
        "xtick.direction": "out", "ytick.direction": "out",
        "xtick.minor.visible": False, "ytick.minor.visible": False,
        "legend.frameon": False, "savefig.dpi": DPI, "figure.dpi": DPI,
        "pdf.fonttype": 42, "svg.fonttype": "none", "lines.solid_capstyle": "butt",
    })


def new_fig(w, h):
    fig = plt.figure(figsize=(w, h), dpi=DPI)
    fig.canvas.draw()
    return fig


def add_ax(fig, l, b, w, h):
    W, H = fig.get_size_inches()
    return fig.add_axes([l / W, b / H, w / W, h / H])


def text_size_in(fig, s, size=TK_PT, weight=W_MED):
    t = fig.text(0, 0, s, fontsize=size, fontweight=weight)
    bb = t.get_window_extent(fig.canvas.get_renderer())
    t.remove()
    return bb.width / fig.dpi, bb.height / fig.dpi


def max_text_w(fig, labels, size=TK_PT, weight=W_MED):
    return max(text_size_in(fig, s, size, weight)[0] for s in labels)


def ftext(fig, x, y, s, size=TK_PT, weight=W_MED, ha="left", va="center", color=INK, gid=None):
    """그림 좌표(in) 글자."""
    return fig.text(x, y, s, transform=fig.dpi_scale_trans, fontsize=size, fontweight=weight,
                    ha=ha, va=va, color=color, gid=gid)


def key_line(fig, x0, y, items, gap=0.34):
    """그림 안 열쇠 한 줄(지침 2.7, 4.4). items: (kind, label, kw). 끝 x(in)를 돌려준다."""
    x = x0
    tr = fig.dpi_scale_trans
    for kind, label, kw in items:
        if kind in ("dot", "odot"):
            fc = kw.get("fc", "black") if kind == "dot" else "white"
            ec = kw.get("ec", kw.get("fc", "black"))
            fig.add_artist(Line2D([x + 0.07], [y], transform=tr, marker="o", ms=kw.get("ms", MS), mfc=fc,
                                  mec=ec, mew=1.5, ls="none", alpha=kw.get("alpha", 1.0)))
            gw = 0.14
        elif kind == "line":
            fig.add_artist(Line2D([x, x + 0.46], [y, y], transform=tr, color=kw.get("c", "black"),
                                  lw=kw.get("lw", LW_AUX), ls=kw.get("ls", "-"), solid_capstyle="butt",
                                  dash_capstyle="butt"))
            if kw.get("marker"):
                fig.add_artist(Line2D([x + 0.23], [y], transform=tr, marker="o", ms=MS, mfc=kw.get("c"),
                                      mec=kw.get("c"), ls="none"))
            gw = 0.46
        elif kind == "band":
            fig.add_artist(Rectangle((x, y - 0.09), 0.46, 0.18, transform=tr, fc=kw.get("fc", BAND), ec="none",
                                     gid="key_swatch"))
            gw = 0.46
        else:
            raise ValueError(kind)
        t = ftext(fig, x + gw + 0.09, y, label, gid="key")
        w = t.get_window_extent(fig.canvas.get_renderer()).width / fig.dpi
        x = x + gw + 0.09 + w + gap
    return x - gap


def cat_axis(ax, labels, positions, weight=W_MED):
    ax.set_yticks(positions)
    ax.set_yticklabels(labels, fontsize=TK_PT, fontweight=weight)
    ax.tick_params(axis="y", length=0)
    ax.spines["left"].set_visible(False)
    for t in ax.get_yticklabels():
        t.set_gid("category")


def in_per_data_y(fig, ax):
    bb = ax.get_window_extent(fig.canvas.get_renderer())
    y0, y1 = ax.get_ylim()
    return (bb.height / fig.dpi) / abs(y1 - y0)


def forest_row(ax, y, d, lo, hi, dbe, lob, hib, color, off):
    """점, 셀 가중 CI(굵게), 블록 등가중 CI(가늘게, 바로 아래)(지침 2.7)."""
    ax.plot([lo, hi], [y, y], color=color, lw=LW_MAIN, solid_capstyle="butt", zorder=3)
    if np.isfinite(lob) and np.isfinite(hib):
        ax.plot([lob, hib], [y + off, y + off], color=color, lw=LW_AUX, solid_capstyle="butt", zorder=3)
    ax.plot([d], [y], marker="o", ms=MS, mfc=color, mec=color, ls="none", zorder=4)


def ref_band_zero(ax):
    ax.axvspan(-0.5, 0.5, fc=BAND, ec="none", zorder=0)
    ax.axvline(0, color=AUX, lw=LW_ZERO, zorder=1)


def forest_key(fig, x0, y, color="black"):
    return key_line(fig, x0, y, [
        ("line", "셀 가중 95% CI", dict(c=color, lw=LW_MAIN)),
        ("line", "블록 등가중 95% CI", dict(c=color, lw=LW_AUX)),
        ("band", "±0.5 cm 동등 범위", {}),
    ])


# ---------------------------------------------------------------- 그림 점검(지침 부록 A.1 의 슬라이드판)
def audit_chart(fig):
    fig.canvas.draw()
    r = fig.canvas.get_renderer()
    W, H = fig.bbox.width, fig.bbox.height
    out = dict(sizes=set(), fonts=set(), thin=[], titles=[], boxed=0, codes=[], comma4=[], dash=[],
               clipped=[], overlaps=[], loose_numbers=[], n_text=0, chars=0)
    texts = [t for t in fig.findobj(Text) if t.get_visible() and t.get_text().strip()]
    ticks = set()
    for a in fig.axes:
        for axis in (a.xaxis, a.yaxis):
            ticks.update(axis.get_ticklabels())
            ticks.add(axis.label)
    boxes = []
    for t in texts:
        s = t.get_text().strip()
        out["n_text"] += 1
        out["chars"] += len(s.replace(" ", ""))
        out["sizes"].add(round(t.get_fontsize(), 2))
        out["fonts"].add(Path(fm.findfont(t.get_fontproperties())).name)
        if t.get_bbox_patch() is not None:
            out["boxed"] += 1
        if CODE_RE.search(s):
            out["codes"].append(s)
        if re.search(r"(?<![\d,.])\d,\d{3}(?![\d,])", s):
            out["comma4"].append(s)
        if DASH_RE.search(s):
            out["dash"].append(s)
        bb = t.get_window_extent(r)
        if bb.x0 < -1 or bb.y0 < -1 or bb.x1 > W + 1 or bb.y1 > H + 1:
            out["clipped"].append(s)
        if t not in ticks and t.get_gid() not in ("key", "category", "label_num") and re.search(r"\d", s):
            out["loose_numbers"].append(s)
        boxes.append((s, bb))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i][1], boxes[j][1]
            ix = min(a.x1, b.x1) - max(a.x0, b.x0)
            iy = min(a.y1, b.y1) - max(a.y0, b.y0)
            if ix > 1 and iy > 1:
                out["overlaps"].append((boxes[i][0], boxes[j][0]))
    for ln in fig.findobj(Line2D):
        if ln.get_visible() and ln.get_linestyle() not in ("None", "none", "") and ln.get_linewidth() < 1.0 - 1e-6:
            out["thin"].append(ln.get_label())
    for c in fig.findobj(Collection):
        lws = [w for w in (c.get_linewidths() if hasattr(c, "get_linewidths") else []) if w > 0]
        if lws and min(lws) < 1.0 - 1e-6:
            out["thin"].append(type(c).__name__)
    out["tick_gap"] = []
    for a in fig.axes:
        if a.get_title().strip():
            out["titles"].append(a.get_title())
        vis = [sp for sp in a.spines.values() if sp.get_visible()]
        if vis and min(sp.get_linewidth() for sp in vis) < 1.0 - 1e-6:
            out["thin"].append("spine")
        # 눈금 라벨 사이 간격 0.10 in 이상(원형 1.3). 같은 축의 이웃 라벨끼리 잰다
        for axis, horiz in ((a.xaxis, True), (a.yaxis, False)):
            bbs = sorted([t.get_window_extent(r) for t in axis.get_ticklabels() if t.get_visible() and t.get_text().strip()],
                         key=lambda b: b.x0 if horiz else b.y0)
            for p, q in zip(bbs[:-1], bbs[1:]):
                gap = (q.x0 - p.x1) if horiz else (q.y0 - p.y1)
                if gap < 0.10 * fig.dpi - 0.5:
                    out["tick_gap"].append(round(gap / fig.dpi, 3))
    rects = [p for p in fig.findobj(Patch) if isinstance(p, Rectangle) and p.get_visible()
             and p not in [a.patch for a in fig.axes] and p is not fig.patch and p.get_gid() != "key_swatch"]
    for p in rects:
        pb = p.get_window_extent(r)
        for s, tb in boxes:
            if pb.width > 2 and pb.height > 2 and pb.x0 <= tb.x0 and tb.x1 <= pb.x1 and pb.y0 <= tb.y0 and tb.y1 <= pb.y1:
                out["boxed"] += 1
    fails = []
    if out["sizes"] - {16.0, 18.0}:
        fails.append("size")
    if out["fonts"] - {"Pretendard-Medium.otf", "Pretendard-SemiBold.otf"}:
        fails.append("font")
    for k in ("thin", "titles", "codes", "comma4", "dash", "clipped", "overlaps", "loose_numbers", "tick_gap"):
        if out[k]:
            fails.append(k)
    if out["boxed"]:
        fails.append("boxed")
    out["sizes"] = sorted(out["sizes"])
    out["fonts"] = sorted(out["fonts"])
    out["fails"] = fails
    return out


MANIFEST = {"charts": {}, "tables": {}}


def finish(fig, sid, name, df, meta):
    """점검 뒤 배치 크기 그대로 300 dpi PNG 와 그림 값 CSV 를 저장한다(bbox_inches 미사용)."""
    a = audit_chart(fig)
    png = OUT / f"{sid}_{name}.png"
    fig.savefig(png, dpi=DPI, facecolor="white")
    w, h = fig.get_size_inches()
    plt.close(fig)
    csv = SRC_DIR / f"{sid}_{name}.csv"
    df.to_csv(csv, index=False, encoding="utf-8")
    MANIFEST["charts"][sid] = dict(
        file=rel(png), source_data=rel(csv), format="PNG(경로 B, Pretendard, 300 dpi, 배치 크기 그대로)",
        size_in=[round(w, 3), round(h, 3)], size_px=[int(round(w * DPI)), int(round(h * DPI))], **meta,
        audit={k: a[k] for k in ("fails", "sizes", "fonts", "n_text", "chars", "overlaps", "clipped",
                                 "loose_numbers", "codes", "dash", "thin", "boxed", "tick_gap")})
    status = "OK" if not a["fails"] else f"FAIL {a['fails']}"
    print(f"[chart] {sid:4s} {png.name:34s} {w:.2f}x{h:.2f} in  text {a['n_text']:3d}  {status}")
    if a["fails"]:
        for k in a["fails"]:
            print("        ", k, a.get(k))
    return a


# ================================================================ S03 연구 흐름 타임라인
def chart_s03(spec):
    sp = spec["S03"]
    reg_path = ROOT / "paper/registry/experiments.csv"
    reg = pd.read_csv(reg_path)
    T0 = datetime(2026, 10, 4, 14, 22)   # 추가 실험 사전 등록 커밋 2678100(2026-10-04 14:22:39 +0900)
    rows = [("H0", "대회 탐색"), ("H1", "사전 등록 실험"), ("H2-5", "통합 실험"), ("T", "전이 시험"),
            ("W", "지역 내·워크플로 시험"), ("X", "추가 실험")]
    assert [r[1] for r in rows] == sp["visible_text"]["row_labels"]

    def day(s):
        return datetime.strptime(s, "%Y-%m-%d") + timedelta(hours=12)

    items = []
    for _, r in reg.iterrows():
        s, st = str(r["dates"]), str(r["status"])
        if "등록 예정" in s:
            a = b = T0
        else:
            m = re.search(r"≤(\d{4}-\d{2}-\d{2})", s)
            if m:                                   # 'YYYY-MM(≤…)' 는 끝점만
                a = b = day(m.group(1))
            else:
                ds = re.findall(r"\d{4}-\d{2}-\d{2}", s)
                a, b = day(ds[0]), day(ds[-1])     # 범위·나열은 시작과 끝
        cls = "done" if st.startswith("done") else ("planned" if st.startswith("planned") else "gray")
        ph = r["phase"]
        key = "H2-5" if ph in ("H2", "H3", "H4", "H5") else ph
        items.append(dict(row=key, new_name=r["new_name"], old_ids=r["old_ids"], dates=s, status=st,
                          cls=cls, start=a, end=b))
    df = pd.DataFrame(items)

    # 같은 행에서 날짜가 겹치는 묶음은 줄(lane)을 나눈다(같은 날 묶음이 서로 가리지 않게)
    gap_days = 1.7
    df["lane"] = -1
    nl = {}
    for key, _ in rows:
        sub = df[df.row == key].sort_values(["start", "end"])
        ends = []
        for i, it in sub.iterrows():
            for li, e in enumerate(ends):
                if (it.start - e).total_seconds() / 86400 > gap_days:
                    ends[li] = it.end
                    df.at[i, "lane"] = li
                    break
            else:
                ends.append(it.end)
                df.at[i, "lane"] = len(ends) - 1
        nl[key] = len(ends)

    W, H = 12.0, 5.2
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, [r[1] for r in rows])
    left, right, bottom, top_band = lab_w + 0.30, 0.25, 0.80, 0.10
    plot_h = H - bottom - top_band
    row_gap = 0.8                                  # 행 사이 간격(줄 단위)
    min_rows = 4                                   # 줄이 적은 행도 최소 4줄 높이(행 이름 16 pt 가 서로 닿지 않게)
    ax = add_ax(fig, left, bottom, W - left - right, plot_h)
    ycen, ypos, y = {}, {}, 0.0
    for key, lab in rows:
        h_rows = max(nl[key], min_rows)
        ypos[key] = y + (h_rows - nl[key]) / 2       # 줄을 행 높이 안에서 가운데 정렬
        ycen[key] = y + (h_rows - 1) / 2
        y += h_rows + row_gap
    last = y - row_gap - 1
    ax.set_ylim(last + 0.7, -0.7)
    pitch = plot_h / (last + 1.4)
    assert min(abs(ycen[a] - ycen[b]) for a, b in zip([k for k, _ in rows][:-1], [k for k, _ in rows][1:])) * pitch >= 0.36
    # 같은 날 묶음이 최대 12개(W 행 10월 1–2일)라 8 pt 원은 겹친다. 원 지름을 줄 간격의 92 % 로 줄인다
    ms = min(MS, round(pitch * 72 * 0.92, 1))
    for i, it in df.iterrows():
        yy = ypos[it.row] + it.lane
        if it.cls == "gray":
            lc, fc, ec = GRAY_ST, GRAY_ST, GRAY_ST
        elif it.cls == "planned":
            lc, fc, ec = "black", "white", "black"
        else:
            lc, fc, ec = "black", "black", "black"
        if it.end > it.start:
            ax.plot([it.start, it.end], [yy, yy], color=lc, lw=LW_AUX, solid_capstyle="butt", zorder=2)
        ax.plot([it.end], [yy], marker="o", ms=ms, mfc=fc, mec=ec, mew=1.0, ls="none", zorder=4)
    ax.axvline(T0, color=AUX, lw=LW_AXIS, zorder=1)
    cat_axis(ax, [r[1] for r in rows], [ycen[k] for k, _ in rows])
    ax.set_xlim(datetime(2026, 6, 26), datetime(2026, 10, 12))
    months = [datetime(2026, m, 1) for m in (7, 8, 9, 10)]
    ax.xaxis.set_major_locator(FixedLocator(mdates.date2num(months)))
    ax.set_xticklabels([f"{m.month}월" for m in months])
    ax.set_xlabel("날짜 (2026년)")
    # 열쇠: 자료가 없는 8월 구역 안 세로 3줄(지침 2.8 '자료가 없는 패널 안 빈 곳'). 사건 표지: 등록선 왼쪽 위
    fig.canvas.draw()
    xk = ax.transData.transform((mdates.date2num(datetime(2026, 8, 4)), 0))[0] / fig.dpi
    yk0 = ax.transData.transform((0, ypos["T"]))[1] / fig.dpi
    for k, item in enumerate([("dot", "판정 기록", dict(fc="black", ms=ms)), ("odot", "결과 대기", dict(ms=ms)),
                              ("dot", "대체·폐기·미실행", dict(fc=GRAY_ST, ec=GRAY_ST, ms=ms))]):
        key_line(fig, xk, yk0 - 0.36 * k, [item])
    xl = ax.transData.transform((mdates.date2num(T0), 0))[0] / fig.dpi
    ftext(fig, xl - 0.10, H - top_band - 0.16, sp["visible_text"]["event_label"], ha="right", color=AUX)

    out = df.assign(start=df.start.dt.strftime("%Y-%m-%d %H:%M"), end=df.end.dt.strftime("%Y-%m-%d %H:%M"))
    out = out.assign(row_label=out.row.map(dict(rows)), source=rel(reg_path))
    meta = dict(slide="S03", page=3, archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(reg_path), "git log 2678100(2026-10-04 14:22:39 +0900)"],
                notes=[f"줄 간격 {pitch:.3f} in, 원 {ms} pt(8 pt 원 지름 0.111 in 는 같은 날 묶음 12개를 겹침 없이 쌓지 못해 줄임)",
                       f"행별 줄 수 {nl}",
                       "범위·나열 날짜는 시작과 끝, '2026-08(≤2026-08-31)'은 끝점만, 추가 실험은 등록 시각",
                       "대체·폐기·미실행(superseded, discarded, not run)은 회색 선과 회색 원"])
    finish(fig, "S03", "timeline", out, meta)
    return pitch


# ================================================================ S06 입력별 공간 규모
def chart_s06(spec):
    sp = spec["S06"]
    labels = sp["visible_text"]["chart_row_labels"]
    q8 = "docs/QA_FINAL_REVIEW_2026-10-02.md Q8 표"
    fu = "docs/QA_FOLLOWUP_2026-10-04.md 3.4 표(지형·SAR 정정)"
    vals = [  # (원자료 해상도 m, 모형 입력 단위 m, 근거 문구, 원천)
        (11, 1000, "점 관측(수 m)을 좌표 소수점 4자리(약 11 m)로 묶어 평균한 뒤 1 km 셀", q8),
        (30, 30, "Copernicus DEM 30 m 의 점 화소 값", fu),
        (30, 990, "30 m 화소 33 × 33 창(약 1 km)", fu),
        (11100, 11100, "ERA5-Land 0.1°(남북 약 11 km), 가장 가까운 격자값", q8),
        (250, 5000, "SoilGrids 250 m 제품을 약 5 km 로 요청해 추출", q8),
        (1000, 1000, "ESA CCI 약 1 km, 가장 가까운 화소", q8),
        (30, 150, "30 m, PolSAR 150 m 창(InSAR 창 [미확인])", fu),
    ]
    W, H = 7.9, 4.9
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, labels)
    left, right, bottom, top_band = lab_w + 0.30, 0.30, 0.86, 0.50
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    n = len(labels)
    for i, (raw, unit, _, _) in enumerate(vals):
        if unit != raw:
            ax.plot([raw, unit], [i, i], color="black", lw=LW_AUX, zorder=2)
        ax.plot([raw], [i], marker="o", ms=MS, mfc="white", mec="black", mew=1.5, ls="none", zorder=3)
        ax.plot([unit], [i], marker="o", ms=MS, mfc="black", mec="black", ls="none", zorder=4)
    ax.axvline(1000, color=AUX, lw=LW_AXIS, zorder=1)
    ax.set_xscale("log")
    ax.set_xlim(7, 25000)
    ax.xaxis.set_major_locator(FixedLocator([10, 100, 1000, 10000]))
    ax.xaxis.set_minor_locator(NullLocator())
    ax.set_xticklabels(["10", "100", "1000", "10,000"])
    ax.set_ylim(n - 0.5, -0.5)
    cat_axis(ax, labels, range(n))
    ax.set_xlabel("공간 규모 (m)")
    yk = H - 0.24
    key_line(fig, 0.05, yk, [("odot", "원자료 해상도", {}), ("dot", "모형 입력 단위", dict(fc="black"))])
    fig.canvas.draw()
    xl = ax.transData.transform((1000, 0))[0] / fig.dpi
    ftext(fig, xl, H - top_band + 0.08, "출력 셀", ha="center", va="bottom", color=AUX)
    df = pd.DataFrame([dict(row=lab, raw_m=v[0], model_unit_m=v[1], basis=v[2], source=v[3])
                       for lab, v in zip(labels, vals)])
    meta = dict(slide="S06", page=6, archetype="A3", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[q8, fu, "docs/EXPERIMENT_PLAN_FINAL_BATCH_2026-10-04.md 0.4"],
                notes=["값은 문서 표에서 옮겼다(원천 CSV 없음). 라벨 원자료는 '점(수 m)'이라 문서의 묶음 단위 11 m 로 찍었다",
                       "ERA5-Land 0.1° 는 남북 11,100 m 로 찍었다. SAR 은 PolSAR 150 m 창(InSAR 창 미확인)",
                       "스펙 높이 5.1 in 를 4.9 in 로 줄여 보조 표와 높이를 맞췄다(A3 규칙, tables/S06_table.json)"])
    finish(fig, "S06", "spatial_scale", df, meta)


# ================================================================ S08 검증 사다리
def chart_s08(spec):
    sp = spec["S08"]
    labels = sp["visible_text"]["row_labels"]
    schemes = ["V-R", "V-P", "V-S", "V-B", "V-C0", "V-C100", "V-C500", "V-G"]
    tpath = CLAIMS / "C7_evaluation_design/tables/lgv_tests.csv"
    mpath = CLAIMS / "C7_evaluation_design/tables/lgv_metrics.csv"
    t = pd.read_csv(tpath)
    m = pd.read_csv(mpath)
    rows = []
    for s, lab in zip(schemes, labels):
        d0 = t[(t.test_id == "L20") & (t["item"] == f"D0[catboost_lo] RMSE@{s}") & (t.scope == "MEAN5")]
        ps = m[(m.scheme == s) & (m.method == "PS") & (m.scope == "MEAN5")]
        assert len(d0) == 1 and len(ps) == 1, s
        rows.append(dict(rung=lab, scheme_code=s, direct_ml_rmse=float(d0.rmse_A.iloc[0]),
                         stefan_ls_rmse=float(ps.rmse.iloc[0])))
    df = pd.DataFrame(rows)
    W, H = 12.0, 5.1
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, labels)
    names = sp["visible_text"]["direct_labels"]   # ['직접 ML', 'Stefan 최소제곱']
    left, right, bottom, top = lab_w + 0.30, 1.45, 0.86, 0.20
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top)
    n = len(df)
    yy = np.arange(n)
    ax.plot(df.direct_ml_rmse, yy, color=METHOD["direct"], lw=LW_AUX, ls=(0, (2, 1.5)), zorder=2)
    ax.plot(df.direct_ml_rmse, yy, marker="o", ms=MS, mfc=METHOD["direct"], mec=METHOD["direct"], ls="none", zorder=3)
    ax.plot(df.stefan_ls_rmse, yy, color=METHOD["source"], lw=LW_AUX, zorder=2)
    ax.plot(df.stefan_ls_rmse, yy, marker="o", ms=MS, mfc=METHOD["source"], mec=METHOD["source"], ls="none", zorder=3)
    ax.set_xlim(20, 35)
    ax.set_xticks([20, 25, 30, 35])
    ax.set_ylim(n - 0.5, -0.5)
    cat_axis(ax, labels, yy)
    ax.set_xlabel("RMSE (cm)")
    # 직접 라벨: 마지막 행 한 번(직접 ML 은 점 오른쪽, Stefan 은 점 왼쪽)
    ax.annotate(names[0], (df.direct_ml_rmse.iloc[-1], n - 1), xytext=(10, 0), textcoords="offset points",
                ha="left", va="center", fontsize=TK_PT, fontweight=W_SEMI, color=INK, annotation_clip=False)
    ax.annotate(names[1], (df.stefan_ls_rmse.iloc[-1], n - 1), xytext=(-10, 0), textcoords="offset points",
                ha="right", va="center", fontsize=TK_PT, fontweight=W_SEMI, color=INK)
    df["source"] = f"{rel(tpath)} (L20, item 'D0[catboost_lo] RMSE@*', MEAN5, rmse_A); {rel(mpath)} (method PS, MEAN5, rmse)"
    meta = dict(slide="S08", page=8, archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(tpath), rel(mpath)],
                notes=["다섯 지역 등가중 평균(MEAN5). 직접 ML = CatBoost(catboost_lo)",
                       "Stefan 최소제곱은 원천 계수 Stefan 의 색과 선(#4d4d4d 실선)"])
    finish(fig, "S08", "validation_ladder", df, meta)


# ================================================================ S20 대회 결과와 이번 시험
def chart_s20(spec):
    sp = spec["S20"]
    s4p = CLAIMS / "C4_sufficient_labels/tables/s4_residual_results.csv"
    cvp = CLAIMS / "C4_sufficient_labels/tables/wf2b_curve.csv"
    flp = CLAIMS / "C4_sufficient_labels/tables/lgx_floor.csv"
    s4 = pd.read_csv(s4p)
    base = (s4.part == "D_indomain") & (s4.cv == "spatial_block_AK") & (s4.model == "ridge") & \
           (s4.featset == "shared25") & (s4.seed == 0)
    st = float(s4[base & (s4.lam == 0.0)].rmse_cm.iloc[0])
    rs = float(s4[base & (s4.lam == 0.75)].rmse_cm.iloc[0])
    cv = pd.read_csv(cvp)
    c = cv[(cv.exp == "wf6") & (cv.target == "Alaska") & (cv.n == 1000)]
    p1 = float(c[c.method == "P1"].rmse.iloc[0])
    r1 = float(c[(c.method == "R1") & (c.lam == -1.0)].rmse.iloc[0])
    fl = pd.read_csv(flp)
    floor = float(fl[(fl.region == "Alaska") & (fl.scope == "eval")].floor_rmse_cm.iloc[0])
    rows = sp["visible_text"]["row_labels"]       # ['대회 설정', '이번 시험']
    W, H = 12.0, 4.6
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, rows)
    left, right, bottom, top = lab_w + 0.30, 0.35, 0.86, 0.25
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top)
    ax.set_xlim(10.5, 15)
    ax.set_ylim(1.6, -0.75)
    ax.axvspan(10.5, floor, fc=GRAY_ST, ec="none", zorder=0)
    pts = [(0, st, METHOD["source"], "Stefan"), (0, rs, METHOD["resid"], "물리 잔차 결합"),
           (1, p1, METHOD["recal"], "재보정 Stefan"), (1, r1, METHOD["resid"], None)]
    for yv, (a, b) in enumerate([(st, rs), (p1, r1)]):
        ax.plot([a, b], [yv, yv], color="black", lw=LW_AUX, zorder=2)
    for yv, x, col, lab in pts:
        ax.plot([x], [yv], marker="o", ms=MS, mfc=col, mec=col, ls="none", zorder=3)
        if lab:
            ax.annotate(lab, (x, yv), xytext=(0, 13), textcoords="offset points", ha="center", va="bottom",
                        fontsize=TK_PT, fontweight=W_SEMI, color=INK)
    ax.text((10.5 + floor) / 2, -0.55, "오차 하한", ha="center", va="center", fontsize=TK_PT, fontweight=W_SEMI,
            color=INK)
    cat_axis(ax, rows, [0, 1])
    ax.set_xticks([11, 12, 13, 14, 15])
    ax.set_xlabel("RMSE (cm)")
    df = pd.DataFrame([
        dict(row=rows[0], method="Stefan(대회, 지역 내 0.5° 블록 6겹)", rmse_cm=st,
             source=f"{rel(s4p)} part D_indomain, cv spatial_block_AK, ridge, shared25, seed 0, lam 0"),
        dict(row=rows[0], method="물리 잔차 결합(대회 최선 설정)", rmse_cm=rs,
             source=f"{rel(s4p)} 같은 조건, lam 0.75"),
        dict(row=rows[1], method="재보정 Stefan(라벨 1000개, 분할 25회)", rmse_cm=p1,
             source=f"{rel(cvp)} exp wf6, Alaska, n 1000, method P1"),
        dict(row=rows[1], method="물리 잔차 결합(교차검증 λ)", rmse_cm=r1,
             source=f"{rel(cvp)} exp wf6, Alaska, n 1000, method R1, lam −1(교차검증)"),
        dict(row="오차 하한", method="알래스카 오차 하한", rmse_cm=floor, source=f"{rel(flp)} Alaska, eval"),
    ])
    red = lambda x: x ** 2 - floor ** 2  # noqa: E731
    df["reducible_sq_cm2"] = [red(st), red(rs), red(p1), red(r1), np.nan]
    meta = dict(slide="S20", page=20, archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(s4p), rel(cvp), rel(flp)],
                notes=[f"줄일 수 있는 오차 제곱 감소 대회 {100 * (red(st) - red(rs)) / red(st):.2f}%, "
                       f"이번 {100 * (red(p1) - red(r1)) / red(p1):.2f}%(결론 줄 39%, 19% 와 일치 확인)",
                       "직접 라벨은 '재보정 Stefan'을 더했다(스펙의 'Stefan' 하나로는 두 기준 점의 색이 구분되지 않음)"])
    assert round(100 * (red(st) - red(rs)) / red(st)) == 39 and round(100 * (red(p1) - red(r1)) / red(p1)) == 19
    finish(fig, "S20", "contest_vs_strict", df, meta)


# ================================================================ S24 독립 지역 포레스트(오른쪽 그래프)
def chart_s24(spec):
    p = CLAIMS / "SI_learners_new_regions/tables/lgd_tests_lic.csv"
    t = pd.read_csv(p)
    pe1 = "MEAN[Lena|x,Canada|x,Russia_W|x,Russia_E|x,Russia_C|x]"

    def pick(q, label, group, desc):
        assert len(q) == 1, label
        r = q.iloc[0]
        return dict(group=group, row=label, contrast=desc, delta=r.delta, ci_lo=r.ci_lo, ci_hi=r.ci_hi,
                    delta_blockeq=r.delta_blockeq, ci_lo_beq=r.ci_lo_beq, ci_hi_beq=r.ci_hi_beq,
                    verdict4=r.verdict4 if isinstance(r.verdict4, str) else "",
                    filter=f"test_id={r.test_id}, item={r['item']}, scope={r.scope}, target={r.target}, pool={r.pool}, n={r.n}")

    g1, g2 = "원천 계수 Stefan 대비(전량)", "재보정 Stefan 대비(라벨 10개)"
    rows = [
        pick(t[(t.test_id == "L8e") & (t["item"] == "R1-P0") & (t.scope == "MEAN") & (t.pool == "PE1") & (t.n == -1)],
             "5지역 평균", g1, "물리 잔차 결합 − 원천 계수 Stefan, 전량"),
        pick(t[(t.test_id == "L38") & (t["item"] == "(b) R1-P0") & (t.target == "Russia_C|x")],
             "러시아 중부", g1, "물리 잔차 결합 − 원천 계수 Stefan, 전량"),
        pick(t[(t.test_id == "L4e") & (t.scope == "compare_pool") & (t.pool == "PE1") & (t.n == 10) & (t.target == pe1)],
             "5지역 평균", g2, "물리 잔차 결합 − 재보정 Stefan(잔차 추가 몫), 라벨 10개"),
        pick(t[(t.test_id == "L38") & (t["item"] == "(c) R1-P1") & (t.target == "Russia_C|x") & (t.n == 10)],
             "러시아 중부", g2, "물리 잔차 결합 − 재보정 Stefan(잔차 추가 몫), 라벨 10개"),
    ]
    df = pd.DataFrame(rows)
    W, H = 5.85, 4.80
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, df.row.tolist())
    left, right, bottom, top_band = lab_w + 0.28, 0.18, 0.86, 0.92
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    ypos = [0.85, 1.85, 3.85, 4.85]
    ax.set_ylim(5.45, 0.0)
    ax.set_xlim(-9.6, 2.2)
    fig.canvas.draw()
    off = 0.11 / in_per_data_y(fig, ax)
    ref_band_zero(ax)
    for y, r in zip(ypos, rows):
        forest_row(ax, y, r["delta"], r["ci_lo"], r["ci_hi"], r["delta_blockeq"], r["ci_lo_beq"], r["ci_hi_beq"],
                   "black", off)
    cat_axis(ax, df.row.tolist(), ypos)
    ax.set_xticks([-8, -6, -4, -2, 0, 2])
    ax.set_xlabel("물리 잔차 결합의 오차 변화 (cm)")
    fig.canvas.draw()
    for gy, g in ((0.0, g1), (3.0, g2)):
        yin = ax.transData.transform((0, gy + 0.05))[1] / fig.dpi
        ftext(fig, 0.05, yin, g, weight=W_SEMI, va="top", gid="category")
    yk1, yk2 = H - 0.22, H - 0.56
    key_line(fig, 0.05, yk1, [("line", "셀 가중 95% CI", dict(c="black", lw=LW_MAIN)),
                              ("line", "블록 등가중 95% CI", dict(c="black", lw=LW_AUX))])
    key_line(fig, 0.05, yk2, [("band", "±0.5 cm 동등 범위", {})])
    df["source"] = rel(p)
    meta = dict(slide="S24", page=24, archetype="A5(오른쪽 그래프)", geometry_in=dict(x=6.817, y=1.55, w=W, h=H),
                sources=[rel(p)],
                notes=["행을 비교 기준·라벨 수 두 묶음 머리와 지역 행으로 나누었다(같은 문구 반복 줄임, 행 이름 4단어 이하). 스펙 행 4개와 같은 대비",
                       "러시아 중부 전량은 L38 (b) 판 CI(스펙 numbers 지시)",
                       "5지역 라벨 10개는 compare_pool 행(판정어 '동등'이 붙은 판)",
                       "왼쪽 지도는 v3 Fig 1a 판 재사용(이 스크립트 범위 밖)"])
    finish(fig, "S24", "independent_forest", df, meta)


# ================================================================ S25 예측 구간 포함률
def chart_s25(spec):
    sp = spec["S25"]
    hp = CLAIMS / "SI_uncertainty/tables/h4__c2_coverage.csv"
    up = CLAIMS / "SI_uncertainty/tables/m1__transfer_uq_summary.csv"
    bp = CLAIMS / "SI_uncertainty/tables/lgu__lgu_b_tests.csv"
    h = pd.read_csv(hp)
    hh = h[(h.method == "hier2_cdf") & (h.test == "label0")]
    mean = float(hh[hh.target.str.startswith("MEAN6")].coverage.iloc[0])
    regs = ["Lena", "Canada", "Russia_W", "Russia_E"]
    hreg = {r: float(hh[hh.target == r].coverage.iloc[0]) for r in regs}
    u = pd.read_csv(up)
    uu = u[(u.method == "cqr_ak") & u.region.isin(regs) & u.cond.isin(["noinfo", "covonly"])]
    assert len(uu) == 8
    b = pd.read_csv(bp)
    nf = float(b[(b.test == "LGU-B1") & b.contrast.str.contains("cqr_pool:nflow@lam1.0", regex=False)].coverage_pool.iloc[0])
    cf = float(b[(b.test == "LGU-B1") & b.contrast.str.contains("cqr_pool:cfm@lam1.0", regex=False)].coverage_pool.iloc[0])
    rows = ["지역 단위 보정", "알래스카 보정 구간", "생성 분위 구간"]
    W, H = 12.0, 4.8
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, rows)
    left, right, bottom, top_band = lab_w + 0.30, 0.35, 0.86, 0.52
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    ax.set_xlim(0.25, 1.0)
    ax.set_ylim(2.55, -0.55)
    ax.axvline(0.9, color=AUX, lw=LW_ZERO, ls=(0, (1.2, 1.4)), zorder=1)
    # 행 0: 계층 conformal(9월 26일 규약) 지역별 값과 주 4지역 평균
    ax.plot(list(hreg.values()), [0] * 4, marker="o", ms=MS_SMALL, mfc="black", mec="black", alpha=0.5, ls="none", zorder=3)
    ax.plot([mean], [0], marker="o", ms=MS + 2, mfc="black", mec="black", ls="none", zorder=4)
    # 행 1: 알래스카 보정 CQR, 주 4지역 × 두 조건 8행의 범위
    ax.plot([uu.coverage.min(), uu.coverage.max()], [1, 1], color="black", lw=LW_MAIN, zorder=2)
    ax.plot(uu.coverage, [1] * len(uu), marker="o", ms=MS_SMALL, mfc="black", mec="black", alpha=0.5, ls="none", zorder=3)
    # 행 2: 원천 셀 교환성으로 보정한 생성 분위 구간 2종
    for x, lab in ((nf, "정규화 흐름"), (cf, "흐름 정합")):
        ax.plot([x], [2], marker="o", ms=MS + 2, mfc="black", mec="black", ls="none", zorder=4)
        ax.annotate(lab, (x, 2), xytext=(0, 12), textcoords="offset points", ha="center", va="bottom",
                    fontsize=TK_PT, fontweight=W_SEMI, color=INK)
    cat_axis(ax, rows, [0, 1, 2])
    ax.set_xticks([0.3, 0.5, 0.7, 0.9])
    ax.set_xticklabels(["0.3", "0.5", "0.7", "0.9"])
    ax.set_xlabel("포함률")
    fig.canvas.draw()
    xl = ax.transData.transform((0.9, 0))[0] / fig.dpi
    ftext(fig, xl, H - top_band + 0.08, "목표", ha="center", va="bottom", color=AUX)
    key_line(fig, 0.05, H - 0.22, [("dot", "주 4지역 평균", dict(fc="black", ms=MS + 2)),
                                   ("dot", "지역별 값", dict(fc="black", ms=MS_SMALL, alpha=0.5)),
                                   ("line", "범위", dict(c="black", lw=LW_MAIN))])
    recs = [dict(row=rows[0], item="주 4지역 평균(지역 등가중)", coverage=mean,
                 source=f"{rel(hp)} test label0, method hier2_cdf, MEAN6")]
    recs += [dict(row=rows[0], item=r, coverage=v, source=f"{rel(hp)} label0, hier2_cdf, {r}") for r, v in hreg.items()]
    recs += [dict(row=rows[1], item=f"{r.region}, {r.cond}", coverage=r.coverage,
                  source=f"{rel(up)} method cqr_ak") for r in uu.itertuples()]
    recs += [dict(row=rows[2], item="정규화 흐름(nflow, λ 1.0)", coverage=nf, source=f"{rel(bp)} LGU-B1 cqr_pool:nflow@lam1.0"),
             dict(row=rows[2], item="흐름 정합(cfm, λ 1.0)", coverage=cf, source=f"{rel(bp)} LGU-B1 cqr_pool:cfm@lam1.0")]
    df = pd.DataFrame(recs)
    meta = dict(slide="S25", page=25, archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(hp), rel(up), rel(bp)],
                notes=["점추정만 그린다(LGU 와 CI 를 나란히 두지 않음)",
                       "행 이름 '생성 모델 구간'을 '생성 분위 구간'으로 바꾸었다(SI_uncertainty README 1.2 용어 축소)",
                       "목표 0.90 은 눈금 0.9 와 '목표' 표지로 보인다(그림 안 수치 0개)",
                       "스펙 노트 대본의 '가로축은 구간 폭, 세로축은 실제 포함률'은 이 차트와 맞지 않는다(가로축 포함률, 세로 구간 종류)"])
    finish(fig, "S25", "coverage", df, meta)


# ================================================================ S29 원고 분량
def chart_s29(spec):
    sp = spec["S29"]
    vals = sp["evidence"]["chart"]["values"]          # {절: [현재, 목표]}
    check = {"서론": (797, 700), "결과": (7500, 2800), "고찰": (2012, 850), "방법": (6998, 4240)}
    for k, v in check.items():   # 문서 표(scirep_format_and_drafts.md 3.2, MANUSCRIPT_SPEC.md 1.1)와 대조
        assert tuple(vals[k]) == v, k
    rows = sp["visible_text"]["chart_rows"]
    W, H = 7.9, 4.75
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, rows)
    left, right, bottom, top_band = lab_w + 0.30, 0.40, 0.86, 0.50
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    for i, k in enumerate(rows):
        cur, tgt = vals[k]
        ax.plot([cur, tgt], [i, i], color=GRAY_ST, lw=LW_AUX, zorder=2)
        ax.plot([cur], [i], marker="o", ms=MS, mfc="white", mec="black", mew=1.5, ls="none", zorder=3)
        ax.plot([tgt], [i], marker="o", ms=MS, mfc="black", mec="black", ls="none", zorder=4)
    i = rows.index("고찰")
    ax.annotate("골격", (vals["고찰"][0], i), xytext=(12, 0), textcoords="offset points", ha="left", va="center",
                fontsize=TK_PT, fontweight=W_SEMI, color=INK)
    ax.set_xlim(0, 8000)
    ax.set_xticks([0, 2000, 4000, 6000, 8000])
    ax.set_xticklabels(["0", "2000", "4000", "6000", "8000"])
    ax.set_ylim(len(rows) - 0.5, -0.5)
    cat_axis(ax, rows, range(len(rows)))
    ax.set_xlabel("단어 수")
    key_line(fig, 0.05, H - 0.24, [("odot", "현재 초안", {}), ("dot", "목표", dict(fc="black"))])
    df = pd.DataFrame([dict(section=k, current_words=vals[k][0], target_words=vals[k][1]) for k in rows])
    df["source"] = ("docs/research/2026-10-04/scirep_format_and_drafts.md 3.2(현재, 마크다운 공백 토큰), "
                    "paper/manuscript/MANUSCRIPT_SPEC.md 1.1·6절(목표)")
    meta = dict(slide="S29", page=29, archetype="A3", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=["docs/research/2026-10-04/scirep_format_and_drafts.md 3.2", "paper/manuscript/MANUSCRIPT_SPEC.md 1.1"],
                notes=["스펙 높이 4.3 in 를 4.75 in 로 늘렸다(근거 블록 4.5 in 이상, 보조 표와 높이 일치)",
                       "고찰 현재 값은 골격 판이라 직접 라벨 '골격'"])
    finish(fig, "S29", "word_budget", df, meta)


# ================================================================ S32 추가 실험 실행 순서와 원고 자리
def chart_s32(spec):
    sp = spec["S32"]
    labels = sp["visible_text"]["row_labels"]
    places = sp["visible_text"]["right_labels"]
    # 등록 실행 순서(EXPERIMENT_PLAN_FINAL_BATCH 6절): XA → XD-alg → XC → XG → XH → XB → XE 1단계 → XF → XD-learn
    #   → XE 2단계 → XI → XJ. 본문/SI 는 같은 문서 0.1·6절(결과 전 고정)
    order = {"XA": [1], "XD": [2, 9], "XC": [3], "XG": [4], "XH": [5], "XB": [6], "XE": [7, 10], "XF": [8],
             "XI": [11], "XJ": [12]}
    ids = ["XA", "XD", "XC", "XG", "XH", "XB", "XE", "XF", "XI", "XJ"]
    main = {"XA", "XC", "XG", "XH"}
    extra = {("XD", 9): "학습형 정책", ("XE", 7): "1단계", ("XE", 10): "2단계"}
    W, H = 12.0, 5.2
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, labels)
    pl_w = max_text_w(fig, places + ["원고 자리"], weight=W_SEMI)
    left, right, bottom, top_band = lab_w + 0.30, pl_w + 0.45, 0.86, 0.50
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    n = len(ids)
    for i, e in enumerate(ids):
        xs = order[e]
        filled = e in main
        if len(xs) > 1:
            ax.plot([xs[0], xs[-1]], [i, i], color="black", lw=LW_AUX, zorder=2)
        for x in xs:
            ax.plot([x], [i], marker="o", ms=MS, mfc="black" if filled else "white", mec="black", mew=1.5,
                    ls="none", zorder=3)
            if (e, x) in extra:
                t = ax.annotate(extra[(e, x)], (x, i), xytext=(0, 9), textcoords="offset points", ha="center",
                                va="bottom", fontsize=TK_PT, fontweight=W_SEMI, color=INK)
                t.set_gid("label_num")
    ax.set_xlim(0.5, 12.5)
    ax.set_xticks(range(1, 13))
    ax.set_ylim(n - 0.5, -0.5)
    cat_axis(ax, labels, range(n))
    ax.set_xlabel("등록 실행 순서")
    fig.canvas.draw()
    xr = (left + (W - left - right)) + 0.30
    for i, pl in enumerate(places):
        yin = ax.transData.transform((0, i))[1] / fig.dpi
        ftext(fig, xr, yin, pl, gid="category")
    ftext(fig, xr, H - 0.24, "원고 자리", weight=W_SEMI, gid="category")
    key_line(fig, 0.05, H - 0.24, [("dot", "본문", dict(fc="black")), ("odot", "보충 자료", {})])
    df = pd.DataFrame([dict(row=lab, experiment_spec_only=e, order=";".join(map(str, order[e])),
                            main_or_si="본문" if e in main else "보충", manuscript_place=pl)
                       for lab, e, pl in zip(labels, ids, places)])
    df["source"] = "docs/EXPERIMENT_PLAN_FINAL_BATCH_2026-10-04.md 0.1(본문/SI), 6(가치 순서, 원고 절)"
    meta = dict(slide="S32", page=32, archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=["docs/EXPERIMENT_PLAN_FINAL_BATCH_2026-10-04.md 0.1, 5, 6"],
                notes=["실험 코드(XA–XJ)는 CSV 의 experiment_spec_only 열에만 둔다(보이는 글 0)",
                       "두 번 실행되는 묶음의 두 번째 점에 '학습형 정책', '2단계', 첫 점에 '1단계' 표지"])
    finish(fig, "S32", "registered_order", df, meta)


# ================================================================ AP3 라벨 0 물리 기준선
def chart_ap3(spec):
    lp = CLAIMS / "C1_label0_safety/tables/lgx_tests.csv"
    bp = CLAIMS / "C1_label0_safety/tables/lgw_bundle.csv"
    t = pd.read_csv(lp)
    b = pd.read_csv(bp)
    mean4 = "MEAN[Lena|x,Canada|x,Russia_W|x,Russia_E|x]"

    def rec(q, row, group, desc, f):
        assert len(q) == 1, row
        r = q.iloc[0]
        return dict(group=group, row=row, contrast=desc, delta=r.delta, ci_lo=r.ci_lo, ci_hi=r.ci_hi,
                    delta_blockeq=r.delta_blockeq, ci_lo_beq=r.ci_lo_beq, ci_hi_beq=r.ci_hi_beq,
                    verdict4=r.verdict4, pending=False, filter=f)

    g1, g2 = "원천 계수 Stefan 대비", "연도 정합 Stefan 대비"
    rows = [
        rec(t[(t.test_id == "L29") & (t.contrast == "P0@tddm-P0|n0") & (t.scope == "MEAN")], "연도 정합 Stefan", g1,
            "연도 정합 Stefan − 원천 계수 Stefan", f"{rel(lp)} test_id L29, contrast P0@tddm-P0|n0, MEAN"),
        rec(b[(b.ab == "AB2") & (b.scope == "MEAN")], "Stefan·위성 제품 앙상블", g1,
            "Stefan·위성 제품 앙상블 − 원천 계수 Stefan", f"{rel(bp)} ab AB2, MEAN"),
        dict(group=g1, row="기존 ALT 지도", contrast="기존 ALT 지도 − 원천 계수 Stefan", pending=True,
             filter="[XG: R2 | 결과 열람 뒤 설계, 등록 이탈(WRAPUP 10) | 2.7 사전 고정 해석 문장 XG-1, XG-2, XG-3, 공통 문장]"),
        rec(t[(t.test_id == "L28") & (t.contrast == "앵커 tddm|D0-P0|n0") & (t.scope == "MEAN")], "직접 ML", g2,
            "직접 ML − 연도 정합 Stefan", f"{rel(lp)} test_id L28, item X9, contrast 앵커 tddm|D0-P0|n0, MEAN"),
    ]
    for r in rows:
        if not r["pending"]:
            assert r.get("delta") is not None
    df = pd.DataFrame(rows)
    W, H = 12.0, 5.2
    fig = new_fig(W, H)
    lab_w = max(max_text_w(fig, df.row.tolist()), max_text_w(fig, [g1, g2], weight=W_SEMI) - 0.2)
    left, right, bottom, top_band = lab_w + 0.40, 0.35, 0.86, 0.55
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    ypos = [0.9, 1.9, 2.9, 4.9]
    ax.set_ylim(5.5, 0.0)
    ax.set_xlim(-4.5, 5.0)
    fig.canvas.draw()
    off = 0.11 / in_per_data_y(fig, ax)
    ref_band_zero(ax)
    for y, r in zip(ypos, rows):
        if r["pending"]:
            ax.text(0.75, y, "결과 대기", ha="left", va="center", fontsize=TK_PT, color=AUX, gid="category")
            continue
        forest_row(ax, y, r["delta"], r["ci_lo"], r["ci_hi"], r["delta_blockeq"], r["ci_lo_beq"], r["ci_hi_beq"],
                   "black", off)
    cat_axis(ax, df.row.tolist(), ypos)
    ax.set_xticks([-4, -2, 0, 2, 4])
    ax.set_xlabel("오차 변화 (cm)")
    fig.canvas.draw()
    for gy, g in ((0.05, g1), (4.05, g2)):
        yin = ax.transData.transform((0, gy))[1] / fig.dpi
        ftext(fig, 0.05, yin, g, weight=W_SEMI, va="top", gid="category")
    forest_key(fig, 0.05, H - 0.24)
    df["source"] = df["filter"]
    meta = dict(slide="AP3", page="부록 3", archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(lp), rel(bp)],
                notes=["행을 비교 기준 두 묶음으로 나누었다(행 이름 4단어 이하). 스펙의 4행과 같은 대비",
                       "기존 ALT 지도 행은 값 없이 '결과 대기' 표지(추가 실험 자리표시는 CSV 의 filter 열)"])
    finish(fig, "AP3", "label0_baselines", df, meta)


# ================================================================ AP5 알래스카 검증 방식 비교
def chart_ap5(spec):
    sp = spec["AP5"]
    p = CLAIMS / "C7_evaluation_design/tables/cv_scheme_comparison.csv"
    t = pd.read_csv(p)
    schemes = [("random_cell", "무작위 셀"), ("site_0.05", "지점 묶음"), ("block_0.5_canonical", "0.5° 블록"),
               ("knndm", "kNNDM")]
    assert [s[1] for s in schemes] == sp["visible_text"]["row_labels"]
    models = [("catboost_lo", "직접 CatBoost", METHOD["direct"], (0, (2, 1.5))),
              ("ridge", "선형 ML", METHOD["direct"], (0, (0.6, 1.4))),
              ("stefan", "Stefan", METHOD["source"], "-"),
              ("stefan_ridge_l075", "물리 잔차 결합", METHOD["resid"], "-")]
    assert [m[1] for m in models] == sp["visible_text"]["direct_labels"]
    recs = []
    for s, sl in schemes:
        for m, ml, _, _ in models:
            q = t[(t.scheme == s) & (t.model == m)]
            assert len(q) == 3
            recs.append(dict(scheme=sl, scheme_code=s, method=ml, model_code=m, rmse_cm_mean_3seeds=q.rmse_cm.mean(),
                             nnd_median_km_mean=q.nnd_median_km.mean()))
    df = pd.DataFrame(recs)
    chk = df.set_index(["scheme_code", "model_code"]).rmse_cm_mean_3seeds
    assert round(chk[("random_cell", "catboost_lo")], 2) == 11.55 and round(chk[("knndm", "catboost_lo")], 2) == 17.91
    W, H = 12.0, 5.2
    fig = new_fig(W, H)
    lab_w = max_text_w(fig, [s[1] for s in schemes])
    left, right, bottom, top_band = lab_w + 0.30, 0.35, 0.86, 0.55
    ax = add_ax(fig, left, bottom, W - left - right, H - bottom - top_band)
    dod = [-0.21, -0.07, 0.07, 0.21]
    for (m, ml, col, ls), dy in zip(models, dod):
        v = df[df.model_code == m].set_index("scheme_code").loc[[s for s, _ in schemes]].rmse_cm_mean_3seeds.values
        yy = np.arange(len(schemes)) + dy
        ax.plot(v, yy, color=col, lw=LW_AUX, ls=ls, zorder=2, dash_capstyle="butt")
        ax.plot(v, yy, marker="o", ms=MS, mfc=col, mec=col, ls="none", zorder=3)
    ax.set_xlim(10, 20)
    ax.set_xticks([10, 12, 14, 16, 18, 20])
    ax.set_ylim(len(schemes) - 0.5, -0.5)
    cat_axis(ax, [s[1] for s in schemes], range(len(schemes)))
    ax.set_xlabel("RMSE (cm)")
    key_line(fig, 0.05, H - 0.24, [("line", ml, dict(c=col, lw=LW_AUX, ls=ls, marker=True))
                                   for _, ml, col, ls in models])
    df["source"] = f"{rel(p)} (seed 0–2 평균, 0.5° 블록은 block_0.5_canonical = 대회 방식)"
    meta = dict(slide="AP5", page="부록 5", archetype="A2", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(p), "docs/QA_FINAL_REVIEW_2026-10-02.md Q15 표"],
                notes=["계열 4개가 0.1–0.3 cm 안에 몰린 행이 있어 직접 라벨 대신 열쇠 한 줄(4항목)과 행 안 세로 어긋남을 썼다",
                       "직접 CatBoost 와 선형 ML 은 직접 ML 색(#6b7280)을 공유하고 선 모양(파선, 점선)으로 구분한다",
                       "Stefan 은 지역 내 최소제곱 계수, 원천 계수 Stefan 의 색(#4d4d4d)"])
    finish(fig, "AP5", "cv_schemes", df, meta)


# ================================================================ AP8 원천 교차검증 조정(신경망 4종)
def chart_ap8(spec):
    sp = spec["AP8"]
    p = CLAIMS / "C5_method_selection/tables/lgfn_tests.csv"
    t = pd.read_csv(p)
    learners = [("mlp", "MLP"), ("tabm", "다중 헤드 MLP"), ("ftt", "축소 FT-Transformer"), ("realmlp", "RealMLP")]
    assert [l[1] for l in learners] == sp["visible_text"]["row_groups"]
    ns = [(0, "0", "라벨 0개"), (10, "10", "라벨 10개"), (-1, "all", "전량")]
    structs = [("D0", "1", "직접 ML", METHOD["direct"]), ("R1", "0.25", "물리 잔차 결합", METHOD["resid"])]
    recs = []
    for lc, ln in learners:
        for nv, ntag, nl in ns:
            for sc, lam, sl, col in structs:
                c = f"{sc}[{lc}*]-{sc}[{lc}]|n{ntag}|lam{lam}"
                q = t[(t.test_id == "LGF-N2") & (t.scope == "MEAN") & (t.contrast == c) & (t["item"] == "aux")]
                assert len(q) == 1, c
                r = q.iloc[0]
                recs.append(dict(learner=ln, n_label=nl, structure=sl, contrast_code=c, delta=r.delta, ci_lo=r.ci_lo,
                                 ci_hi=r.ci_hi, delta_blockeq=r.delta_blockeq, ci_lo_beq=r.ci_lo_beq,
                                 ci_hi_beq=r.ci_hi_beq, verdict4=r.verdict4, verdict4_common=r.verdict4_common,
                                 role=r.role))
    df = pd.DataFrame(recs)
    W, H = 12.0, 5.2
    fig = new_fig(W, H)
    rlab = [x[2] for x in ns]
    lab_w = max_text_w(fig, rlab)
    left, right, bottom, top_band, gap = lab_w + 0.30, 0.20, 0.95, 1.00, 0.32
    pw = (W - left - right - gap * 3) / 4
    ph = H - bottom - top_band
    axes = []
    for k, (lc, ln) in enumerate(learners):
        ax = add_ax(fig, left + k * (pw + gap), bottom, pw, ph)
        ax.set_xlim(-2.4, 4.6)
        ax.set_ylim(2.55, -0.55)
        axes.append(ax)
    fig.canvas.draw()
    off = 0.10 / in_per_data_y(fig, axes[0])
    for k, (lc, ln) in enumerate(learners):
        ax = axes[k]
        ref_band_zero(ax)
        for i, (nv, ntag, nl) in enumerate(ns):
            for (sc, lam, sl, col), dy in zip(structs, (-0.2, 0.2)):
                r = df[(df.learner == ln) & (df.n_label == nl) & (df.structure == sl)].iloc[0]
                forest_row(ax, i + dy - off / 2, r.delta, r.ci_lo, r.ci_hi, r.delta_blockeq, r.ci_lo_beq,
                           r.ci_hi_beq, col, off)
        ax.set_xticks([-2, 0, 2, 4])
        if k == 0:
            cat_axis(ax, rlab, range(3))
        else:
            ax.set_yticks([])
            ax.spines["left"].set_visible(False)
        fig.canvas.draw()
        bb = ax.get_window_extent(fig.canvas.get_renderer())
        ftext(fig, (bb.x0 + bb.x1) / 2 / fig.dpi, H - top_band + 0.10, ln, size=AX_PT, weight=W_SEMI, ha="center",
              va="bottom")
    x_mid = left + (W - left - right) / 2
    ftext(fig, x_mid, 0.22, "조정판 − 기본판 (cm)", size=AX_PT, ha="center")
    # 방향 표지 1개(지침 R-11): 첫 패널 아래 왼쪽, 축 방향 화살표 하나
    fig.canvas.draw()
    b0 = axes[0].get_window_extent(fig.canvas.get_renderer())
    x0 = b0.x0 / fig.dpi
    fig.add_artist(Line2D([x0 + 0.42, x0 + 0.02], [0.22, 0.22], transform=fig.dpi_scale_trans, color=AUX, lw=LW_AXIS,
                          marker=None))
    fig.add_artist(Line2D([x0 + 0.02], [0.22], transform=fig.dpi_scale_trans, color=AUX, marker="<", ms=7,
                          mfc=AUX, mec=AUX, ls="none"))
    ftext(fig, x0 + 0.52, 0.22, sp["visible_text"]["direction_marker"], color=AUX)
    key_line(fig, 0.05, H - 0.24, [("line", "직접 ML", dict(c=METHOD["direct"], lw=LW_MAIN, marker=True)),
                                   ("line", "물리 잔차 결합", dict(c=METHOD["resid"], lw=LW_MAIN, marker=True)),
                                   ("line", "셀 가중 95% CI", dict(c="black", lw=LW_MAIN)),
                                   ("line", "블록 등가중 95% CI", dict(c="black", lw=LW_AUX)),
                                   ("band", "±0.5 cm", {})], gap=0.30)
    df["source"] = f"{rel(p)} (test_id LGF-N2, scope MEAN, item aux = 주 대비, 잔차 λ 0.25)"
    meta = dict(slide="AP8", page="부록 8", archetype="A2(작은 다중 그림 4개, 같은 x)",
                geometry_in=dict(x=0.667, y=1.3, w=W, h=H), sources=[rel(p)],
                notes=["학습기 4패널 × 라벨 수 3행 × 구조 2계열(직접 ML, 물리 잔차 결합 λ 0.25 주 대비)",
                       "점 색은 대비가 속한 방법(직접 ML, 물리 잔차 결합)을 뜻한다. λ 1.0 보조 대비는 CSV 밖(스펙 노트 참고 줄에 FT-T −0.76)",
                       "N4 의 20점 가운데 14점은 그림에 넣지 않는다(서술, 노트)"])
    finish(fig, "AP8", "nn_tuning", df, meta)


# ================================================================ S09 Stefan 개념 좌표도(A8, 실제 레나델타 라벨 위)
S09_SEED, S09_N_LAB, S09_KAPPA = 0, 10, 10


def lena_labels():
    """레나델타 직접 라벨 행(ALT, √TDD)과 원천 계수 E0. 실험과 같은 행 단위(load_base 의 F4_direct 행)."""
    bp = ROOT / "data/processed/fidelity_base_v3.csv"
    sp = ROOT / "data/processed/paper_figs/fig1_source.csv"
    d = pd.read_csv(bp, usecols=["loc_id", "region", "source_id", "alt_cm", "e5_sqrt_tdd"])
    d = d[(d.region == "Lena_RU") & (d.source_id == "F4_direct") & d.alt_cm.notna() & d.e5_sqrt_tdd.notna()]
    d = d.reset_index(drop=True)
    e0 = float(pd.read_csv(sp).set_index("target").loc["Lena", "E0"])
    return d, e0, bp, sp


def chart_s09(spec):
    sp = spec["S09"]
    names = sp["visible_text"]["direct_labels"]          # ['원천 계수 Stefan', '재보정 Stefan', '잔차']
    d, E0, bp, srcp = lena_labels()
    assert round(E0, 2) == 1.62 and len(d) == 3037
    s_all, y_all = d.e5_sqrt_tdd.values, d.alt_cm.values
    idx = np.sort(np.random.RandomState(S09_SEED).choice(len(d), S09_N_LAB, replace=False))
    s_l, y_l = s_all[idx], y_all[idx]
    E_ls = float((s_l * y_l).sum() / (s_l ** 2).sum())           # 선택 라벨의 최소제곱 계수(원점 통과)
    E_n = (S09_N_LAB * E_ls + S09_KAPPA * E0) / (S09_N_LAB + S09_KAPPA)   # κ 수축 재보정(LG 계획 §3)
    resid = y_l - E_n * s_l
    # 잔차 선분 4개: 라벨 10개를 √TDD 순으로 늘어놓고 서로 다른 √TDD 를 가진 것 가운데 고르게 4개
    order = np.argsort(s_l)
    uniq = [i for k, i in enumerate(order) if k == 0 or s_l[i] != s_l[order[k - 1]]]
    pick = [uniq[int(round(q))] for q in np.linspace(0, len(uniq) - 1, 4)]
    pick = sorted(set(pick), key=lambda i: s_l[i])
    W, H = 12.0, 4.75
    XMAX, YMAX = 40.0, 120.0
    fig = new_fig(W, H)
    right_w = max_text_w(fig, names[:2], weight=W_SEMI) + 0.25
    left, bottom, top = 0.95, 0.86, 0.22
    ax = add_ax(fig, left, bottom, W - left - right_w, H - bottom - top)
    ax.plot(s_all, y_all, marker="o", ms=3.5, mfc="black", mec="none", alpha=0.35, ls="none", zorder=2, rasterized=False)
    xs = np.array([0.0, XMAX])
    ax.plot(xs, E0 * xs, color=METHOD["source"], lw=LW_AUX, zorder=3)
    ax.plot(xs, E_n * xs, color=METHOD["recal"], lw=LW_MAIN, zorder=4)
    for i in pick:
        ax.plot([s_l[i], s_l[i]], [E_n * s_l[i], y_l[i]], color=METHOD["resid"], lw=LW_AUX, zorder=5)
    ax.plot(s_l, y_l, marker="o", ms=MS, mfc="black", mec="black", ls="none", zorder=6)
    ax.set_xlim(0, XMAX)
    ax.set_ylim(YMAX, 0)                                   # 아래로 깊어짐
    ax.set_xticks([0, 10, 20, 30, 40])
    ax.set_yticks([0, 50, 100])
    ax.set_xlabel("√TDD (√(°C·일))")
    ax.set_ylabel("ALT (cm)")
    # 직접 라벨 3개: 두 직선은 오른쪽 끝(원천 계수는 화면 아래쪽 선), 잔차는 가장 긴 선분 옆
    ax.annotate(names[0], (XMAX, E0 * XMAX), xytext=(8, -5), textcoords="offset points", ha="left", va="top",
                fontsize=TK_PT, fontweight=W_SEMI, color=INK, annotation_clip=False)
    ax.annotate(names[1], (XMAX, E_n * XMAX), xytext=(8, 5), textcoords="offset points", ha="left", va="bottom",
                fontsize=TK_PT, fontweight=W_SEMI, color=INK, annotation_clip=False)
    j = max(pick, key=lambda i: abs(resid[i]))
    ax.annotate(names[2], (s_l[j], E_n * s_l[j] + resid[j] / 2), xytext=(9, 0), textcoords="offset points",
                ha="left", va="center", fontsize=TK_PT, fontweight=W_SEMI, color=INK)
    # 수식 1개(18 pt, 자료가 없는 왼쪽 아래 구역)
    ax.text(0.30, 0.16, sp["visible_text"]["formula_in_figure"], transform=ax.transAxes, ha="center", va="center",
            fontsize=AX_PT, fontweight=W_MED, color=INK)
    fig.canvas.draw()
    bb = ax.get_window_extent(fig.canvas.get_renderer())
    x0, y0, y1 = bb.x0 / fig.dpi, bb.y0 / fig.dpi, bb.y1 / fig.dpi
    # 축 범위 밖(ALT > YMAX) 행은 세로축 아래 끝의 삼각 표지 하나(지침 2.7 '축 끝의 작은 삼각 표지'). 수는 노트·CSV
    n_over = int((y_all > YMAX).sum())
    if n_over:
        fig.add_artist(Line2D([x0], [y0 - 0.09], transform=fig.dpi_scale_trans, color="black", marker="v", ms=5,
                              mfc="black", mec="black", ls="none"))
    # 방향 표지 1개(16 pt #4d4d4d, 쪽당 1개): 세로축 위쪽, 원점 아래의 빈 구역
    ya, yb = y1 - 0.95, y1 - 1.30
    fig.add_artist(Line2D([x0 + 0.16, x0 + 0.16], [ya, yb], transform=fig.dpi_scale_trans, color=AUX, lw=LW_AXIS))
    fig.add_artist(Line2D([x0 + 0.16], [yb], transform=fig.dpi_scale_trans, color=AUX, marker="v", ms=7, mfc=AUX,
                          mec=AUX, ls="none"))
    ftext(fig, x0 + 0.28, (ya + yb) / 2, "아래로 깊어짐", color=AUX)
    df = pd.DataFrame(dict(loc_id=d.loc_id.values[idx], sqrt_tdd=s_l, alt_cm=y_l, residual_cm=resid,
                           residual_segment=[i in pick for i in range(S09_N_LAB)]))
    df["E0"], df["E_ls_10"], df["E_recal_10"] = E0, E_ls, E_n
    df["n_rows_lena"], df["n_rows_over_ymax"], df["seed"], df["kappa"] = len(d), int((y_all > YMAX).sum()), S09_SEED, S09_KAPPA
    df["source"] = f"{rel(bp)} (region Lena_RU, source_id F4_direct, alt_cm, e5_sqrt_tdd); {rel(srcp)} (target Lena, E0)"
    meta = dict(slide="S09", page=9, archetype="A8", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(bp), rel(srcp), "docs/EXPERIMENT_PLAN_LG_2026-09-29.md §3(E_n = (n·E_ls + κ·E0)/(n + κ), κ = 10)"],
                notes=[f"예시 라벨 10개는 레나델타 행 {len(d)}개에서 RandomState({S09_SEED}) 추출(판정과 무관한 예시). "
                       f"E_ls {E_ls:.3f}, 재보정 E {E_n:.3f}, 원천 E0 {E0:.3f}",
                       f"세로축 0–{YMAX:.0f} cm(아래로 깊어짐). 범위 밖 행 {int((y_all > YMAX).sum())}개는 아래 축 끝 삼각 표지",
                       "잔차 선분은 서로 다른 √TDD 를 가진 라벨 가운데 4개. '잔차' 라벨은 가장 긴 선분 옆",
                       "기호 열쇠 줄('E Stefan 계수, TDD 융해 도일')은 그림 밖 슬라이드 글(A8, y 6.10)"])
    finish(fig, "S09", "stefan_concept", df, meta)


# ================================================================ S26·S27 라벨 수별 워크플로 경로(A9)
def workflow_values():
    """경로 값(그림 명세 9.3 d, 지침 6.7 d): 레나델타·캐나다 전이 풀. 라벨 3·10 은 재보정 앵커 + 저가중 잔차(λ 0.25)의
    풀 층화 평균, 라벨 40·160·전량은 교차검증 선정 규칙의 두 지역 단순 평균(새 집계, 판정 없음)."""
    pfp = ROOT / "data/processed/paper_figs/pool_fixed_curve.csv"
    wfp = ROOT / "results/rescale_wf/data/processed/wf/wf_curve.csv"
    pf = pd.read_csv(pfp)
    q = pf[(pf.edition == "E2_LenaCanada_all_n") & (pf.method == "R1") & (pf.lam == 0.25) & (pf.baseline == "P0")
           & (pf.scope == "MEAN")]
    recs = []
    for n in (3, 10):
        r = q[q.n == n]
        assert len(r) == 1
        recs.append(dict(interval="3–10", method="재보정과 저가중 잔차", n=n, value_cm=float(r.delta.iloc[0]),
                         ci_lo=float(r.ci_lo.iloc[0]), ci_hi=float(r.ci_hi.iloc[0]), verdict4=str(r.verdict4.iloc[0]),
                         source=f"{rel(pfp)} edition E2_LenaCanada_all_n, method R1, lam 0.25, baseline P0, scope MEAN, n {n}"))
    wf = pd.read_csv(wfp)
    w = wf[(wf.exp == "wf4") & (wf.method == "W") & (wf["mode"] == "x") & wf.target.isin(["Lena", "Canada"])]
    for n, iv in ((40, "40–160"), (160, "40–160"), (-1, "전량")):
        r = w[w.n == n]
        assert len(r) == 2 and set(r.target) == {"Lena", "Canada"}
        recs.append(dict(interval=iv, method="교차검증 선정", n=n, value_cm=float(r.d_p0.mean()),
                         lena_cm=float(r[r.target == "Lena"].d_p0.iloc[0]), canada_cm=float(r[r.target == "Canada"].d_p0.iloc[0]),
                         verdict4="새 집계, 판정 없음",
                         source=f"{rel(wfp)} exp wf4, method W, mode x, target Lena·Canada, n {n}, d_p0 두 지역 단순 평균"))
    recs.insert(0, dict(interval="0", method="원천 계수 Stefan", n=0, value_cm=0.0, verdict4="정의상 0", source="기준 방법 자체"))
    return pd.DataFrame(recs), pfp, wfp


def draw_workflow_path(fig, box, vals, labels, intervals, small=False):
    """경로 그림: n = 0 별도 축, n ≥ 3 로그 축, 전량 별도 축(지침 2.12 R-24). 돌려주는 값은 띠 그리기용 x 위치(in).
    small=True 는 결론 쪽 축소판(3.80 in): 눈금 숫자(320 생략)와 두 줄 라벨, 세로축 눈금 3개."""
    L, B, Wd, Hd = box
    n0_w, gap, all_w = (0.28, 0.24, 0.28) if small else (0.55, 0.15, 0.55)
    main_w = Wd - n0_w - all_w - 2 * gap
    ax0 = add_ax(fig, L, B, n0_w, Hd)
    axm = add_ax(fig, L + n0_w + gap, B, main_w, Hd)
    axa = add_ax(fig, L + n0_w + gap + main_w + gap, B, all_w, Hd)
    ylim = (-0.55, 1.45) if small else (-0.7, 1.3)
    for a in (ax0, axm, axa):
        a.set_ylim(*ylim)
        a.axhline(0, color=AUX, lw=LW_ZERO, zorder=1)
        a.tick_params(axis="x", length=5)
    for a in (axm, axa):
        a.spines["left"].set_visible(False)
        a.set_yticks([])
    for a in (ax0, axa):
        a.set_xlim(0, 1)
        a.set_xticks([])
    if small:
        ax0.set_yticks([0, 0.5, 1.0])
        ax0.set_yticklabels(["0", "0.5", "1.0"])
    else:
        ax0.set_yticks([-0.5, 0, 0.5, 1.0])
        ax0.set_yticklabels(["−0.5", "0", "0.5", "1.0"])
    ax0.set_ylabel("오차 변화 (cm)")
    axm.set_xscale("log")
    axm.set_xlim(2.2, 1500 if small else 1400)
    grid = [3, 10, 40, 160, 320, 1000]
    axm.xaxis.set_major_locator(FixedLocator(grid))
    axm.xaxis.set_minor_locator(NullLocator())
    axm.set_xticklabels(["3", "10", "40", "", "", "1000"] if small else [""] * len(grid))   # 축소판: 160·320 라벨 생략(간격 0.10 in)
    v = vals.set_index(["method", "n"]).value_cm
    # 구간 0: 원천 계수 Stefan 의 값 0(정의). 축 폭 전체의 3.0 pt 선분(측정점 없음)
    ax0.plot([0, 1], [0, 0], color=METHOD["source"], lw=LW_MAIN, solid_capstyle="butt", zorder=3)
    # 구간 3–10: 재보정과 저가중 잔차(재보정 앵커 + 잔차 ML 색, 실선)
    r1 = [v[("재보정과 저가중 잔차", 3)], v[("재보정과 저가중 잔차", 10)]]
    axm.plot([3, 10], r1, color=METHOD["resid"], lw=LW_MAIN, zorder=3)
    axm.plot([3, 10], r1, marker="o", ms=MS, mfc=METHOD["resid"], mec=METHOD["resid"], ls="none", zorder=4)
    # 구간 40–160과 전량: 교차검증 선정(고유 색 없음, 검정 경로). 320–1000 은 시험하지 않아 선분 없음
    wv = [v[("교차검증 선정", 40)], v[("교차검증 선정", 160)]]
    axm.plot([40, 160], wv, color="black", lw=LW_MAIN, zorder=3)
    axm.plot([40, 160], wv, marker="o", ms=MS, mfc="black", mec="black", ls="none", zorder=4)
    axa.plot([0.5], [v[("교차검증 선정", -1)]], marker="o", ms=MS, mfc="black", mec="black", ls="none", zorder=4)
    fig.canvas.draw()
    r = fig.canvas.get_renderer()
    b0, bm, ba = (a.get_window_extent(r) for a in (ax0, axm, axa))
    fx = lambda a, x: a.transData.transform((x, 0))[0] / fig.dpi  # noqa: E731  자료 x → 그림 in
    fy = lambda a, y: a.transData.transform((0, y))[1] / fig.dpi  # noqa: E731
    # 방법 라벨 3개(16 pt SemiBold, 각 1회). 그림 좌표로 두어 축 경계에 잘리지 않게 한다
    lab0, lab1, lab2 = labels
    if small:
        ftext(fig, b0.x1 / fig.dpi + 0.06, fy(ax0, 0) - 0.07, lab0.replace(" Stefan", "\nStefan"), weight=W_SEMI, va="top")
        ftext(fig, fx(axm, 10), fy(axm, r1[1]) + 0.12, lab1.replace(" 저가중", "\n저가중"), weight=W_SEMI, ha="center", va="bottom")
        ftext(fig, fx(axm, 160) + 0.12, fy(axm, wv[1]) + 0.06, lab2.replace(" ", "\n"), weight=W_SEMI, ha="left", va="center")
    else:
        ftext(fig, b0.x1 / fig.dpi + 0.10, fy(ax0, 0) - 0.08, lab0, weight=W_SEMI, va="top")
        ftext(fig, fx(axm, 10), fy(axm, r1[1]) + 0.14, lab1, weight=W_SEMI, ha="center", va="bottom")
        ftext(fig, fx(axm, np.sqrt(40 * 160)), fy(axm, (wv[0] + wv[1]) / 2) + 0.14, lab2, weight=W_SEMI, ha="center", va="bottom")
    pos = dict(x_left=b0.x0 / fig.dpi, x0_right=b0.x1 / fig.dpi, x_main_left=bm.x0 / fig.dpi, x_main_right=bm.x1 / fig.dpi,
               x_all_right=ba.x1 / fig.dpi, x3=fx(axm, 3), y_axis=bm.y0 / fig.dpi,
               x_centers=[(b0.x0 + b0.x1) / 2 / fig.dpi, fx(axm, np.sqrt(3 * 10)), fx(axm, np.sqrt(40 * 160)),
                          fx(axm, np.sqrt(320 * 1000)), (ba.x0 + ba.x1) / 2 / fig.dpi],
               x_bounds=[fx(axm, np.sqrt(10 * 40)), fx(axm, np.sqrt(160 * 320))])
    ya = pos["y_axis"]
    if small:
        for x, name in ((pos["x_centers"][0], intervals[0]), (pos["x_centers"][4], intervals[4])):
            ftext(fig, x, ya - 0.26, name, ha="center", gid="category")
        ftext(fig, (pos["x_main_left"] + pos["x_main_right"]) / 2, ya - 0.60, "라벨 수", size=AX_PT, ha="center")
    else:
        for xb in pos["x_bounds"]:                     # 구간 경계 긴 눈금 0.15 in, 1.5 pt
            fig.add_artist(Line2D([xb, xb], [ya, ya - 0.15], transform=fig.dpi_scale_trans, color=INK, lw=LW_AXIS))
        for x, name in zip(pos["x_centers"], intervals):
            ftext(fig, x, ya - 0.30, name, ha="center", gid="category")
        ftext(fig, (pos["x_main_left"] + pos["x_main_right"]) / 2, ya - 0.66, "라벨 수", size=AX_PT, ha="center")
    return pos


def chart_s26(spec):
    sp = spec["S26"]
    vt = sp["visible_text"]
    vals, pfp, wfp = workflow_values()
    W, H = 12.0, 5.2
    fig = new_fig(W, H)
    # 경로 그림 12.00 × 3.30(본문 위끝), 띠 머리 y0 + 3.60·4.35, 띠 높이 0.60(원형 A9). 모두 한 PNG 안에서 같은 x 기준
    path_h = 3.30
    left, top, bottom_in_path = 1.25, 0.12, 0.98
    pos = draw_workflow_path(fig, (left, H - path_h + bottom_in_path, W - left - 0.30, path_h - top - bottom_in_path),
                             vals, vt["method_labels"], vt["interval_names"])
    heads, place, diag = vt["band_heads"], vt["placement_band"], vt["diagnosis_band"]
    y_head1, y_line1 = H - 3.80, H - 4.08            # 띠 1: 다음 관측(y0 + 3.60 … 4.20)
    y_head2, y_line2 = H - 4.55, H - 4.83            # 띠 2: 진단(y0 + 4.35 … 4.95)
    tr = fig.dpi_scale_trans
    ftext(fig, 0.0, y_head1, heads[0], color="#6B6B6B", gid="category")
    ftext(fig, 0.0, y_head2, heads[1], color="#6B6B6B", gid="category")
    # 관측 띠: 분산 배치 0–1000(선 하나, 라벨 하나)
    fig.add_artist(Line2D([pos["x_left"], pos["x_main_right"]], [y_line1, y_line1], transform=tr, color=INK, lw=LW_AUX,
                          solid_capstyle="butt"))
    ftext(fig, (pos["x_left"] + pos["x_main_right"]) / 2, y_head1, place[0], weight=W_SEMI, ha="center", gid="category")
    # 진단 띠: 외삽 영역 표시(0), 계수 편향 진단(3–전량, 선 하나)
    fig.add_artist(Line2D([pos["x_left"], pos["x0_right"]], [y_line2, y_line2], transform=tr, color=INK, lw=LW_AUX,
                          solid_capstyle="butt"))
    ftext(fig, pos["x_left"], y_head2, diag[0], weight=W_SEMI, ha="left", gid="category")
    fig.add_artist(Line2D([pos["x3"], pos["x_all_right"]], [y_line2, y_line2], transform=tr, color=INK, lw=LW_AUX,
                          solid_capstyle="butt"))
    ftext(fig, (pos["x3"] + pos["x_all_right"]) / 2, y_head2, diag[1], weight=W_SEMI, ha="center", gid="category")
    n_text = len(vt["interval_names"]) + len(vt["method_labels"]) + len(place) + len(diag) + len(heads)
    assert n_text == vt["text_object_count"] == 13
    meta = dict(slide="S26", page=26, archetype="A9", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(pfp), rel(wfp), "docs/RESEARCH_CLAIMS_WORKFLOW_2026-10-01.md 4절", "FIGURE_SPEC_v3.md 9.3 d"],
                notes=["경로 그림(12.00 × 3.30)과 띠 머리·관측 띠·진단 띠를 한 PNG(12.00 × 5.20)에 그려 x 기준을 공유한다",
                       "세로축은 원천 계수 Stefan 대비 오차 변화(cm). 축 이름에 기준 방법을 넣지 않은 이유는 n = 0 직접 라벨과의 "
                       "문구 반복(H13)을 피하기 위함",
                       "라벨 3·10 은 레나델타·캐나다 풀 층화 평균(동등, 미결정). 40·160·전량은 교차검증 선정의 두 지역 단순 평균(새 집계, 판정 없음). "
                       "320–1000 은 이 풀에서 시험하지 않아 선분 없음",
                       "이 풀에서는 경로가 0 위에 놓인다(캐나다 재보정 오차 증가). y 기준·풀 선택은 사용자 결정(그림 명세 D-13)",
                       f"텍스트 객체 {n_text}개(구간 이름 5, 방법 라벨 3, 관측 1, 진단 2, 띠 머리 2). 축 제목 2개와 눈금은 별도",
                       "배치 띠 '분산 배치'는 추가 실험(사전 지정 배치 절차) 결과 뒤 다시 본다(스펙 flags)"])
    finish(fig, "S26", "workflow_path", vals, meta)


def chart_s27(spec):
    """결론 쪽 열 1: 워크플로 경로 축소판(3.80 × 3.40, 눈금 16 pt). 열 2·3(Fig 7b 지도, Fig 5 지도)은 v3 그림 범위."""
    sp = spec["S26"]
    vt = sp["visible_text"]
    vals, pfp, wfp = workflow_values()
    W, H = 3.80, 3.40
    fig = new_fig(W, H)
    left, top, bottom = 0.88, 0.14, 0.80
    draw_workflow_path(fig, (left, bottom, W - left - 0.08, H - top - bottom), vals, vt["method_labels"],
                       vt["interval_names"], small=True)
    meta = dict(slide="S27", page=27, archetype="A10(열 1 그림)", geometry_in=dict(x=0.667, y=1.3, w=W, h=H),
                sources=[rel(pfp), rel(wfp)],
                notes=["S26 과 같은 값·같은 색. 눈금 라벨은 3, 10, 40, 1000(160·320 은 눈금만, 라벨 간격 0.10 in 규칙)과 0·전량",
                       "띠(다음 관측, 진단)는 축소판에 두지 않는다. 열 2(알래스카 보정 지도)·열 3(캐나다 선정 지점 지도)은 v3 그림 범위"])
    finish(fig, "S27", "workflow_small", vals, meta)


# ================================================================ 수치 대조(스펙 numbers ↔ 원천 CSV)
def _nums(s):
    """문자열의 수(부호 포함)를 뽑는다. U+2212 와 '-' 를 음수로 본다. 백분율·단위는 무시."""
    return [float(x.replace("−", "-").replace(",", "")) for x in re.findall(r"[−-]?\d[\d,]*\.?\d*", s)]


def values_check(spec):
    """스펙의 numbers 값과 원천 CSV 값을 대조하고 values_check.txt 에 적는다. 모든 수치는 CSV 를 다시 열어 계산한다."""
    lines, n_ok, n_bad = [], 0, 0

    def chk(sid, label, spec_val, got, tol=0.0051, src=""):
        nonlocal n_ok, n_bad
        want = _nums(spec_val) if isinstance(spec_val, str) else list(spec_val)
        got = [float(g) for g in (got if isinstance(got, (list, tuple, np.ndarray, pd.Series)) else [got])]
        ok = len(want) == len(got) and all(abs(w - g) <= tol for w, g in zip(want, got))
        n_ok, n_bad = n_ok + ok, n_bad + (not ok)
        lines.append(f"[{'OK' if ok else 'MISMATCH'}] {sid} {label}\n    spec: {spec_val}\n    csv : "
                     f"{', '.join(f'{g:.4f}' for g in got)}\n    from: {src}")

    # S03 등록부 집계
    reg = pd.read_csv(ROOT / "paper/registry/experiments.csv")
    st = reg.status.astype(str).str.split("(").str[0].str.strip()
    cnt = lambda ph, s: int(((reg.phase == ph) & (st == s)).sum())  # noqa: E731
    chk("S03", "등록부 행 수(스펙 '47 묶음(등록부 64행)')", [64], [len(reg)], src="paper/registry/experiments.csv 행 수")
    chk("S03", "H0 25행: done 21, superseded 3, discarded 1",
        [25, 21, 3, 1], [int((reg.phase == "H0").sum()), cnt("H0", "done"), cnt("H0", "superseded"), cnt("H0", "discarded")],
        src="experiments.csv phase, status 첫 낱말")
    chk("S03", "T 8행 done, W 12행(done 11, not run 1), X 10행 planned",
        [8, 8, 12, 11, 1, 10, 10], [int((reg.phase == "T").sum()), cnt("T", "done"), int((reg.phase == "W").sum()),
                                   cnt("W", "done"), cnt("W", "not run"), int((reg.phase == "X").sum()), cnt("X", "planned")],
        src="experiments.csv")
    chk("S03", "paper_use main 행 17", [17], [int(reg.paper_use.astype(str).str.startswith("main").sum())], src="experiments.csv paper_use")
    # S06 자료 표 행 수
    m2 = json.load(open(ROOT / "data/processed/fidelity_base_v2_meta.json"))
    m3 = json.load(open(ROOT / "data/processed/fidelity_base_v3_meta.json"))
    m4 = json.load(open(ROOT / "data/processed/fidelity_base_v4_meta.json"))
    f1 = pd.read_csv(ROOT / "data/processed/paper_figs/fig1_meta.csv")
    chk("S06", "17,423 / 17,572 / 17,572 / 18,088", [17423, 17572, 17572, 18088],
        [m2["n_old"], m2["n_total"], m3["n_rows"], m4["n_rows"]["total"]], src="fidelity_base_v{2,3,4}_meta.json")
    chk("S06", "17,467 본 실험 직접 라벨 셀", [17467], [int(f1.n_v3_cells.iloc[0])], src="paper_figs/fig1_meta.csv n_v3_cells")
    nd = pd.read_csv(ROOT / "data/processed/paper_figs/fig1_not_drawn.csv")
    chk("S06", "약관 미확인 셀 38, 19, 1(캐나다, 북대서양, 러시아 서부)", [38, 19, 1],
        [int(nd[nd.iloc[:, 0].astype(str).str.contains(k)].iloc[:, -1].iloc[0]) for k in ("Canada", "NAtl", "Russia_W")]
        if len(nd) >= 3 else [np.nan], src="paper_figs/fig1_not_drawn.csv")
    # S08 검증 사다리
    v = pd.read_csv(CLAIMS / "C7_evaluation_design/tables/lgv_tests.csv")
    m = pd.read_csv(CLAIMS / "C7_evaluation_design/tables/lgv_metrics.csv")
    schemes = ["V-R", "V-P", "V-S", "V-B", "V-C0", "V-C100", "V-C500", "V-G"]
    d0 = [float(v[(v.test_id == "L20") & (v["item"] == f"D0[catboost_lo] RMSE@{s}") & (v.scope == "MEAN5")].rmse_A.iloc[0])
          for s in schemes]
    chk("S08", "직접 ML RMSE 22.91, 24.28, 25.61, 27.17, 30.31, 30.97, 31.13, 33.25", "22.91, 24.28, 25.61, 27.17, 30.31, 30.97, 31.13, 33.25", d0,
        src="lgv_tests.csv L20 rmse_A MEAN5")
    l19 = v[(v.test_id == "L19") & (v.contrast == "D0[catboost_lo]:V-G-V-R") & (v.scope == "MEAN5")].iloc[0]
    chk("S08", "지역 홀드아웃 − 셀 무작위 +10.34 [7.29, 12.84], 블록 등가중 +6.15 [4.70, 7.57]", "10.34, 7.29, 12.84, 6.15, 4.70, 7.57",
        [l19.delta, l19.ci_lo, l19.ci_hi, l19.delta_blockeq, l19.ci_lo_beq, l19.ci_hi_beq], src="lgv_tests.csv L19 MEAN5")
    l20 = v[(v.test_id == "L20") & (v.contrast == "D0[catboost_lo]:V-C500-V-R") & (v.scope == "MEAN5")].iloc[0]
    chk("S08", "500 km 군집 − 셀 무작위 +8.22 [5.65, 10.38]", "8.22, 5.65, 10.38", [l20.delta, l20.ci_lo, l20.ci_hi], src="lgv_tests.csv L20 MEAN5")
    ps = m[(m.method == "PS") & (m.scope == "MEAN5") & m.scheme.isin(schemes)].rmse
    chk("S08", "Stefan 최소제곱 26.81–26.96", "26.81, 26.96", [ps.min(), ps.max()], src="lgv_metrics.csv PS MEAN5")
    # S09
    _, e0, _, _ = lena_labels()
    chk("S09", "레나델타 원천 계수 E0 1.62", "1.62", [e0], src="paper_figs/fig1_source.csv target Lena")
    # S20
    s20 = pd.read_csv(SRC_DIR / "S20_contest_vs_strict.csv")
    st_, rs_, p1_, r1_, fl_ = s20.rmse_cm.values
    red = lambda x: x ** 2 - fl_ ** 2  # noqa: E731
    chk("S20", "대회 14.46 → 13.33 cm (7.80%)", "14.46, 13.33, 7.80", [st_, rs_, 100 * (st_ - rs_) / st_], src="s4_residual_results.csv")
    chk("S20", "대회 줄일 수 있는 오차 제곱 81.13 → 49.82 cm², 38.60%", "81.13, 49.82, 38.60",
        [red(st_), red(rs_), 100 * (red(st_) - red(rs_)) / red(st_)], src="파생(rmse² − 11.31²)")
    chk("S20", "이번 79.71 → 64.42 cm², 19.17% (재보정 Stefan 14.41 → 13.87, 3.75%)", "79.71, 64.42, 19.17, 14.41, 13.87, 3.75",
        [red(p1_), red(r1_), 100 * (red(p1_) - red(r1_)) / red(p1_), p1_, r1_, 100 * (p1_ - r1_) / p1_], src="wf2b_curve.csv, lgx_floor.csv")
    chk("S20", "오차 하한 11.31", "11.31", [fl_], src="lgx_floor.csv Alaska eval")
    # S24
    t = pd.read_csv(CLAIMS / "SI_learners_new_regions/tables/lgd_tests_lic.csv")
    vr = t[(t.test_id == "L1e") & (t.pool == "PE1") & (t["item"] == "verdict")].iloc[0]
    chk("S24", "5지역 풀 직접 ML 유의 개선 지역 0/5", [0, 5], [vr.k_improve, vr.n_regions], src="lgd_tests_lic.csv L1e verdict PE1")
    pe1 = "MEAN[Lena|x,Canada|x,Russia_W|x,Russia_E|x,Russia_C|x]"
    r = t[(t.test_id == "L8e") & (t["item"] == "R1-P0") & (t.scope == "MEAN") & (t.pool == "PE1") & (t.n == -1)].iloc[0]
    chk("S24", "5지역 전량 −3.19 [−4.19, −2.22], 블록 −2.62 [−3.71, −1.53]", "-3.19, -4.19, -2.22, -2.62, -3.71, -1.53",
        [r.delta, r.ci_lo, r.ci_hi, r.delta_blockeq, r.ci_lo_beq, r.ci_hi_beq], src="lgd_tests_lic.csv L8e PE1 MEAN n -1")
    r = t[(t.test_id == "L4e") & (t.scope == "compare_pool") & (t.pool == "PE1") & (t.n == 10) & (t.target == pe1)].iloc[0]
    chk("S24", "5지역 라벨 10 −0.13", "-0.13", [r.delta], src="lgd_tests_lic.csv L4e compare_pool PE1 n 10")
    r = t[(t.test_id == "L38") & (t["item"] == "(b) R1-P0") & (t.target == "Russia_C|x")].iloc[0]
    chk("S24", "러시아 중부 전량 −5.42 [−8.96, −1.64], 블록 −2.55 [−6.83, 1.48]", "-5.42, -8.96, -1.64, -2.55, -6.83, 1.48",
        [r.delta, r.ci_lo, r.ci_hi, r.delta_blockeq, r.ci_lo_beq, r.ci_hi_beq], src="lgd_tests_lic.csv L38 (b) Russia_C|x")
    r = t[(t.test_id == "L8e") & (t.target == "Tibet_LGD|x") & (t.n == -1)].iloc[0]
    chk("S24", "티베트 원천 계수 Stefan RMSE 244.51", "244.51", [r.rmse_B], src="lgd_tests_lic.csv L8e Tibet_LGD|x rmse_B")
    # S25
    h = pd.read_csv(CLAIMS / "SI_uncertainty/tables/h4__c2_coverage.csv")
    hh = h[(h.method == "hier2_cdf") & (h.test == "label0")]
    r = hh[hh.target.str.startswith("MEAN6")].iloc[0]
    chk("S25", "계층 conformal 0.86 [0.80, 0.92], 러시아 서부 0.75", "0.86, 0.80, 0.92, 0.75",
        [r.coverage, r.coverage_lo, r.coverage_hi, float(hh[hh.target == "Russia_W"].coverage.iloc[0])], src="h4__c2_coverage.csv")
    u = pd.read_csv(CLAIMS / "SI_uncertainty/tables/m1__transfer_uq_summary.csv")
    uu = u[(u.method == "cqr_ak") & u.region.isin(["Lena", "Canada", "Russia_W", "Russia_E"]) & u.cond.isin(["noinfo", "covonly"])]
    chk("S25", "알래스카 보정 CQR 포함률 0.32–0.73", "0.32, 0.73", [uu.coverage.min(), uu.coverage.max()], src="m1__transfer_uq_summary.csv cqr_ak")
    b = pd.read_csv(CLAIMS / "SI_uncertainty/tables/lgu__lgu_b_tests.csv")
    nf = float(b[(b.test == "LGU-B1") & b.contrast.str.contains("cqr_pool:nflow@lam1.0", regex=False)].coverage_pool.iloc[0])
    cf = float(b[(b.test == "LGU-B1") & b.contrast.str.contains("cqr_pool:cfm@lam1.0", regex=False)].coverage_pool.iloc[0])
    chk("S25", "생성 분위 구간 0.76, 0.65", "0.76, 0.65", [nf, cf], src="lgu__lgu_b_tests.csv LGU-B1")
    e4 = pd.read_csv(CLAIMS / "SI_uncertainty/tables/e4_interval_score.csv").set_index("interval")
    chk("S25", "알래스카 지역 내 CQR 0.45 → 0.93, 구간 점수 68.85 대 65.39(그림 밖, 노트)", "0.45, 0.93, 68.85, 65.39",
        [e4.loc["raw", "coverage"], e4.loc["cqr", "coverage"], e4.loc["cqr", "interval_score"], e4.loc["const_width", "interval_score"]],
        src="e4_interval_score.csv")
    # S26
    vals, _, _ = workflow_values()
    vv = vals.set_index(["method", "n"])
    chk("S26", "3–10 구간 +0.18, +0.78", "0.18, 0.78", [vv.loc[("재보정과 저가중 잔차", 3), "value_cm"], vv.loc[("재보정과 저가중 잔차", 10), "value_cm"]],
        src="pool_fixed_curve.csv E2_LenaCanada_all_n R1 0.25 MEAN")
    chk("S26", "교차검증 선정 +0.55, +0.18, 0.00(레나델타 +0.20, +0.02, −0.32, 캐나다 +0.90, +0.35, +0.32)",
        "0.55, 0.18, 0.00, 0.20, 0.02, -0.32, 0.90, 0.35, 0.32",
        [vv.loc[("교차검증 선정", n), "value_cm"] for n in (40, 160, -1)] + [vv.loc[("교차검증 선정", n), "lena_cm"] for n in (40, 160, -1)]
        + [vv.loc[("교차검증 선정", n), "canada_cm"] for n in (40, 160, -1)], src="wf_curve.csv wf4 W x d_p0")
    pf = pd.read_csv(ROOT / "data/processed/paper_figs/pool_fixed_curve.csv")
    q = pf[(pf.edition == "E1_P4_n_le_10") & (pf.method == "R1") & (pf.lam == 0.25) & (pf.baseline == "P0") & (pf.scope == "MEAN")]
    chk("S26", "대안 (c) 혼합 풀 −0.99, −2.63", "-0.99, -2.63", [float(q[q.n == 3].delta.iloc[0]), float(q[q.n == 10].delta.iloc[0])],
        src="pool_fixed_curve.csv E1_P4_n_le_10 R1 0.25 MEAN")
    # S29(문서 값: 문서 본문에 그 수가 있는지)
    doc = (ROOT / "docs/research/2026-10-04/scirep_format_and_drafts.md").read_text(encoding="utf-8")
    ms = (ROOT / "paper/manuscript/MANUSCRIPT_SPEC.md").read_text(encoding="utf-8")
    fmt = lambda v: (f"{v:,}", str(v))  # noqa: E731  문서는 네 자리 수에 쉼표를 쓴다(7,500)
    for val in (797, 7500, 2012, 6998):
        found = any(f in doc for f in fmt(val))
        lines.append(f"[{'OK' if found else 'MISMATCH'}] S29 현재 초안 단어 수 {val}\n    from: scirep_format_and_drafts.md 3.2 표 검색")
        n_ok, n_bad = n_ok + found, n_bad + (not found)
    for val in (700, 2800, 850, 4240):
        found = any(f in ms or f in doc for f in fmt(val))
        lines.append(f"[{'OK' if found else 'MISMATCH'}] S29 목표 단어 수 {val}\n    from: MANUSCRIPT_SPEC.md·scirep_format_and_drafts.md 본문 검색")
        n_ok, n_bad = n_ok + found, n_bad + (not found)
    # S32
    s32 = pd.read_csv(SRC_DIR / "S32_registered_order.csv")
    chk("S32", "추가 실험 10묶음, 본문 4, 보충 6", [10, 4, 6], [len(s32), int((s32.main_or_si == "본문").sum()), int((s32.main_or_si == "보충").sum())],
        src="source_data/S32_registered_order.csv(문서 0.1·6절 옮김)")
    # AP3
    lp = pd.read_csv(CLAIMS / "C1_label0_safety/tables/lgx_tests.csv")
    bp = pd.read_csv(CLAIMS / "C1_label0_safety/tables/lgw_bundle.csv")
    r = lp[(lp.test_id == "L29") & (lp.contrast == "P0@tddm-P0|n0") & (lp.scope == "MEAN")].iloc[0]
    chk("AP3", "연도 정합 − 원천 −1.00 [−1.35, −0.32], 블록 −0.41", "-1.00, -1.35, -0.32, -0.41",
        [r.delta, r.ci_lo, r.ci_hi, r.delta_blockeq], src="lgx_tests.csv L29 MEAN")
    r = bp[(bp.ab == "AB2") & (bp.scope == "MEAN")].iloc[0]
    chk("AP3", "앙상블 −2.73 [−3.33, −0.91], 블록 +0.05", "-2.73, -3.33, -0.91, 0.05", [r.delta, r.ci_lo, r.ci_hi, r.delta_blockeq],
        src="lgw_bundle.csv AB2 MEAN")
    r = lp[(lp.test_id == "L28") & (lp.contrast == "앵커 tddm|D0-P0|n0") & (lp.scope == "MEAN")].iloc[0]
    chk("AP3", "직접 ML − 연도 정합 +3.24 [1.86, 4.29]", "3.24, 1.86, 4.29", [r.delta, r.ci_lo, r.ci_hi], src="lgx_tests.csv L28 MEAN")
    # AP5
    a5 = pd.read_csv(SRC_DIR / "AP5_cv_schemes.csv")
    g = lambda mc: a5[a5.model_code == mc].rmse_cm_mean_3seeds  # noqa: E731
    chk("AP5", "직접 CatBoost 11.55 → 17.91", "11.55, 17.91", [g("catboost_lo").min(), g("catboost_lo").max()], src="cv_scheme_comparison.csv seed 평균")
    chk("AP5", "Stefan 14.24–14.46", "14.24, 14.46", [g("stefan").min(), g("stefan").max()], src="cv_scheme_comparison.csv")
    chk("AP5", "물리 잔차 결합 12.77–13.70", "12.77, 13.70", [g("stefan_ridge_l075").min(), g("stefan_ridge_l075").max()], src="cv_scheme_comparison.csv")
    chk("AP5", "선형 ML 12.65–14.11", "12.65, 14.11", [g("ridge").min(), g("ridge").max()], src="cv_scheme_comparison.csv")
    # AP8
    f = pd.read_csv(CLAIMS / "C5_method_selection/tables/lgfn_tests.csv")
    n2 = f[(f.test_id == "LGF-N2") & (f.scope == "MEAN") & (f["item"] == "aux")]
    gv = lambda c: float(n2[n2.contrast == c].delta.iloc[0])  # noqa: E731
    chk("AP8", "두 CI 일치 오차 증가 +2.96(MLP 10), +1.33, +2.56, +1.67(다중 헤드 0, 10, 전량)", "2.96, 1.33, 2.56, 1.67",
        [gv("D0[mlp*]-D0[mlp]|n10|lam1"), gv("D0[tabm*]-D0[tabm]|n0|lam1"), gv("D0[tabm*]-D0[tabm]|n10|lam1"),
         gv("D0[tabm*]-D0[tabm]|nall|lam1")], src="lgfn_tests.csv LGF-N2 MEAN aux")
    chk("AP8", "축소 FT-Transformer 전량 잔차 −0.19(λ 0.25), −0.76(λ 1.0)", "-0.19, -0.76",
        [gv("R1[ftt*]-R1[ftt]|nall|lam0.25"), gv("R1[ftt*]-R1[ftt]|nall|lam1")], src="lgfn_tests.csv LGF-N2 MEAN aux")
    n4 = f[(f.test_id == "LGF-N4")]
    txt = " ".join(str(x) for x in n4.astype(str).values.ravel())
    found = "14/20" in txt or ("14" in txt and "20" in txt)
    lines.append(f"[{'OK' if found else 'MISMATCH'}] AP8 14/20(조정이 설정을 바꾼 점 가운데 대상 오차가 커진 점)\n    from: lgfn_tests.csv LGF-N4 행 문자열 검색")
    n_ok, n_bad = n_ok + found, n_bad + (not found)
    head = [f"덱 전용 차트 수치 대조 {datetime.now().strftime('%Y-%m-%d %H:%M')}  일치 {n_ok}  불일치 {n_bad}",
            "스펙 deck/deck_spec_paper_report.json 의 numbers 와 원천 CSV(다시 열어 계산)의 대조. 허용 오차 0.005(반올림 두 자리)", ""]
    (OUT / "values_check.txt").write_text("\n".join(head + lines) + "\n", encoding="utf-8")
    print(f"[values] 일치 {n_ok} 불일치 {n_bad} → {rel(OUT / 'values_check.txt')}")
    return n_ok, n_bad


# ================================================================ 헤어라인 표
TBL = dict(header_h=0.50, top_lw=1.25, top_c="4A4A4A", mid_lw=0.75, mid_c="DDDDDD", head_pt=18, cell_pt=18,
           head_font="Pretendard SemiBold", cell_font="Pretendard Medium", head_color="4A4A4A", cell_color="111111",
           accent="EA851B", col_gap=0.18, line_h=0.30, safety=0.95)
_PIL = {}


def text_w_pt(s, weight="Medium", pt=18):
    key = (weight, pt)
    if key not in _PIL:
        _PIL[key] = ImageFont.truetype(str(FONT_DIR / f"Pretendard-{weight}.otf"), size=pt * 10)
    return _PIL[key].getlength(s) / 10.0


NBSP_UNITS = r"(cm|cm²|단어|in|pt|km|m|%|쪽|편|개|대비|종)"


def nbsp(s):
    """숫자와 단위를 줄바꿈 없는 공백으로 묶는다(지침 5.9)."""
    return re.sub(r"(\d)\s" + NBSP_UNITS + r"(?![가-힣A-Za-z])", lambda m: m.group(1) + " " + m.group(2), s)


def wrap(s, width_in, weight="Medium", pt=18):
    """낱말 단위 줄바꿈(keep-all). 공백에서만 나눈다(NBSP 는 묶음). 두 줄이면 구 경계(쉼표, 빗금 뒤)를
    우선하는 균형 분할, 세 줄 이상이면 한두 글자 고아 줄을 앞 줄 낱말과 합친다(지침 5.9, S-09)."""
    maxw = width_in * 72 * TBL["safety"]
    W = lambda x: text_w_pt(x, weight, pt)  # noqa: E731
    words = s.split(" ")
    lines, cur = [], ""
    for w in words:
        cand = w if not cur else cur + " " + w
        if W(cand) <= maxw or not cur:
            cur = cand
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    if len(lines) == 2:
        best = None
        for k in range(1, len(words)):
            a, b = " ".join(words[:k]), " ".join(words[k:])
            if W(a) > maxw or W(b) > maxw:
                continue
            score = max(W(a), W(b))
            if words[k - 1].endswith((",", "/")) or words[k] == "/":
                score -= 0.15 * maxw
            if b.startswith(("/", ",", ")")):
                score += maxw
            if len(b.replace("\u00a0", "")) <= 2:
                score += maxw
            if best is None or score < best[0]:
                best = (score, [a, b])
        if best:
            lines = best[1]
    elif len(lines) >= 3:
        last = lines[-1]
        if len(last.replace("\u00a0", "").strip()) <= 2 and " " in lines[-2]:
            head, tail = lines[-2].rsplit(" ", 1)
            if W(tail + " " + last) <= maxw:
                lines[-2], lines[-1] = head, tail + " " + last
    over = [ln for ln in lines if W(ln) > maxw]
    return lines, over


TABLE_IDS = ["S02", "S04", "S06", "S10", "S29", "S30", "S31", "S33", "S34", "AP2", "AP4", "AP6", "AP7", "AP9", "AP10"]


def table_defs(spec):
    """스펙 visible_text.table 에서 표 정의를 만든다. 문구는 스펙 그대로이고 바꾼 곳은 changes 에 적는다."""
    defs = {}
    acc = {"S10": ("기본 답(앵커)",), "S31": ("수치만(기본안)",)}
    for sid in TABLE_IDS:
        s = spec[sid]
        vt = s["visible_text"]
        tb = vt["support_table"] if sid in ("S06", "S29") else vt["table"]
        geo = dict(tb["geometry_in"])
        cols = list(tb["columns"])
        rows = [list(r) for r in tb["rows"]]
        changes = []
        if sid == "S06":       # 행 수를 원천 meta 에서 다시 읽어 대조(손 수치 0)
            m2 = json.load(open(ROOT / "data/processed/fidelity_base_v2_meta.json"))
            m3 = json.load(open(ROOT / "data/processed/fidelity_base_v3_meta.json"))
            m4 = json.load(open(ROOT / "data/processed/fidelity_base_v4_meta.json"))
            f1 = pd.read_csv(ROOT / "data/processed/paper_figs/fig1_meta.csv")
            got = [m2["n_old"], m2["n_total"], m3["n_rows"], m4["n_rows"]["total"], int(f1.n_v3_cells.iloc[0])]
            for r, v in zip(rows, got):
                assert r[1] == f"{v:,}", (r, v)
            geo.update(y=1.30 + 0.50, row_h=0.79)   # 차트 자료 영역 위끝에 맞추고 차트 아래끝과 맞춘다(A3)
            changes.append("y 1.30 → 1.80, 행 높이 0.75 → 0.79 in(A3: 보조 요소를 주 그림 자료 영역 위끝과 아래끝에 맞춤)")
        if sid == "S29":
            fix = {"200단어 이하, 계획 199(판 B)와 196(판 A)": "200단어 이하, 계획 196–199",
                   "350단어 이하, v3 초안 268–343": "350단어 이하, 초안 268–343"}
            for r in rows:
                if r[1] in fix:
                    changes.append(f"'{r[1]}' → '{fix[r[1]]}'(18 pt, 열 2.6 in 에서 두 줄 안. 'v3' 는 보이는 글의 내부 판 이름이라 뺌, 지침 H3. 판 A·B 구분은 노트)")
                    r[1] = fix[r[1]]
            geo.update(y=1.30 + 0.50, row_h=0.76)
            changes.append("y 1.30 → 1.80, 행 높이 0.80 → 0.76 in(A3 높이 일치, 차트 4.75 in)")
        align = ["r" if (sid == "S06" and j == 1) else "l" for j in range(len(cols))]
        support = None
        if vt.get("support_line"):
            sl = vt["support_line"]
            if "  " in sl:
                lead, body = sl.split("  ", 1)
            else:
                lead, body = "", sl
            support = dict(lead=lead, text=body)
        accent = []
        for r_i, r in enumerate(rows):
            for c_i, cell in enumerate(r):
                if cell in acc.get(sid, ()):
                    accent.append([r_i, c_i])
        defs[sid] = dict(slide=sid, page=s["page"], title=s["title"], subtitle=s["subtitle"], takeaway=s["takeaway"],
                         archetype=s["archetype"], kind="native_hairline_table(python-pptx 텍스트 상자 + 규칙선)",
                         geometry_in=geo, columns=cols, rows=rows, align=align, accent_cells=accent,
                         support_line=support, changes_from_spec=changes, style=TBL,
                         data_sources=s["evidence"].get("data_sources", []))
    return defs


def layout_table(t):
    """열 폭에 맞춰 줄을 나누고 넘침을 점검한다."""
    g = t["geometry_in"]
    cw = g["col_w"]
    gap = TBL["col_gap"]
    issues = []
    head = []
    for j, c in enumerate(t["columns"]):
        ls, over = wrap(nbsp(c), cw[j] - gap, "SemiBold")
        head.append(ls)
        if len(ls) > 1:
            issues.append(f"머리 '{c}' {len(ls)}줄")
    body = []
    max_lines = max(1, int((g["row_h"] - 0.08) / TBL["line_h"]))
    for i, r in enumerate(t["rows"]):
        rr = []
        for j, c in enumerate(r):
            ls, over = wrap(nbsp(c), cw[j] - gap, "Medium")
            rr.append(ls)
            if over:
                issues.append(f"행 {i + 1} 열 {j + 1} 낱말 하나가 열보다 넓음: {over}")
            if len(ls) > max_lines:
                issues.append(f"행 {i + 1} 열 {j + 1} {len(ls)}줄 > {max_lines}줄: '{c}'")
            if any(len(ln.replace(' ', ' ').strip()) <= 1 for ln in ls[1:]):
                issues.append(f"행 {i + 1} 열 {j + 1} 고아 줄")
        body.append(rr)
    for txt in [*t["columns"], *[c for r in t["rows"] for c in r]]:
        if CODE_RE.search(txt):
            issues.append(f"내부 코드 의심: {txt}")
        if DASH_RE.search(re.sub(r"수[십백천]–수[십백천]", "", txt)):
            issues.append(f"연결어 대시: {txt}")
        if re.search(r"(다|니다)\.?$", txt.strip()):
            issues.append(f"'다' 종결: {txt}")
    t["wrapped_header"] = head
    t["wrapped_rows"] = body
    t["height_in"] = round(g.get("header_h", TBL["header_h"]) + g["row_h"] * len(t["rows"]), 3)
    t["max_lines_per_cell"] = max_lines
    t["issues"] = issues
    return t


def _pptx_imports():
    sys.path.insert(0, str(DECK))
    import final_lib as fl  # noqa: E402  (읽기 전용 사용: text, hline, _set_fonts)
    return fl


def build_hairline_table(slide, t, fl=None):
    """원어 헤어라인 표(A7): 위·아래 1.25 pt #4A4A4A, 머리 아래 0.75 pt #DDDDDD, 머리 18 pt SemiBold #4A4A4A,
    셀 18 pt Medium #111111, 셀은 세로 가운데, 줄은 미리 낱말 단위로 나눈 문단. 채움 0, 세로선 0."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    fl = fl or _pptx_imports()
    g = t["geometry_in"]
    x, y, w = g["x"], g["y"], g["w"]
    hh, rh, cw = g.get("header_h", TBL["header_h"]), g["row_h"], g["col_w"]
    gap = TBL["col_gap"]
    col = lambda h: RGBColor.from_string(h)  # noqa: E731
    fl.hline(slide, x, y, w, color=col(TBL["top_c"]), lw=TBL["top_lw"])

    def cell(xx, yy, ww, hgt, lines, font, colr, align):
        paras = [[{"t": ln, "size": 18, "color": colr, "font": font}] for ln in lines]
        fl.text(slide, xx, yy, ww, hgt, paras, align=PP_ALIGN.RIGHT if align == "r" else PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0, space_after=0, wrap=False)

    xx = x
    for j, ls in enumerate(t["wrapped_header"]):
        cell(xx if t["align"][j] == "l" else xx + gap, y, cw[j] - gap, hh, ls, TBL["head_font"],
             col(TBL["head_color"]), t["align"][j])
        xx += cw[j]
    fl.hline(slide, x, y + hh, w, color=col(TBL["mid_c"]), lw=TBL["mid_lw"])
    acc = {tuple(a) for a in t["accent_cells"]}
    for i, rr in enumerate(t["wrapped_rows"]):
        xx = x
        yy = y + hh + i * rh
        for j, ls in enumerate(rr):
            c = col(TBL["accent"]) if (i, j) in acc else col(TBL["cell_color"])
            cell(xx if t["align"][j] == "l" else xx + gap, yy, cw[j] - gap, rh, ls, TBL["cell_font"], c,
                 t["align"][j])
            xx += cw[j]
    yb = y + hh + rh * len(t["wrapped_rows"])
    fl.hline(slide, x, yb, w, color=col(TBL["top_c"]), lw=TBL["top_lw"])
    if t.get("support_line"):
        s = t["support_line"]
        width = w
        runs = []
        if s["lead"]:
            runs.append({"t": s["lead"] + "  ", "size": 18, "color": col("111111"), "font": "Pretendard SemiBold"})
        runs.append({"t": nbsp(s["text"]), "size": 18, "color": col("111111"), "font": "Pretendard Medium"})
        fl.text(slide, x, yb + 0.22, width, 0.80, [runs], line_spacing=1.1, space_after=0)
    return yb


# ================================================================ 미리보기(머리, 근거, 결론 줄)
def preview(spec, defs, charts):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches
    fl = _pptx_imports()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    ORANGE, GRAY = RGBColor(0xEA, 0x85, 0x1B), RGBColor(0x6B, 0x6B, 0x6B)
    order = ["S02", "S03", "S04", "S06", "S08", "S09", "S10", "S20", "S24", "S25", "S26", "S27", "S29", "S30", "S31",
             "S32", "S33", "S34", "AP2", "AP3", "AP4", "AP5", "AP6", "AP7", "AP8", "AP9", "AP10"]
    pages = []
    for sid in order:
        s = spec[sid]
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        fl.text(sl, 0.667, 0.24, 11.25, 0.55, [{"t": s["title"], "size": 26, "color": ORANGE, "font": "Pretendard ExtraBold"}])
        fl.text(sl, 0.667, 0.80, 11.25, 0.34, [{"t": s["subtitle"], "size": 14, "color": GRAY, "font": "Pretendard Medium"}])
        fl.text(sl, 12.067, 0.22, 0.60, 0.30, [{"t": str(s["page"]), "size": 12, "color": GRAY, "font": "Pretendard Medium"}],
                align=PP_ALIGN.RIGHT)
        if sid in charts:
            c = MANIFEST["charts"][sid]
            g = c["geometry_in"]
            sl.shapes.add_picture(str(ROOT / c["file"]), Inches(g["x"]), Inches(g["y"]), Inches(c["size_in"][0]),
                                  Inches(c["size_in"][1]))
        if sid in defs:
            build_hairline_table(sl, defs[sid], fl)
        if sid == "S09":      # A8 기호 열쇠 줄(수식이 있는 쪽): y0 + H − 0.40, 18 pt Medium 한 줄
            fl.text(sl, 0.667, 6.10, 12.0, 0.40, [{"t": s["visible_text"]["symbol_key_line"], "size": 18,
                                                   "color": RGBColor(0x11, 0x11, 0x11), "font": "Pretendard Medium"}])
        if s.get("takeaway"):
            fl.hline(sl, 0.667, 6.82, 12.0, color=RGBColor(0xDD, 0xDD, 0xDD), lw=0.9)
            fl.text(sl, 0.667, 6.88, 12.0, 0.36, [{"t": s["takeaway"], "size": 15, "color": RGBColor(0x11, 0x11, 0x11),
                                                    "font": "Pretendard ExtraBold"}], align=PP_ALIGN.CENTER)
        pages.append(sid)
    pptx_path = OUT / "preview_evidence.pptx"
    prs.save(pptx_path)
    return pptx_path, pages


def export_pdf_no_autospace(pptx_path, pdf_dir):
    """PPTX → ODP → 문단 속성에 style:text-autospace="none" → PDF(지침 5.6, R-22)."""
    tmp = Path(tempfile.mkdtemp(prefix="prev_", dir=str(PREV_DIR)))
    prof = f"file://{tmp}/lo_profile"
    soff = shutil.which("soffice")
    subprocess.run([soff, f"-env:UserInstallation={prof}", "--headless", "--convert-to", "odp", "--outdir", str(tmp),
                    str(pptx_path)], check=True, capture_output=True, timeout=600)
    odp = tmp / (pptx_path.stem + ".odp")
    patched = tmp / (pptx_path.stem + "_na.odp")
    with zipfile.ZipFile(odp) as zin, zipfile.ZipFile(patched, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ("styles.xml", "content.xml"):
                s = data.decode("utf-8")
                s = re.sub(r'style:text-autospace="[^"]*"', "", s)
                s = s.replace("<style:paragraph-properties", '<style:paragraph-properties style:text-autospace="none"')
                data = s.encode("utf-8")
            ct = zipfile.ZIP_STORED if item.filename == "mimetype" else zipfile.ZIP_DEFLATED
            zout.writestr(item, data, compress_type=ct)
    subprocess.run([soff, f"-env:UserInstallation={prof}", "--headless", "--convert-to", "pdf", "--outdir", str(tmp),
                    str(patched)], check=True, capture_output=True, timeout=600)
    pdf = pdf_dir / "preview_evidence.pdf"
    shutil.move(str(tmp / (patched.stem + ".pdf")), pdf)
    shutil.rmtree(tmp, ignore_errors=True)
    return pdf


# ================================================================ 실행
CHARTS = {"S03": chart_s03, "S06": chart_s06, "S08": chart_s08, "S09": chart_s09, "S20": chart_s20, "S24": chart_s24,
          "S25": chart_s25, "S26": chart_s26, "S27": chart_s27, "S29": chart_s29, "S32": chart_s32, "AP3": chart_ap3,
          "AP5": chart_ap5, "AP8": chart_ap8}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default="")
    ap.add_argument("--no-preview", action="store_true")
    a = ap.parse_args()
    for d in (OUT, SRC_DIR, TAB_DIR, PREV_DIR):
        d.mkdir(parents=True, exist_ok=True)
    setup_mpl()
    spec = load_spec()
    only = [s for s in a.only.split(",") if s]
    old = OUT / "MANIFEST.json"
    if only and old.exists():
        prev = json.load(open(old, encoding="utf-8"))
        MANIFEST["charts"].update(prev.get("charts", {}))
        MANIFEST["tables"].update(prev.get("tables", {}))
    for sid, fn in CHARTS.items():
        if not only or sid in only:
            fn(spec)
    defs = table_defs(spec)
    for sid, t in defs.items():
        layout_table(t)
        p = TAB_DIR / f"{sid}_table.json"
        with open(p, "w", encoding="utf-8") as f:
            json.dump(t, f, ensure_ascii=False, indent=1)
        MANIFEST["tables"][sid] = dict(file=rel(p), format=t["kind"], geometry_in=t["geometry_in"],
                                       height_in=t["height_in"], issues=t["issues"], changes_from_spec=t["changes_from_spec"])
        st = "OK" if not t["issues"] else f"ISSUES {t['issues']}"
        print(f"[table] {sid:4s} h {t['height_in']:.2f} in  {st}")
    if not only or any(s in only for s in CHARTS):
        n_ok, n_bad = values_check(spec)
        MANIFEST["values_check"] = dict(file=rel(OUT / "values_check.txt"), ok=n_ok, mismatch=n_bad)
    MANIFEST["meta"] = dict(
        generated=datetime.now().strftime("%Y-%m-%d %H:%M"), script=rel(Path(__file__)),
        spec=rel(SPEC), chart_rule="지침 5.6 경로 B, 배치 크기 그대로 300 dpi, Pretendard Medium·SemiBold, 축 18 pt, 눈금·직접 라벨 16 pt",
        format_choice="차트 14개 모두 경로 B PNG(한글 포함, 지침 5.6: 경로 A EMF 는 한글이 Noto Sans CJK 로 대체되어 쓰지 않음). "
                      "표 15개는 python-pptx 원어 헤어라인 표(편집 가능). python-pptx 원어 차트는 쓰지 않음(Pretendard 축·눈금 글자와 "
                      "방법 색·선 모양·두 CI 막대를 원어 차트 객체로 재현할 수 없음)",
        table_rule="지침 3.2, 원형 A7: 원어 헤어라인 표, 18 pt, 채움 0, 세로선 0",
        not_made=dict(S28="오른쪽 '라벨 수 계급별 면적 비율'은 [DECISION](지침 8절 13). 대체안(Fig 1a 슬라이드판 + 굵은 리드 줄)이 기본이라 차트를 만들지 않았다",
                      S27_col2_col3="결론 쪽 열 2(알래스카 보정 지도, Fig 7b 판)·열 3(캐나다 선정 지점 지도, Fig 5 판)은 v3 그림 범위. 열 1 축소판만 S27_workflow_small.png",
                      S16_S18="왼쪽 SI 지도 슬라이드판(대상 중심 좌표 [미확인])과 오른쪽 Fig 4b·4c 는 v3 그림 범위",
                      v3_panels="S01, S02 보조, S05, S07, S11–S15, S17, S19, S21–S23, AP1 의 근거는 v3 논문 그림 패널의 슬라이드판(medium='slide')으로 이 스크립트 범위 밖"))
    if not a.no_preview:
        pptx_path, pages = preview(spec, defs, set(MANIFEST["charts"]))
        pdf = export_pdf_no_autospace(pptx_path, PREV_DIR)
        for old_png in PREV_DIR.glob("page-*.png"):
            old_png.unlink()
        subprocess.run(["pdftoppm", "-r", "110", "-png", str(pdf), str(PREV_DIR / "page")], check=True)
        MANIFEST["preview"] = dict(pptx=rel(pptx_path), pdf=rel(pdf), pages=pages,
                                   png=sorted(rel(p) for p in PREV_DIR.glob("page-*.png")))
        print(f"[preview] {rel(pdf)} ({len(pages)} 쪽)")
    with open(OUT / "MANIFEST.json", "w", encoding="utf-8") as f:
        json.dump(MANIFEST, f, ensure_ascii=False, indent=1, default=str)


if __name__ == "__main__":
    main()
