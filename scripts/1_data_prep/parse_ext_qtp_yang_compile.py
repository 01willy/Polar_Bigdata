"""qtp_yang_compile: Yang & Qiu 2026(Zenodo 18150789)의 티베트 고원 ALT 현장 관측 집계표를 표준 점 자료로 만든다.

근거
----
계획서 `docs/EXPERIMENT_PLAN_LG_2026-09-29.md` 6B 절(LGD, 개정 10), 형식 정의
`data/processed/ext_labels/_schema.json`(버전 1.0), 자료원 계획 `data/processed/ext_labels/_sources_plan.json`.
이 자료원은 보조 라벨이다(방법 미확인). 확인적 풀에 넣지 않는다.

입력(data/raw/qtp_yang_compile/)
--------------------------------
in-situ measurements of ALT and MAGT.xlsx   Zenodo 18150789. 시트 ALT(152지점), 시트 MAGT(쓰지 않는다)
41612_2024_866_MOESM1_ESM.pdf               Chang 등 2024 npj Clim Atmos Sci 보충 자료. 표 S1(6-7쪽)
41612_2024_866_MOESM1_ESM_p6_bbox.html      위 PDF 6쪽의 pdftotext -bbox-layout 산출(없으면 만든다)
41612_2024_866_MOESM1_ESM_p7_bbox.html      위 PDF 7쪽의 산출. poppler 0.62 는 여러 쪽을 한 번에 주면 첫 쪽만
                                            내므로 쪽마다 따로 만든다
읽기 전용 대조 자료: fidelity_base_v3.csv, CALM PANGAEA 972777, 같은 폴더의 다른 티베트 점 자료
(qtp_du_gpr, qtp_fu_temp, qtp_liux_compile). 점 자료가 아직 없으면 원자료 표를 읽는다.

규칙
----
1. 시트 ALT 의 152지점을 모두 넣는다. 지우는 행은 없다. 단위는 머리글대로 cm 다.
2. method = unknown, label_def = unknown. 출처가 CALM 인 6지점은 method = borehole_temp,
   label_def = temp_derived 다(CALM 이벤트 머리말의 Method 가 Ground temperature measurements).
3. 연도는 Chang 표 S1 과 같은 80지점만 채운다. 관측 기간이 한 해면 year 도 채우고
   value_kind = annual_value, 여러 해면 year 는 비우고 value_kind = multiyear_mean 이다.
   나머지 72지점은 연도가 없다. 연도 규칙(1990년 이후)을 확인할 수 없으므로 qc_flag 에 date 를 붙인다.
4. 표 S1 의 Source 열은 병합 칸이다. 칸의 범위는 PDF 의 글자 위치로 복원한다(출처 글자 묶음의
   세로 중심과 행 묶음의 세로 중심이 같다는 조건). 복원 결과가 EXPECTED_GROUPS 와 다르면 중단한다.
5. site_id 는 시트의 ID 로 만든다(<src_id>_<4자리>). 지점 이름은 중복이 있어 쓰지 않는다(KKXL 4건).
6. 좌표는 원자료 값을 그대로 쓴다. 의심스러운 좌표도 고치지 않고 qc_flag 에 coord 를 붙인다.
7. dup_of: 다른 자료원의 지점이 체비쇼프 0.01° 이내에 있으면 적는다. 여러 개면 v3 F4_direct 셀,
   같은 지점의 근거가 있는 대상(값이 1 cm 이내로 같거나 이름이 같다), 자료원 우선순위(v3 CALM 지온,
   qtp_du_gpr, qtp_fu_temp, qtp_liux_compile), 거리 순으로 하나를 고른다. 나머지는 notes 와 메타에
   적는다. 근거가 없는 대상(dup_kind=near_only)은 거리 기준일 뿐이고 같은 지점이 아닐 수 있다.
   연 값이 4개 이상인 대상(연별 계열)은 이 표의 관측 기간 평균과 비교한다. 기간이 없으면 값 비교를
   하지 않는다(여러 해 가운데 한 해와 우연히 같은 경우를 막는다).
   표 안에서 0.01° 이내이고 값이 같은 뒤쪽 기록은 외부 대상이 없을 때 앞 기록을 dup_of 로 적는다.
8. 우측 절단: 시트에 '>' 표기는 없다. CALM 6지점은 CALM 기록과 대조해, 기록이 전부 '>' 인 지점(CN6)만
   right_censored = 1 로 둔다.
9. 범위 QC 는 0 < alt_cm <= 600 이다.

산출
----
data/processed/ext_labels/qtp_yang_compile_points.csv
data/processed/ext_labels/qtp_yang_compile_meta.json

실행(ROOT): OMP_NUM_THREADS=1 python scripts/1_data_prep/parse_ext_qtp_yang_compile.py
"""
from __future__ import annotations

import csv
import hashlib
import html
import json
import math
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

SRC_ID = "qtp_yang_compile"
SRC_NAME = ("Yang & Qiu 2026, 티베트 고원 ALT·MAGT 현장 관측 집계표(Zenodo 18150789)와 "
            "Chang 등 2024 보충 표 S1")
URL = "https://zenodo.org/records/18150789"
DOI = "10.5281/zenodo.18150789"
XLSX_URL = ("https://zenodo.org/api/records/18150789/files/"
            "in-situ%20measurements%20of%20ALT%20and%20MAGT.xlsx/content")
XLSX_MD5 = "1401277607ec2a93bf7e56cf611aff6f"      # Zenodo API 의 파일 checksum(2026-09-29 조회)
CHANG_DOI = "10.1038/s41612-024-00866-0"
CHANG_SI_URL = ("https://media.springernature.com/original/springer-static/esm/"
                "art%3A10.1038%2Fs41612-024-00866-0/MediaObjects/41612_2024_866_MOESM1_ESM.pdf")
ACCESSED = "2026-09-29"

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw" / SRC_ID
XLSX = RAW / "in-situ measurements of ALT and MAGT.xlsx"
SI_PDF = RAW / "41612_2024_866_MOESM1_ESM.pdf"
SI_PAGES = (6, 7)
SI_BBOX = {p: RAW / f"41612_2024_866_MOESM1_ESM_p{p}_bbox.html" for p in SI_PAGES}
ZENODO_JSON = RAW / "zenodo_record_18150789.json"
OUT_DIR = ROOT / "data" / "processed" / "ext_labels"
POINTS = OUT_DIR / f"{SRC_ID}_points.csv"
META = OUT_DIR / f"{SRC_ID}_meta.json"
SCHEMA = OUT_DIR / "_schema.json"
V3 = ROOT / "data" / "processed" / "fidelity_base_v3.csv"
CALM_TAB = ROOT / "data" / "raw" / "calm" / "PANGAEA_972777_CALM_ALT_NH.tab"
EXPLORE = ROOT / "data" / "raw" / "qtp_alt_open_2026-09-29"
FU_XLSX = [ROOT / "data" / "raw" / "qtp_fu_temp" / "Permafrost_Profile_QTP_2001-2020.xlsx",
           EXPLORE / "figshare29206613" / "Permafrost_Profile_QTP_2001-2020.xlsx"]
LIUX_XLSX = [ROOT / "data" / "raw" / "qtp_liux_compile" / "ds01.xlsx",
             EXPLORE / "figshare21444879" / "ds01.xlsx"]
DU_CSV = [EXPLORE / "zenodo21999366" / "du_pit_gpr_points.csv"]

LICENSE_YANG = "CC BY 4.0 (Zenodo 10.5281/zenodo.18150789)"
LICENSE_CHANG_PART = ("year, year_min, year_max, orig_source from Chang et al. 2024 Table S1 "
                      "(article licence CC BY-NC-ND 4.0)")
CITATION = ("Yang, G., Qiu, H. (2026). Permafrost thermal dynamics and thawed soil organic carbon on the "
            "Tibetan Plateau [Data set]. Zenodo. https://doi.org/10.5281/zenodo.18150789")
CITATION_CHANG = ("Chang, T., Yi, Y., Jiang, H., Li, R., Lu, P., Liu, L., Wang, L., Zhao, L., Zwieback, S., "
                  "Zhao, J. (2024). Unraveling the non-linear relationship between seasonal deformation and "
                  "permafrost active layer thickness. npj Climate and Atmospheric Science 7, 308. "
                  "https://doi.org/10.1038/s41612-024-00866-0 (Supplementary Table S1)")

COLUMNS = ["src_id", "site_id", "site_name", "lat", "lon", "year", "month", "alt_cm", "method", "label_def",
           "n_obs", "country", "macro", "citation", "license",
           "subunit", "date", "eos_basis", "value_kind", "year_min", "year_max", "alt_sd_cm",
           "coord_prec_deg", "disturbed", "disturb_type", "right_censored", "orig_source", "dup_of",
           "qc_flag", "notes"]
SHEET_HEADER = ["ID", "Site name", "Lat", "Lon", "Elevation\n(m)", "ALT\n(cm)", "Source"]
N_EXPECTED = 152
N_CHANG = 80
CHANG_LABEL_IN_SHEET = "Chang et al. 2025"      # 시트의 표기. 논문의 게재 연도는 2024 다
CALM_LABEL_IN_SHEET = "CALM-ALT"

# 표 S1 의 Source 병합 칸(쪽 순서, 위에서 아래). 글자 위치로 복원한 결과와 대조하는 기준이다.
EXPECTED_GROUPS = [
    ("Bai et al. (2023)", 1), ("Zhao & Sheng (2019)", 22), ("Chen (2018)", 1), ("Qin et al. (2017)", 3),
    ("China Railway First Survey and Design Institute Group", 7), ("Wu et al. (2008)", 2),
    ("Wu et al. (2012)", 3),
    ("Wu et al. (2012)", 15), ("Wu et al. (2015)", 3), ("Wu et al. (2017)", 3), ("Xie et al. (2012)", 3),
    ("Xie et al. (2015)", 6), ("Niu (2021)", 3), ("Zhao et al. (2021)", 8),
]
# 표 S1 의 참고 문헌 번호 3-14(보충 자료 8-9쪽). 메타에만 적는다.
CHANG_REFS = {
    "Bai et al. (2023)": "Bai, Y. et al. Methods in Ecology and Evolution 14, 1732-1746 (2023)",
    "Zhao & Sheng (2019)": "Zhao, L. & Sheng, Y. Permafrost and Its Change on the Qinghai-Tibet Plateau "
                           "(Science Press, 2019)",
    "Chen (2018)": "Chen, J. doi:10.11888/Geocry.tpdc.270459 (2018)",
    "Qin et al. (2017)": "Qin, Y. et al. J. Geophys. Res. Atmos. 122, 11,604-11,620 (2017)",
    "China Railway First Survey and Design Institute Group": "참고 문헌 번호 없음",
    "Wu et al. (2008)": "Wu, Q. & Zhang, T. J. Geophys. Res. Atmos. 113 (2008)",
    "Wu et al. (2012)": "Wu, Q., Zhang, T. & Liu, Y. The Cryosphere 6, 607-612 (2012)",
    "Wu et al. (2015)": "Wu, Q., Hou, Y., Yun, H. & Liu, Y. Global and Planetary Change 124, 149-155 (2015)",
    "Wu et al. (2017)": "Wu, Q., Yu, W. & Jin, H. Sci Rep 7, 1544 (2017)",
    "Xie et al. (2012)": "Xie, C., Zhao, L., Wu, T. & Dong, X. J. Mt. Sci. 9, 483-491 (2012)",
    "Xie et al. (2015)": "Xie, C. et al. Arctic, Antarctic, and Alpine Research 47, 719-728 (2015)",
    "Niu (2021)": "Niu, F. doi:10.11888/Cryos.tpdc.271933 (2021)",
    "Zhao et al. (2021)": "Zhao, L. et al. Earth System Science Data 13, 4207-4218 (2021)",
}
# 철도 노반 관측 자료에서 온 출처. 시설 부지 여부는 확인하지 못했다(notes 에만 적는다).
RAILWAY_SOURCES = {"Niu (2021)", "China Railway First Survey and Design Institute Group"}
# 같은 관측망(Wu Q. 의 청장 철도·도로 시추공)이라 qtp_fu_temp 와 이름이 같으면 같은 지점으로 본다.
WU_NETWORK = {"Wu et al. (2008)", "Wu et al. (2012)", "Wu et al. (2015)", "Wu et al. 2012"}

# Tibet 의 지리 범위(6B.4)와 하위 단위 규칙. parse_ext_qtp_du_gpr.py 와 같은 상자와 이름이다.
TIBET_BOX = dict(lat_min=26.0, lat_max=40.0, lon_min=73.0, lon_max=105.0)
SUBUNIT_RULE = [
    ("Qilian", "치롄", dict(lat_min=37.0, lat_max=40.0, lon_min=96.0, lon_max=103.0, lon_max_incl=True)),
    ("YellowSrc_East", "황하원·동부", dict(lat_min=33.5, lat_max=36.5, lon_min=95.0, lon_max=101.0,
                                     lon_max_incl=True)),
    ("QTEC", "회랑", dict(lat_min=31.0, lat_max=36.2, lon_min=90.0, lon_max=95.0, lon_max_incl=False)),
    ("West_QTP", "서부", dict(lat_min=30.0, lat_max=38.0, lon_min=77.0, lon_max=90.0, lon_max_incl=False)),
]
DUP_DEG = 0.01            # 체비쇼프 거리 기준(6B.3)
EPS = 1e-9
SAME_VALUE_CM = 1.0       # 값 일치 판단 허용차
SERIES_MIN_N = 4          # 대상의 연 값이 이 수 이상이면 연별 계열로 보고 기간 평균과 비교한다
NAME_COORD_DEG = 0.02     # 이름이 같은 지점의 좌표 차이가 이 값을 넘으면 coord 표시
CELL = 0.009              # 약 1 km 셀(6B.3)
RANK = {"v3_f4_direct": 0, "v3_other": 1, "qtp_du_gpr": 2, "qtp_fu_temp": 3, "qtp_liux_compile": 4}
CENTER_TOL_PT = 3.0       # 병합 칸 복원의 세로 중심 허용차(pt). 행 간격은 14-15.5 pt 다


# ---------------------------------------------------------------- 공통
def file_hash(path: Path, algo: str) -> str:
    h = hashlib.new(algo)
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def subunit_of(lat: float, lon: float) -> str:
    for name, _, b in SUBUNIT_RULE:
        lon_ok = (b["lon_min"] <= lon <= b["lon_max"]) if b["lon_max_incl"] else (b["lon_min"] <= lon < b["lon_max"])
        if b["lat_min"] <= lat <= b["lat_max"] and lon_ok:
            return name
    return "QTP_other"


def in_tibet(lat: float, lon: float) -> bool:
    b = TIBET_BOX
    return b["lat_min"] <= lat <= b["lat_max"] and b["lon_min"] <= lon <= b["lon_max"]


def n_decimals(x: float) -> int:
    s = repr(float(x))
    if "e" in s or "E" in s:
        return 0
    return len(s.split(".")[1].rstrip("0")) if "." in s else 0


def coord_prec(lat: float, lon: float) -> float:
    """좌표 정밀도. 소수 4자리 이상이고 위도와 경도가 모두 분 단위의 배수면 0.0167, 아니면 소수 자리 수."""
    d = max(n_decimals(lat), n_decimals(lon))
    if d >= 4 and all(abs(v * 60 - round(v * 60)) < 0.01 for v in (lat, lon)):
        return 0.0167
    return float(10 ** (-max(d, 2)))


def num_str(x) -> str:
    """원자료 정밀도를 유지한 숫자 표기. 정수 값은 소수점 없이 적는다."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    f = float(x)
    return str(int(f)) if f == int(f) else repr(f)


def norm_name(s: str) -> str:
    return re.sub(r"[\s \-–_]", "", str(s)).upper()


def cheb(lat: float, lon: float, rlat, rlon):
    return np.maximum(np.abs(np.asarray(rlat, float) - lat), np.abs(np.asarray(rlon, float) - lon))


def first_existing(paths) -> Path | None:
    for p in paths:
        if Path(p).exists():
            return Path(p)
    return None


# ---------------------------------------------------------------- Yang 시트
def read_yang() -> tuple[pd.DataFrame, dict]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    info = dict(sheets={ws.title: dict(dimensions=ws.dimensions) for ws in wb.worksheets})
    ws = wb["ALT"]
    header = [c.value for c in ws[1]][:7]
    if header != SHEET_HEADER:
        raise SystemExit(f"시트 ALT 의 머리글이 예상과 다르다: {header}")
    # Source 열은 병합 칸이다. 병합 범위의 첫 칸 값을 범위 안의 모든 행에 준다
    src_of_row = {}
    merged = []
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 7 and rng.max_col == 7:
            v = ws.cell(rng.min_row, 7).value
            merged.append(dict(range=str(rng), n_rows=rng.max_row - rng.min_row + 1,
                               label=str(v).strip().split("\n")[0] if v else None))
            for r in range(rng.min_row, rng.max_row + 1):
                src_of_row[r] = v
        else:
            raise SystemExit(f"Source 열 밖의 병합 칸이 있다: {rng}")
    rows, n_blank, censored_marks = [], 0, 0
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        if all(v is None for v in vals[:6]):
            n_blank += 1
            continue
        src = src_of_row.get(r, vals[6])
        if src is None:
            raise SystemExit(f"시트 ALT {r}행의 Source 가 비어 있다")
        lines = [ln.strip() for ln in str(src).split("\n") if ln.strip()]
        label = lines[0]
        url = next((ln for ln in lines if ln.startswith("http")), "")
        alt_raw, rc = vals[5], 0
        if isinstance(alt_raw, str):
            s = alt_raw.strip()
            if s.startswith(">"):
                rc, censored_marks = 1, censored_marks + 1
                s = s[1:]
            alt_raw = float(s)          # 숫자가 아니면 여기서 중단된다
        for k, v in (("Lat", vals[2]), ("Lon", vals[3]), ("ALT", alt_raw)):
            if not isinstance(v, (int, float)):
                raise SystemExit(f"시트 ALT {r}행의 {k} 가 숫자가 아니다: {v!r}")
        name_raw = str(vals[1])
        rows.append(dict(ID=int(vals[0]), excel_row=r, name_raw=name_raw,
                         name=name_raw.replace(" ", " ").strip(), lat=float(vals[2]), lon=float(vals[3]),
                         elev_m=vals[4], alt_cm=float(alt_raw), rc_mark=rc, src_label=label, src_url=url))
    df = pd.DataFrame(rows)
    if len(df) != N_EXPECTED or df.ID.tolist() != list(range(1, N_EXPECTED + 1)):
        raise SystemExit(f"시트 ALT 의 행 수 또는 ID 가 예상과 다르다: {len(df)}행")
    wsm = wb["MAGT"]
    n_magt = sum(1 for r in range(2, wsm.max_row + 1) if wsm.cell(r, 1).value is not None)
    info.update(n_rows_alt=len(df), n_blank_rows_alt=n_blank, n_rows_magt_not_used=n_magt,
                merged_source_cells=sorted(merged, key=lambda m: int(re.findall(r"\d+", m["range"])[0])),
                n_gt_marks=censored_marks,
                n_name_with_nbsp=int(df.name_raw.str.contains(" ").sum()))
    return df, info


# ---------------------------------------------------------------- Chang 표 S1
WORD_RE = re.compile(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>')


def ensure_bbox() -> str:
    """쪽마다 따로 만든다(poppler 0.62 의 -bbox-layout 은 여러 쪽을 주면 첫 쪽의 글자만 낸다)."""
    state = "기존 파일"
    for p, f in SI_BBOX.items():
        if not f.exists():
            subprocess.run(["pdftotext", "-bbox-layout", "-f", str(p), "-l", str(p), str(SI_PDF), str(f)],
                           check=True)
            state = "pdftotext 로 생성"
    return state


def read_chang() -> tuple[pd.DataFrame, dict]:
    pages = []
    for p in SI_PAGES:
        part = SI_BBOX[p].read_text(encoding="utf-8").split("<page ")[1:]
        if len(part) != 1:
            raise SystemExit(f"bbox 파일의 쪽 수가 예상과 다르다(쪽 {p}): {len(part)}")
        pages.append(part[0])
    out, groups, max_dev = [], [], 0.0
    for pi, page in enumerate(pages):
        words = [dict(x0=float(a), y0=float(b), x1=float(c), y1=float(d), t=html.unescape(e))
                 for a, b, c, d, e in WORD_RE.findall(page)]
        for w in words:
            w["xc"], w["yc"], w["h"] = (w["x0"] + w["x1"]) / 2, (w["y0"] + w["y1"]) / 2, w["y1"] - w["y0"]
        title = [w for w in words if w["t"] == "Table"]
        if not title:
            raise SystemExit("표 S1 의 제목을 찾지 못했다")
        y_head0 = title[0]["y1"]
        head = {w["t"]: w for w in words if y_head0 < w["y0"] < y_head0 + 12
                and w["t"] in ("Area", "Site", "name", "Period", "Vegetation", "type", "Source",
                               "Latitude", "Longitude", "Elevation", "ALT")}
        if len(head) != 11:
            raise SystemExit(f"표 S1 의 머리글을 찾지 못했다: {sorted(head)}")
        centers = dict(area=head["Area"]["xc"], site=(head["Site"]["x0"] + head["name"]["x1"]) / 2,
                       lat=head["Latitude"]["xc"], lon=head["Longitude"]["xc"], period=head["Period"]["xc"],
                       elev=head["Elevation"]["xc"], alt=head["ALT"]["xc"],
                       veg=(head["Vegetation"]["x0"] + head["type"]["x1"]) / 2, source=head["Source"]["xc"])
        y_body0 = max(w["y1"] for w in words if w["t"] in ("(°)", "(m)", "(cm)") and w["y0"] < y_head0 + 30)
        body = [w for w in words if w["y0"] > y_body0 - 1 and w["x0"] > 60]
        for w in body:
            w["col"] = min(centers, key=lambda k: abs(centers[k] - w["xc"]))
        # 행: 위도 열의 숫자
        lat_words = sorted([w for w in body if w["col"] == "lat" and re.fullmatch(r"\d{2}\.\d+", w["t"])],
                           key=lambda w: w["yc"])
        rows = []
        for lw in lat_words:
            same = [w for w in body if abs(w["yc"] - lw["yc"]) <= 3.5]
            cell = {c: " ".join(w["t"] for w in sorted([w for w in same if w["col"] == c], key=lambda w: w["x0"]))
                    for c in ("site", "lon", "period", "elev", "alt")}
            if not (re.fullmatch(r"\d{2,3}\.\d+", cell["lon"]) and re.fullmatch(r"\d{4}(-\d{4})?", cell["period"])
                    and re.fullmatch(r"\d{3,4}", cell["elev"]) and re.fullmatch(r"\d{2,3}", cell["alt"])
                    and cell["site"]):
                raise SystemExit(f"표 S1 의 행을 읽지 못했다(쪽 {SI_PAGES[pi]}): {lw['t']} {cell}")
            rows.append(dict(page=SI_PAGES[pi], name=cell["site"], lat=float(lw["t"]), lon=float(cell["lon"]),
                             lat_text=lw["t"], lon_text=cell["lon"], period=cell["period"],
                             elev_m=int(cell["elev"]), alt_cm=float(cell["alt"]), y0=lw["y0"], y1=lw["y1"]))
        # 출처 글자 묶음. 위 첨자(참고 문헌 번호)는 뺀다
        sw = [w for w in body if w["col"] == "source" and not (w["h"] < 6.0 and w["t"].isdigit())]
        lines = {}
        for w in sorted(sw, key=lambda w: (w["y0"], w["x0"])):
            key = next((k for k in lines if abs(k - w["y0"]) < 1.0), w["y0"])
            lines.setdefault(key, []).append(w)
        blocks = []
        for y in sorted(lines):
            ws_ = lines[y]
            item = dict(y0=min(w["y0"] for w in ws_), y1=max(w["y1"] for w in ws_),
                        text=" ".join(w["t"] for w in sorted(ws_, key=lambda w: w["x0"])))
            if blocks and item["y0"] - blocks[-1]["last_y0"] <= 11.0:
                b = blocks[-1]
                b["text"] += " " + item["text"]
                b["y1"], b["last_y0"] = max(b["y1"], item["y1"]), item["y0"]
            else:
                blocks.append(dict(y0=item["y0"], y1=item["y1"], last_y0=item["y0"], text=item["text"]))
        # 병합 칸 복원: 출처 묶음의 세로 중심과 행 묶음의 세로 중심이 가장 가까운 끝 행을 고른다
        i = 0
        for b in blocks:
            if i >= len(rows):
                raise SystemExit("표 S1 의 출처 묶음이 행보다 많다")
            bc = (b["y0"] + b["y1"]) / 2
            devs = [abs((rows[i]["y0"] + rows[j]["y1"]) / 2 - bc) for j in range(i, len(rows))]
            k = int(np.argmin(devs))
            if devs[k] > CENTER_TOL_PT:
                raise SystemExit(f"표 S1 의 출처 칸을 복원하지 못했다: {b['text']} (차이 {devs[k]:.1f} pt)")
            max_dev = max(max_dev, devs[k])
            label = re.sub(r"\s+", " ", b["text"]).strip()
            for j in range(i, i + k + 1):
                rows[j]["source"] = label
                rows[j]["source_center_dev_pt"] = round(devs[k], 2)
            groups.append((label, k + 1))
            i += k + 1
        if i != len(rows):
            raise SystemExit(f"표 S1 에 출처가 없는 행이 있다(쪽 {SI_PAGES[pi]}): {len(rows) - i}행")
        out += rows
    if groups != EXPECTED_GROUPS:
        raise SystemExit(f"표 S1 의 출처 묶음이 기준과 다르다: {groups}")
    df = pd.DataFrame(out)
    if len(df) != N_CHANG:
        raise SystemExit(f"표 S1 의 티베트 행 수가 예상과 다르다: {len(df)}")
    info = dict(n_rows=len(df), pages=list(SI_PAGES), groups=[dict(source=g, n_rows=n) for g, n in groups],
                max_center_dev_pt=round(max_dev, 2), center_tol_pt=CENTER_TOL_PT,
                references={k: v for k, v in CHANG_REFS.items()})
    return df, info


def match_chang(y: pd.DataFrame, ch: pd.DataFrame) -> dict:
    """시트의 앞 80행과 표 S1 의 80행을 순서대로 대조한다."""
    head = y.iloc[:N_CHANG]
    if set(head.src_label) != {CHANG_LABEL_IN_SHEET} or (y.iloc[N_CHANG:].src_label == CHANG_LABEL_IN_SHEET).any():
        raise SystemExit("시트에서 출처가 Chang 인 행이 앞 80행이 아니다")
    mism = []
    for (_, a), (_, b) in zip(head.iterrows(), ch.iterrows()):
        diff = [k for k, u, v in (("name", norm_name(a["name"]), norm_name(b["name"])),
                                  ("lat", a.lat, b.lat), ("lon", a.lon, b.lon),
                                  ("elev_m", float(a.elev_m), float(b.elev_m)), ("alt_cm", a.alt_cm, b.alt_cm))
                if (u != v if isinstance(u, str) else abs(u - v) > 1e-9)]
        if diff:
            mism.append(dict(ID=int(a.ID), fields=diff, sheet=[a["name"], a.lat, a.lon, a.elev_m, a.alt_cm],
                             table_s1=[b["name"], b.lat, b.lon, b.elev_m, b.alt_cm]))
    if mism:
        raise SystemExit(f"시트와 표 S1 이 다른 행이 있다: {mism}")
    return dict(n_compared=N_CHANG, n_identical=N_CHANG - len(mism), fields=["name", "lat", "lon", "elev_m", "alt_cm"],
                order="시트 ID 1-80 과 표 S1 의 행 순서가 같다")


# ---------------------------------------------------------------- CALM 대조
def calm_check(y: pd.DataFrame) -> dict:
    if not CALM_TAB.exists():
        return dict(available=False, note="CALM PANGAEA 972777 파일이 없어 대조하지 못했다")
    text = CALM_TAB.read_text(encoding="utf-8")
    method = {m.group(1): m.group(2).strip()
              for m in re.finditer(r"^\tCALM_(CN\d+) \(.*?COMMENT: (.*?)$", text, flags=re.M)}
    body = text.split("*/\n", 1)[1].splitlines()
    cols = body[0].split("\t")
    ie, iy, ia, ila, ilo = (cols.index(c) for c in ("Event", "Date/Time", "ALD [cm]", "Latitude", "Longitude"))
    rec = {}
    for ln in body[1:]:
        p = ln.split("\t")
        if p[ie].startswith("CALM_CN") and p[ia].strip():
            rec.setdefault(p[ie][5:], []).append((int(p[iy][:4]), p[ia].strip(), float(p[ila]), float(p[ilo])))
    out = {}
    for _, r in y[y.src_label == CALM_LABEL_IN_SHEET].iterrows():
        code = norm_name(r["name"])
        if code not in rec:
            out[code] = dict(found=False)
            continue
        vals = rec[code]
        unc = [(yr, float(v)) for yr, v, _, _ in vals if not v.startswith(">")]
        cen = [(yr, float(v[1:])) for yr, v, _, _ in vals if v.startswith(">")]
        mean_unc = float(np.mean([v for _, v in unc])) if unc else None
        all_cens = len(unc) == 0 and len(cen) > 0
        bound = max(v for _, v in cen) if cen else None
        out[code] = dict(
            found=True, site_id=f"{SRC_ID}_{int(r.ID):04d}", yang_alt_cm=r.alt_cm,
            calm_method=method.get(code, ""),
            calm_lat=vals[0][2], calm_lon=vals[0][3],
            cheb_deg=round(float(max(abs(vals[0][2] - r.lat), abs(vals[0][3] - r.lon))), 5),
            years_all=[min(v[0] for v in vals), max(v[0] for v in vals)],
            n_uncensored=len(unc), n_censored=len(cen),
            years_uncensored=[min(v[0] for v in unc), max(v[0] for v in unc)] if unc else None,
            mean_uncensored_cm=round(mean_unc, 2) if unc else None,
            censored_bound_cm=bound, all_censored=all_cens,
            value_matches=(abs(mean_unc - r.alt_cm) <= SAME_VALUE_CM) if unc else
                          (bound is not None and abs(bound - r.alt_cm) <= SAME_VALUE_CM))
    return dict(available=True, file=str(CALM_TAB.relative_to(ROOT)), sites=out,
                rule="시트 값이 CALM 기록의 '>' 없는 연도 평균과 1 cm 이내면 일치로 본다. "
                     "기록이 전부 '>' 이고 시트 값이 그 하한과 같으면 right_censored = 1 이다")


# ---------------------------------------------------------------- 대조 자료(dup_of)
def load_refs() -> tuple[pd.DataFrame, dict]:
    """다른 자료원의 지점. 열: ref_src, ref_id, name, lat, lon, alt_cm, year, rank, group."""
    parts, mode = [], {}
    # v3(읽기 전용)
    v3 = pd.read_csv(V3, usecols=["loc_id", "lat", "lon", "region", "alt_cm", "source_id"], low_memory=False)
    v3 = v3[v3.lat.between(TIBET_BOX["lat_min"] - 1, TIBET_BOX["lat_max"] + 1)
            & v3.lon.between(TIBET_BOX["lon_min"] - 1, TIBET_BOX["lon_max"] + 1)]
    parts.append(pd.DataFrame(dict(
        ref_src="fidelity_base_v3", ref_id=v3.loc_id.astype(int).astype(str), name=v3.region + "/" + v3.source_id,
        lat=v3.lat, lon=v3.lon, alt_cm=v3.alt_cm, year=np.nan, year_end=np.nan,
        group=np.where(v3.source_id == "F4_direct", "v3_f4_direct", "v3_other"))))
    mode["fidelity_base_v3"] = dict(mode="table", file=str(V3.relative_to(ROOT)), n_rows_in_box=int(len(v3)))
    # 다른 티베트 점 자료. 점 자료가 있으면 그것을 쓰고 없으면 원자료 표를 읽는다
    for src in ("qtp_du_gpr", "qtp_fu_temp", "qtp_liux_compile"):
        p = OUT_DIR / f"{src}_points.csv"
        if p.exists():
            d = pd.read_csv(p, dtype=str, keep_default_na=False)
            d = d[d.alt_cm != ""]
            y0 = pd.to_numeric(d.year.where(d.year != "", d.year_min), errors="coerce")
            y1 = pd.to_numeric(d.year.where(d.year != "", d.year_max), errors="coerce")
            parts.append(pd.DataFrame(dict(ref_src=src, ref_id=d.site_id, name=d.site_name,
                                           lat=d.lat.astype(float), lon=d.lon.astype(float),
                                           alt_cm=d.alt_cm.astype(float), year=y0, year_end=y1, group=src)))
            mode[src] = dict(mode="points", file=str(p.relative_to(ROOT)), n_rows=int(len(d)),
                             sha256=file_hash(p, "sha256"))
            continue
        raw = raw_ref(src)
        if raw is None:
            mode[src] = dict(mode="missing", note="점 자료와 원자료 표가 모두 없어 대조하지 못했다")
            continue
        d, f = raw
        parts.append(d.assign(ref_src=src, group=src))
        mode[src] = dict(mode="raw_table", file=str(f.relative_to(ROOT)), n_rows=int(len(d)),
                         note="점 자료가 없어 원자료 표를 읽었다. dup_of 의 site_id 는 원자료 코드다")
    ref = pd.concat(parts, ignore_index=True)
    ref["rank"] = ref.group.map(RANK)
    return ref, mode


def raw_ref(src: str):
    if src == "qtp_fu_temp":
        f = first_existing(FU_XLSX)
        if f is None:
            return None
        site, alt = read_fu(f)
        m = alt.melt(id_vars="Year", var_name="name", value_name="alt_m").dropna(subset=["alt_m"])
        m = m.merge(site, on="name", how="inner")
        return pd.DataFrame(dict(ref_id=m.name, name=m.name, lat=m.lat, lon=m.lon,
                                 alt_cm=m.alt_m.astype(float) * 100.0, year=m.Year.astype(float),
                                 year_end=m.Year.astype(float))), f
    if src == "qtp_liux_compile":
        f = first_existing(LIUX_XLSX)
        if f is None:
            return None
        d = read_liux(f)
        yrs = d.period.astype(str).str.findall(r"\d{4}")
        return pd.DataFrame(dict(ref_id=d.code.astype(str), name=d.name.astype(str), lat=d.lat, lon=d.lon,
                                 alt_cm=d.alt_m.astype(float) * 100.0,
                                 year=[float(v[0]) if v else np.nan for v in yrs],
                                 year_end=[float(v[-1]) if v else np.nan for v in yrs])), f
    if src == "qtp_du_gpr":
        f = first_existing(DU_CSV)
        if f is None:
            return None
        d = pd.read_csv(f)
        d = d[d.alt_m.notna()]
        return pd.DataFrame(dict(ref_id=d.site.astype(str), name=d.site.astype(str), lat=d.lat, lon=d.lon,
                                 alt_cm=d.alt_m * 100.0, year=d.year.astype(float),
                                 year_end=d.year.astype(float))), f
    return None


def read_fu(f: Path):
    x = pd.ExcelFile(f)
    site = x.parse("SITE").iloc[:, :4]
    site.columns = ["name", "region", "lat", "lon"]
    site = site.dropna(subset=["lat", "lon"])
    site["name"] = site.name.astype(str).str.strip()
    alt = x.parse("ALT")
    alt.columns = [str(c).strip() for c in alt.columns]
    alt = alt[pd.to_numeric(alt.Year, errors="coerce").notna()]
    for c in alt.columns:
        alt[c] = pd.to_numeric(alt[c], errors="coerce")
    return site, alt


def read_liux(f: Path) -> pd.DataFrame:
    d = pd.read_excel(f, header=1)
    d.columns = ["ID", "ID_paper", "name", "code", "lat", "lon", "period", "alt_m", "src"]
    return d.dropna(subset=["lat", "lon"])


def find_dups(y: pd.DataFrame, ref: pd.DataFrame) -> dict:
    """행별 0.01° 이내 대상. 같은 (ref_src, ref_id) 는 하나로 묶는다."""
    res = {}
    rlat, rlon = ref.lat.values, ref.lon.values
    for _, r in y.iterrows():
        d = cheb(r.lat, r.lon, rlat, rlon)
        near = ref[d <= DUP_DEG + EPS].assign(dist=d[d <= DUP_DEG + EPS])
        has_period = pd.notna(r.p_min)
        cands = []
        for (src, rid), g in near.groupby(["ref_src", "ref_id"], sort=False):
            g = g[g.alt_cm.notna()]
            if g.empty:
                continue
            vals = [float(v) for v in g.alt_cm]
            y0 = [int(v) for v in g.year if pd.notna(v)]
            y1 = [int(v) for v in g.year_end if pd.notna(v)]
            series = len(vals) >= SERIES_MIN_N
            period_mean = None
            if series:
                sel = g[(g.year >= r.p_min) & (g.year_end <= r.p_max)] if has_period else g.iloc[:0]
                period_mean = round(float(sel.alt_cm.mean()), 1) if len(sel) else None
                same_value = period_mean is not None and abs(period_mean - r.alt_cm) <= SAME_VALUE_CM
            else:
                same_value = any(abs(v - r.alt_cm) <= SAME_VALUE_CM for v in vals)
            same_name = norm_name(r["name"]) in {norm_name(rid), norm_name(g.name.iloc[0])}
            overlap = None
            if has_period and y0 and y1:
                overlap = bool(min(y0) <= r.p_max and max(y1) >= r.p_min)
            kind = ("same_value_and_name" if same_value and same_name else "same_value" if same_value
                    else "same_name" if same_name else "near_only")
            cands.append(dict(
                ref=f"{src}:{rid}", ref_src=src, ref_name=str(g.name.iloc[0]), rank=int(g["rank"].iloc[0]),
                group=str(g.group.iloc[0]), cheb_deg=round(float(g.dist.min()), 5),
                ref_n_values=len(vals), ref_is_series=series,
                ref_alt_cm=[round(v, 1) for v in vals] if not series else
                           dict(min=round(min(vals), 1), mean=round(float(np.mean(vals)), 1), max=round(max(vals), 1)),
                ref_period_mean_cm=period_mean,
                ref_years=[min(y0), max(y1)] if y0 and y1 else None, years_overlap=overlap,
                same_value=bool(same_value), same_name=bool(same_name), kind=kind))
        # v3 F4_direct 가 먼저다(v4 의 제외 규칙). 다음은 같은 지점의 근거가 있는 대상, 우선순위, 거리
        cands.sort(key=lambda c: (c["group"] != "v3_f4_direct", c["kind"] == "near_only", c["rank"],
                                  c["cheb_deg"]))
        res[int(r.ID)] = cands
    return res


def internal_pairs(y: pd.DataFrame) -> dict:
    """표 안에서 0.01° 이내인 다른 기록. same 은 값도 같은 기록(같은 기록의 중복 수록으로 본다)."""
    out = {}
    for _, r in y.iterrows():
        d = cheb(r.lat, r.lon, y.lat.values, y.lon.values)
        near = y[(d <= DUP_DEG + EPS) & (y.ID != r.ID)]
        same = near[(near.alt_cm - r.alt_cm).abs() <= SAME_VALUE_CM]
        out[int(r.ID)] = dict(near=[int(v) for v in near.ID], same=[int(v) for v in same.ID])
    return out


def name_checks(y: pd.DataFrame) -> dict:
    """이름이 같은 지점의 좌표 대조. 원자료 표만 읽는다."""
    out = dict(fu=dict(available=False), liux=dict(available=False), flagged={})
    f = first_existing(FU_XLSX)
    if f is not None:
        site, _ = read_fu(f)
        site = site.assign(key=site.name.map(norm_name))
        rows = []
        for _, r in y[y.net_source.isin(WU_NETWORK)].iterrows():
            for _, s in site[site.key == norm_name(r["name"])].iterrows():
                d = float(max(abs(s.lat - r.lat), abs(s.lon - r.lon)))
                rows.append(dict(ID=int(r.ID), name=r["name"], lat=r.lat, lon=r.lon, ref_lat=float(s.lat),
                                 ref_lon=float(s.lon), cheb_deg=round(d, 4), flagged=d > NAME_COORD_DEG))
                if d > NAME_COORD_DEG:
                    out["flagged"].setdefault(int(r.ID), []).append(
                        f"qtp_fu_temp {s['name']} ({num_str(round(s.lat, 6))}, {num_str(round(s.lon, 6))}) "
                        f"cheb {d:.3f}")
        out["fu"] = dict(available=True, file=str(f.relative_to(ROOT)), n_pairs=len(rows),
                         n_flagged=sum(x["flagged"] for x in rows), pairs=rows,
                         rule="출처가 Wu Q. 관측망(WU_NETWORK)인 행과 이름이 같은 qtp_fu_temp 지점")
    f = first_existing(LIUX_XLSX)
    if f is not None:
        lx = read_liux(f)
        lx = lx.assign(key=lx.name.map(norm_name), alt_cm=lx.alt_m.astype(float) * 100.0)
        rows = []
        for _, r in y.iterrows():
            m = lx[(lx.key == norm_name(r["name"])) & ((lx.alt_cm - r.alt_cm).abs() <= SAME_VALUE_CM)]
            for _, s in m.iterrows():
                d = float(max(abs(s.lat - r.lat), abs(s.lon - r.lon)))
                rows.append(dict(ID=int(r.ID), name=r["name"], lat=r.lat, lon=r.lon, alt_cm=r.alt_cm,
                                 ref_code=str(s.code), ref_lat=float(s.lat), ref_lon=float(s.lon),
                                 ref_period=str(s.period), ref_source=str(s.src), cheb_deg=round(d, 4),
                                 flagged=d > NAME_COORD_DEG))
                if d > NAME_COORD_DEG:
                    out["flagged"].setdefault(int(r.ID), []).append(
                        f"qtp_liux_compile {s.code} {s['name']} ({num_str(round(s.lat, 6))}, "
                        f"{num_str(round(s.lon, 6))}) cheb {d:.3f}")
        out["liux"] = dict(available=True, file=str(f.relative_to(ROOT)), n_pairs=len(rows),
                           n_flagged=sum(x["flagged"] for x in rows), pairs=rows,
                           rule="이름이 같고 값이 1 cm 이내로 같은 qtp_liux_compile 기록(같은 기록으로 본다)")
    return out


# ---------------------------------------------------------------- 조립
def build(y: pd.DataFrame, ch: pd.DataFrame, calm: dict, dups: dict, inner: dict, names: dict):
    recs, detail = [], []
    for idx, r in y.iterrows():
        sid = f"{SRC_ID}_{int(r.ID):04d}"
        is_chang = r.src_label == CHANG_LABEL_IN_SHEET
        is_calm = r.src_label == CALM_LABEL_IN_SHEET
        notes, flags = [f"elev_m={num_str(r.elev_m)}"], []
        year = year_min = year_max = ""
        value_kind = ""
        if is_chang:
            c = ch.iloc[idx]
            a, b = (c.period.split("-") + [c.period])[:2]
            year_min, year_max = int(a), int(b)
            if year_min > year_max:
                raise SystemExit(f"관측 기간이 뒤집혀 있다: {c.period}")
            if year_min == year_max:
                year, value_kind = year_min, "annual_value"
            else:
                value_kind = "multiyear_mean"
            notes.append(f"period_table_s1={c.period}")
            orig = f"Chang et al. 2024 Table S1 (sheet label: {r.src_label}); original source: {c.source}"
            if c.source in RAILWAY_SOURCES:
                notes.append("source is railway subgrade monitoring; site disturbance not verified")
        else:
            orig = f"{r.src_label}; {r.src_url}" if r.src_url else r.src_label
            notes.append("year not given in source table")
        method, label_def, rc = "unknown", "unknown", int(r.rc_mark)
        if is_calm:
            method, label_def = "borehole_temp", "temp_derived"
            ck = calm.get("sites", {}).get(norm_name(r["name"]), {}) if calm.get("available") else {}
            if ck.get("found"):
                if ck["all_censored"] and ck["value_matches"]:
                    rc = 1
                    notes.append(f"CALM record {ck['years_all'][0]}-{ck['years_all'][1]}: all {ck['n_censored']} "
                                 f"values are '>{num_str(ck['censored_bound_cm'])}' (lower bound)")
                else:
                    value_kind = "multiyear_mean" if ck["value_matches"] else ""
                    notes.append(
                        f"CALM record mean of uncensored years {ck['years_uncensored'][0]}-"
                        f"{ck['years_uncensored'][1]} (n={ck['n_uncensored']}) = {ck['mean_uncensored_cm']} cm"
                        + (f"; {ck['n_censored']} censored years excluded from that mean" if ck["n_censored"] else "")
                        + ("" if ck["value_matches"] else "; does not match sheet value"))
            else:
                notes.append("CALM record not compared")
        # 연도 QC
        if year_max == "":
            flags.append("date")
        elif int(year_min) < 1990:
            flags.append("year_pre1990")
        if not (0 < r.alt_cm <= 600):
            flags.append("range")
        if not in_tibet(r.lat, r.lon):
            flags.append("coord")
            notes.append("outside Tibet box")
        if int(r.ID) in names["flagged"]:
            if "coord" not in flags:
                flags.append("coord")
            notes.append("same-name site has different coordinates: " + " | ".join(names["flagged"][int(r.ID)]))
        # dup_of
        cands, inn = dups[int(r.ID)], inner[int(r.ID)]
        dup_of = ""
        if cands:
            top = cands[0]
            dup_of = top["ref"]
            note = f"dup_kind={top['kind']}; dup_cheb_deg={top['cheb_deg']:.5f}; dup_name={top['ref_name']}"
            if top["ref_years"]:
                a, b = top["ref_years"]
                note += f"; dup_years={a}" + (f"-{b}" if b != a else "")
                if top["years_overlap"] is False and top["kind"] != "near_only":
                    note += "; dup_years_differ=1"
            notes.append(note)
            if len(cands) > 1:
                notes.append("other_within_0.01deg=" + ",".join(
                    c["ref"] + ("" if c["kind"] == "near_only" else f"({c['kind']})") for c in cands[1:]))
        earlier_same = [i for i in inn["same"] if i < int(r.ID)]
        if not dup_of and earlier_same:
            dup_of = f"{SRC_ID}:{SRC_ID}_{min(earlier_same):04d}"
            notes.append("dup_kind=same_value_in_table")
        if inn["same"]:
            notes.append("same_value_in_table=" + ",".join(f"{SRC_ID}_{i:04d}" for i in inn["same"]))
        other_near = [i for i in inn["near"] if i not in inn["same"]]
        if other_near:
            notes.append("near_in_table=" + ",".join(f"{SRC_ID}_{i:04d}" for i in other_near))
        macro = "Tibet" if in_tibet(r.lat, r.lon) else "other"
        recs.append(dict(
            src_id=SRC_ID, site_id=sid, site_name=r["name"], lat=num_str(r.lat), lon=num_str(r.lon),
            year=year, month="", alt_cm=num_str(r.alt_cm), method=method, label_def=label_def, n_obs=1,
            country="China", macro=macro,
            citation=CITATION + ("; " + CITATION_CHANG if is_chang else ""),
            license=LICENSE_YANG + ("; " + LICENSE_CHANG_PART if is_chang else ""),
            subunit=subunit_of(r.lat, r.lon) if macro == "Tibet" else "", date="", eos_basis="none",
            value_kind=value_kind, year_min=year_min, year_max=year_max, alt_sd_cm="",
            coord_prec_deg=num_str(coord_prec(r.lat, r.lon)), disturbed="", disturb_type="",
            right_censored=rc, orig_source=orig, dup_of=dup_of,
            qc_flag=";".join(flags) if flags else "ok", notes="; ".join(notes)))
        detail.append(dict(site_id=sid, site_name=r["name"], candidates=cands, in_table=inn))
    df = pd.DataFrame(recs, columns=COLUMNS)
    if not df.site_id.is_unique:
        raise SystemExit("site_id 가 중복된다")
    return df, detail


def cell_estimate(df: pd.DataFrame, ref: pd.DataFrame) -> dict:
    """약 1 km 셀 수의 예비 추정(6B.3 의 색인). v4 조립의 확정 값이 아니다."""
    f4 = ref[ref.group == "v3_f4_direct"]

    def count(d: pd.DataFrame) -> dict:
        if d.empty:
            return dict(n_rows=0, n_cells=0, n_cells_new=0, n_blocks_new=0)
        lat, lon = d.lat.astype(float).values, d.lon.astype(float).values
        ky = np.floor(lat / CELL).astype(int)
        phi = np.deg2rad((ky + 0.5) * CELL)
        kx = np.floor(lon * np.cos(phi) / CELL).astype(int)
        c = pd.DataFrame(dict(ky=ky, kx=kx, lat=lat, lon=lon)).groupby(["ky", "kx"], as_index=False).mean()
        near = np.array([bool((cheb(a, b, f4.lat.values, f4.lon.values) <= DUP_DEG + EPS).any()) if len(f4) else False
                         for a, b in zip(c.lat, c.lon)])
        new = c[~near]
        block = (np.floor(new.lat / 0.5) * 100000 + np.floor(new.lon / 0.5)).astype(int)
        return dict(n_rows=int(len(d)), n_cells=int(len(c)), n_cells_near_v3_f4_direct=int(near.sum()),
                    n_cells_new=int(len(new)), n_blocks_new=int(block.nunique()))
    ext = df.dup_of.ne("") & ~df.dup_of.str.startswith(SRC_ID + ":")
    return dict(
        rule="ky = floor(lat/0.009), kx = floor(lon*cos(phi)/0.009), phi = (ky+0.5)*0.009 도. 셀 좌표는 셀 안 "
             "좌표의 평균. v3 F4_direct 셀과 체비쇼프 0.01° 이내인 셀은 뺀다. 블록은 0.5°",
        all_rows=count(df),
        rows_without_dup_of=count(df[df.dup_of == ""]),
        rows_without_external_dup_of=count(df[~ext]),
        rows_qc_ok_without_dup_of=count(df[(df.dup_of == "") & (df.qc_flag == "ok")]),
        note="자료원 계획의 추정은 새 셀 123개(블록 37개)이고 0.01° 격자로 센 값이다. 격자 정의가 달라 값이 다를 수 있다")


def dist(v: pd.Series) -> dict:
    v = v.astype(float)
    if v.empty:
        return dict(n=0)
    return dict(n=int(len(v)), min=float(v.min()), p25=float(v.quantile(0.25)), median=float(v.median()),
                mean=round(float(v.mean()), 1), p75=float(v.quantile(0.75)), max=float(v.max()))


def git_commit() -> str:
    try:
        return subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True,
                              check=True).stdout.strip()
    except Exception:
        return "unknown"


def main() -> None:
    for p in (XLSX, SI_PDF, SCHEMA, V3):
        if not p.exists():
            raise SystemExit(f"입력 파일이 없다: {p}")
    md5 = file_hash(XLSX, "md5")
    if md5 != XLSX_MD5:
        raise SystemExit(f"xlsx 의 md5 가 Zenodo 기록과 다르다: {md5}")
    bbox_state = ensure_bbox()
    schema = json.loads(SCHEMA.read_text(encoding="utf-8"))
    if [c["name"] for c in schema["columns"]] != COLUMNS:
        raise SystemExit("형식 정의의 열과 파서의 열이 다르다")
    vocab = {c["name"]: c["vocab"] for c in schema["columns"] if "vocab" in c}

    y, sheet_info = read_yang()
    ch, chang_info = read_chang()
    chang_info["match"] = match_chang(y, ch)
    y["net_source"] = [ch.iloc[i].source if i < N_CHANG else y.src_label[i] for i in range(len(y))]
    per = ch.period.str.extract(r"(\d{4})(?:-(\d{4}))?")
    y["p_min"] = pd.Series(per[0].astype(float).values, index=y.index[:N_CHANG]).reindex(y.index)
    y["p_max"] = pd.Series(per[1].fillna(per[0]).astype(float).values, index=y.index[:N_CHANG]).reindex(y.index)
    calm = calm_check(y)
    ref, ref_mode = load_refs()
    dups = find_dups(y, ref)
    inner = internal_pairs(y)
    names = name_checks(y)
    df, detail = build(y, ch, calm, dups, inner, names)

    for col, allowed in vocab.items():
        bad = set(df[col].astype(str)) - {str(a) for a in allowed} - {""}
        if bad:
            raise SystemExit(f"{col} 에 형식 정의 밖의 값이 있다: {bad}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    df.to_csv(POINTS, index=False, encoding="utf-8", quoting=csv.QUOTE_MINIMAL, lineterminator="\n")

    alt = df.alt_cm.astype(float)
    ymin = pd.to_numeric(df.year_min, errors="coerce")
    ymax = pd.to_numeric(df.year_max, errors="coerce")
    files = []
    for p in (XLSX, SI_PDF, *SI_BBOX.values(), RAW / "41612_2024_866_MOESM1_ESM_layout.txt", ZENODO_JSON):
        if p.exists():
            files.append(dict(name=p.name, size=p.stat().st_size, sha256=file_hash(p, "sha256"),
                              md5=file_hash(p, "md5")))
    zen = json.loads(ZENODO_JSON.read_text(encoding="utf-8")) if ZENODO_JSON.exists() else {}
    dup_src = df.dup_of.str.split(":").str[0]
    n_kind = {k: int(df.notes.str.contains(f"dup_kind={k}(?:;|$)").sum())
              for k in ("same_value_and_name", "same_value", "same_name", "near_only", "same_value_in_table")}
    year_conflicts = [dict(site_id=d["site_id"], site_name=d["site_name"], ref=d["candidates"][0]["ref"],
                           kind=d["candidates"][0]["kind"], ref_years=d["candidates"][0]["ref_years"],
                           period_table_s1=ch.period.iloc[i])
                      for i, d in enumerate(detail)
                      if d["candidates"] and d["candidates"][0]["years_overlap"] is False
                      and d["candidates"][0]["kind"] != "near_only"]
    meta = dict(
        src_id=SRC_ID, name=SRC_NAME, url=URL, doi=DOI, related_doi=[CHANG_DOI], accessed=ACCESSED,
        download=dict(xlsx=XLSX_URL, chang_si_pdf=CHANG_SI_URL,
                      local_copy="data/raw/qtp_alt_open_2026-09-29/(zenodo18150789, chang2024_npj_SI) 에서 복사",
                      verified="xlsx 의 md5 는 Zenodo API 의 checksum 과 같다. PDF 는 2026-09-29 에 다시 받아 "
                               "sha256 이 같음을 확인했다",
                      not_downloaded=[f["key"] for f in zen.get("files", []) if f["key"].endswith(".nc")]),
        zenodo=dict(title=zen.get("metadata", {}).get("title"),
                    creators=[c.get("name") for c in zen.get("metadata", {}).get("creators", [])],
                    publication_date=zen.get("metadata", {}).get("publication_date"),
                    license=zen.get("metadata", {}).get("license"),
                    access_right=zen.get("metadata", {}).get("access_right"),
                    description_statement="Borehole measurements of active layer thickness (ALT) and mean annual "
                                          "ground temperature (MAGT) ... compiled from published studies"),
        files=files,
        license="CC BY 4.0 (Zenodo 18150789). Chang 등 2024 의 본문 약관은 CC BY-NC-ND 4.0 이다",
        license_detail=dict(
            yang_table="CC BY 4.0. Zenodo API 의 license.id = cc-by-4.0",
            chang_article="CC BY-NC-ND 4.0. 논문 누리집의 Rights and permissions 문구(2026-09-29 확인). "
                          "자료원 계획과 SOURCES.json 의 'CC BY 4.0' 표기는 맞지 않는다",
            chang_supplement="보충 자료 PDF 안에 약관 문구가 없다. 개별 약관은 확인하지 못했다",
            columns_from_chang=["year", "year_min", "year_max", "value_kind", "orig_source(원출처 이름)",
                                "notes 의 period_table_s1"],
            handling="점 자료의 앞 80행은 Chang 표 S1 에서 옮긴 관측 기간과 원출처 이름을 담는다. 약관이 "
                     "변경물의 공유를 허용하지 않으므로 사용자가 판단할 때까지 점 자료를 커밋하지 않는다"
                     "(ext_labels/.gitignore 에 추가). 좌표, 고도, ALT 값은 Yang 표(CC BY 4.0)에 있다"),
        citation=CITATION, citation_related=CITATION_CHANG,
        parser=str(Path(__file__).resolve().relative_to(ROOT)), parser_git_commit=git_commit(),
        parser_committed=False, schema_version=schema["version"],
        sheet=sheet_info, table_s1=chang_info,
        n_rows_raw=int(sheet_info["n_rows_alt"]), n_rows_points=int(len(df)), n_sites=int(df.site_id.nunique()),
        n_site_names=int(df.site_name.nunique()),
        repeated_site_names={k: int(v) for k, v in df.site_name.value_counts().items() if v > 1},
        n_by_label_def={k: int(v) for k, v in df.label_def.value_counts().items()},
        n_by_method={k: int(v) for k, v in df.method.value_counts().items()},
        n_by_macro={k: int(v) for k, v in df.macro.value_counts().items()},
        n_by_subunit={k: int(v) for k, v in df.subunit.value_counts().items()},
        n_by_value_kind={(k or "blank"): int(v) for k, v in df.value_kind.value_counts().items()},
        n_by_qc_flag={k: int(v) for k, v in df.qc_flag.value_counts().items()},
        n_by_sheet_source={k: int(v) for k, v in y.src_label.value_counts().items()},
        n_by_origin={k: int(v) for k, v in y.net_source.value_counts().items()},
        n_by_coord_prec={k: int(v) for k, v in df.coord_prec_deg.value_counts().items()},
        n_right_censored=int((df.right_censored == 1).sum()),
        n_with_year=int((ymax.notna()).sum()), n_without_year=int(ymax.isna().sum()),
        year_range=[int(ymin.min()), int(ymax.max())],
        n_by_period={k: int(v) for k, v in ch.period.value_counts().sort_index().items()},
        alt_cm=dist(alt),
        alt_cm_by_label_def={k: dist(g.alt_cm) for k, g in df.groupby("label_def")},
        alt_cm_by_subunit={k: dist(g.alt_cm) for k, g in df.groupby("subunit")},
        alt_cm_qc_ok_without_dup_of=dist(df[(df.qc_flag == "ok") & (df.dup_of == "")].alt_cm),
        lat_range=[float(df.lat.astype(float).min()), float(df.lat.astype(float).max())],
        lon_range=[float(df.lon.astype(float).min()), float(df.lon.astype(float).max())],
        n_excluded_by_reason={}, n_excluded_total=0,
        exclusion_note="지운 행은 없다. 시트 MAGT 는 ALT 라벨이 아니므로 읽지 않았다",
        unit_conversion="없다. 시트 머리글이 ALT (cm) 이고 표 S1 의 ALT (cm) 와 80행이 같다",
        missing_value_code="시트 ALT 에 결측 부호와 빈 값이 없다",
        coordinate_conversion="없다. 원자료가 십진 도(북위, 동경)다. 부호 변환도 없다",
        coord_prec_rule="위도와 경도의 소수 자리 수 가운데 큰 값(최소 2자리). 소수 4자리 이상이고 위도와 경도가 "
                        "모두 분의 배수면 0.0167",
        qc_flag_rule=dict(date="연도가 없는 행(72행). 1990년 이후 규칙을 확인할 수 없다",
                          coord=f"Tibet 상자 밖이거나 이름이 같은 다른 자료원 지점과 좌표가 {NAME_COORD_DEG}° 넘게 "
                                "다른 행. 좌표는 고치지 않았다",
                          range="0 < alt_cm <= 600 밖", year_pre1990="관측 기간의 첫 해가 1990년 이전"),
        macro_rule=dict(box=TIBET_BOX, note="6B.4 의 지리 정의. 라벨 값을 쓰지 않는다"),
        subunit_rule=[dict(subunit=n, label=k, box=b) for n, k, b in SUBUNIT_RULE]
                     + [dict(subunit="QTP_other", label="그 밖", box="위 네 상자 밖의 Tibet 범위")],
        calm_check=calm,
        dup_rule=dict(distance="체비쇼프 0.01° 이내(허용차 1e-9)",
                      priority=["fidelity_base_v3 의 F4_direct", "fidelity_base_v3 의 그 밖(CALM 지온 등)",
                                "qtp_du_gpr", "qtp_fu_temp", "qtp_liux_compile", "표 안의 앞 기록(값이 같을 때)"],
                      order="v3 F4_direct 셀, 같은 지점의 근거(값이 1 cm 이내로 같거나 이름이 같다), 자료원 우선순위, 거리",
                      series="연 값이 4개 이상인 대상은 이 표의 관측 기간 평균과 비교한다. 기간이 없으면 값 비교를 "
                             "하지 않는다",
                      caution="거리 기준이다. dup_kind=near_only 는 같은 셀 크기 안의 다른 지점일 수 있다",
                      id_format="<src_id>:<site_id>. v3 는 fidelity_base_v3:<loc_id>"),
        dup_reference=ref_mode,
        n_with_dup_of=int((df.dup_of != "").sum()),
        n_dup_of_by_source={k: int(v) for k, v in dup_src[df.dup_of != ""].value_counts().items()},
        n_dup_kind=n_kind,
        dup_year_conflicts=dict(n=len(year_conflicts), rows=year_conflicts,
                                note="같은 지점으로 본 대상의 연도가 표 S1 의 관측 기간과 겹치지 않는 행"),
        dup_detail=detail,
        name_coordinate_check=names,
        preliminary_cell_estimate=cell_estimate(df, ref),
        unverified_items=[
            "측정 방법. Zenodo 설명문은 시추공 측정(Borehole measurements)이라고 적고, Chang 등 2024 본문은 "
            "토양 센서, 시추공, GPR 이라고 적는다. 기록별 방법은 표에 없다",
            "연도가 없는 72지점의 관측 연도. 원 논문(Jin 2009, Chen 2016, Cao 2018, Qin 2017, Wang 2017, "
            "Wu 2012, Xie 2012, Xie 2015)을 읽지 않았다. 1990년 이전 값이 섞였는지 확인하지 못했다",
            "관측 기간이 여러 해인 행의 값이 기간 평균인지는 표 S1 에 적혀 있지 않다. Wu 2012 출처 지점은 "
            "qtp_fu_temp 의 2006-2010 평균과 가깝다",
            "표 S1 의 XD, BQ 지점의 연도는 2012 다. qtp_liux_compile 은 BQ 를 2011, qtp_du_gpr 는 같은 값의 "
            "GPR 점(R01 등)을 2011 로 적는다. 어느 쪽이 맞는지 확인하지 못했다",
            "Chang 등 2024 본문은 티베트 지점을 81개로 적고 표 S1 에는 80행이 있다",
            "철도 노반 관측 자료에서 온 지점(Niu 2021 3지점, China Railway 7지점)이 시설 부지인지 "
            "확인하지 못했다. disturbed 는 모든 행에서 빈 칸이다",
            "좌표가 다른 자료원과 어긋나는 지점의 옳은 좌표. 고치지 않고 qc_flag 에 coord 를 붙였다",
            "Chang 보충 자료의 개별 약관",
            "계절 말 여부. 관측 날짜가 없다(month, date 는 빈 칸, eos_basis = none)",
        ])
    META.write_text(json.dumps(meta, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8")

    # ------------------------------------------------------------ 품질 요약
    print(f"[{SRC_ID}] 원자료 {meta['n_rows_raw']}행, 점 자료 {len(df)}행, 지점 {meta['n_sites']}개, "
          f"이름 {meta['n_site_names']}종")
    print("bbox 파일:", bbox_state, "| 표 S1 대조:", chang_info["match"]["n_identical"], "/", N_CHANG,
          "| 출처 칸 복원의 최대 중심 차이", chang_info["max_center_dev_pt"], "pt")
    print("label_def:", meta["n_by_label_def"], "| method:", meta["n_by_method"])
    print("macro:", meta["n_by_macro"], "| subunit:", meta["n_by_subunit"])
    print("qc_flag:", meta["n_by_qc_flag"], "| right_censored:", meta["n_right_censored"])
    print("연도 있는 행:", meta["n_with_year"], "| 연도 범위:", meta["year_range"],
          "| value_kind:", meta["n_by_value_kind"])
    print("ALT(cm):", meta["alt_cm"])
    print("dup_of:", meta["n_with_dup_of"], meta["n_dup_of_by_source"], meta["n_dup_kind"])
    print("대조 자료:", {k: v["mode"] for k, v in ref_mode.items()})
    print("이름 대조 coord 표시:", {k: v for k, v in names["flagged"].items()})
    if calm.get("available"):
        for k, v in calm["sites"].items():
            print("CALM", k, {a: v.get(a) for a in ("yang_alt_cm", "mean_uncensored_cm", "n_uncensored",
                                                    "n_censored", "all_censored", "value_matches")})
    print("셀 추정:", json.dumps({k: v for k, v in meta["preliminary_cell_estimate"].items()
                                if isinstance(v, dict)}, ensure_ascii=False))
    print("산출:", POINTS.relative_to(ROOT), META.relative_to(ROOT))


if __name__ == "__main__":
    sys.exit(main())
