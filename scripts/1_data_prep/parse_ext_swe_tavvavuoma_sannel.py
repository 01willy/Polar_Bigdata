"""Sannel 2020(Bolin Centre) Tavvavuoma 이탄지 계절 말 융해 깊이 파싱 → ext_labels 표준 점 자료.

계획서: docs/EXPERIMENT_PLAN_LG_2026-09-29.md 6B 절(LGD). 형식: data/processed/ext_labels/_schema.json (1.0).

입력(data/raw/swe_tavvavuoma_sannel/sannel-2020/)
- sannel-2020-tavvavuoma-snow-and-thaw-depth.xlsx
    'Read me'                 지점 9곳(T1-T9)의 좌표(도와 소수 분)와 경관 단위
    'Thaw depth, 2010-2012'   연 1회 방문(2010-09-02, 2011-09-03, 2012-08-23), 지점별 '평균 (반복 3회)' 문자열
                              표기 '>100' 은 측정 한계 초과, 'No data (in water)' 는 결측
                              각주에 설치 시 융해 깊이 2건(T1 2005-09-23 75 cm, T4 2006-09-04 115 cm)
- sannel-2020-tavvavuoma-ground-temperatures.xlsx
    지점별 3시간 간격 지온 1년(T1-T7 은 2007년, T8-T9 는 2008년). 영구동토 유무 판정에만 쓴다.

라벨 규칙(자료원 규칙: 영구동토가 있는 지점의 계절 말 융해 깊이만 쓴다)
- 영구동토 지점 판정(파일 안의 자료만 쓴다). 아래 가운데 하나를 만족하면 영구동토 지점이다.
    (a) 연중 최고 지온이 0 °C 이하인 센서가 있다.
    (b) 가장 깊은 센서의 연평균 지온이 0 °C 미만이다.
    (c) 2010-2012년 계절 말 측정에서 동결면에 닿은 값(숫자 값)이 있다.
  어느 것도 만족하지 않는 지점(호안, 호수, 펜)의 '>100' 기록과 'No data' 기록은 점 자료에 넣지 않고
  메타의 제외 사유에 수를 적는다. 방문 단위 전체 표는 data/raw/swe_tavvavuoma_sannel/ 에 남긴다.
- method: 파일에 측정 기기 기재가 없으므로 other 다.
- label_def: 관측 월이 8-9월이면 direct_eos, 아니면 direct_dated. eos_basis 는 record_date 다.
- alt_cm: 파일이 적은 지점 값(반복 3회의 평균, 정수 cm)을 쓴다. 반복값의 평균과 표준편차는 따로 적는다.
- '>100' 은 alt_cm = 100, right_censored = 1 로 둔다(값은 하한이다).
- 좌표: 도와 소수 분을 십진 도로 바꾼다(도 + 분/60). 소수 6자리로 적는다(원자료 정밀도 0.001분 = 0.0000167도).

산출
- data/processed/ext_labels/swe_tavvavuoma_sannel_points.csv
- data/processed/ext_labels/swe_tavvavuoma_sannel_meta.json
- data/raw/swe_tavvavuoma_sannel/derived_thaw_depth_visits.csv   방문 단위 전체 표(제외 기록 포함)
- data/raw/swe_tavvavuoma_sannel/derived_ground_temp_summary.csv  지점·깊이별 지온 요약

실행(저장소 루트에서, 스레드 1개)
  OMP_NUM_THREADS=1 OPENBLAS_NUM_THREADS=1 MKL_NUM_THREADS=1 \
    /home/anaconda3/bin/python scripts/1_data_prep/parse_ext_swe_tavvavuoma_sannel.py
"""
import csv
import datetime as dt
import hashlib
import json
import math
import os
import re
import statistics
import subprocess
import sys

import openpyxl

SRC_ID = "swe_tavvavuoma_sannel"
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW_DIR = os.path.join(ROOT, "data", "raw", SRC_ID)
XLSX_DIR = os.path.join(RAW_DIR, "sannel-2020")
F_THAW = os.path.join(XLSX_DIR, "sannel-2020-tavvavuoma-snow-and-thaw-depth.xlsx")
F_TEMP = os.path.join(XLSX_DIR, "sannel-2020-tavvavuoma-ground-temperatures.xlsx")
F_ZIP = os.path.join(RAW_DIR, "sannel-2020.zip")
OUT_DIR = os.path.join(ROOT, "data", "processed", "ext_labels")
F_POINTS = os.path.join(OUT_DIR, f"{SRC_ID}_points.csv")
F_META = os.path.join(OUT_DIR, f"{SRC_ID}_meta.json")
F_VISITS = os.path.join(RAW_DIR, "derived_thaw_depth_visits.csv")
F_GT = os.path.join(RAW_DIR, "derived_ground_temp_summary.csv")
F_V3 = os.path.join(ROOT, "data", "processed", "fidelity_base_v3.csv")

URL = "https://doi.org/10.17043/sannel-2020-temperature-1"
URL_FILE = "https://bolin.su.se/data/uploads/sannel-2020.zip"
DOI = "10.17043/sannel-2020-temperature-1"
ACCESSED = "2026-09-29"
LICENSE = "CC BY 4.0"
CITATION = ("Sannel, A.B.K. (2020) Ground temperature, thaw depth and snow depth in a permafrost "
            "peatland in Tavvavuoma, northern Sweden. Dataset version 1. Bolin Centre Database. "
            "https://doi.org/10.17043/sannel-2020-temperature-1")
ORIG_SOURCE = ("Sannel, A.B.K. (2020) Ground temperature and snow depth variability within a subarctic "
               "peat plateau landscape. Permafrost and Periglacial Processes 31, 255-263, "
               "doi:10.1002/ppp.2045")
COUNTRY = "Sweden"
MACRO = "NAtlantic"
SUBUNIT = "Scandinavia"
PROBE_LIMIT_CM = 100.0          # 파일의 '>100' 표기에서 읽은 측정 한계
COORD_PREC_DEG = 0.001 / 60.0   # 원자료 좌표는 0.001분 단위
EOS_MONTHS = (8, 9)
ALT_MIN, ALT_MAX = 0.0, 600.0
CELL_DEG = 0.009

COLUMNS = ["src_id", "site_id", "site_name", "lat", "lon", "year", "month", "alt_cm", "method",
           "label_def", "n_obs", "country", "macro", "citation", "license",
           "subunit", "date", "eos_basis", "value_kind", "year_min", "year_max", "alt_sd_cm",
           "coord_prec_deg", "disturbed", "disturb_type", "right_censored", "orig_source",
           "dup_of", "qc_flag", "notes"]

RE_DM = re.compile(r"^\s*(\d+)\s*°\s*([\d.]+)\s*[’'′]\s*([NSEW])\s*$")
RE_HDR = re.compile(r"^\s*(T\d+)\s*-\s*late season thaw depth \((cm)\)\s*(\**)\s*$")
RE_VAL = re.compile(r"^\s*(>?)\s*(\d+(?:\.\d+)?)\s*\(([^)]*)\)\s*$")
RE_FOOT = re.compile(r"^(\*+)\s*Thaw depth at installation (\d{4}-\d{2}-\d{2}):\s*(\d+(?:\.\d+)?)\s*cm\s*$")
RE_DEPTH = re.compile(r"^\s*(T\d+),\s*(\d+)\s*cm")


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for blk in iter(lambda: f.read(1 << 20), b""):
            h.update(blk)
    return h.hexdigest()


def dm_to_deg(text):
    """'68°27.720’N' 형식을 십진 도로 바꾼다. 남위와 서경은 음수다."""
    m = RE_DM.match(str(text))
    if not m:
        raise ValueError(f"좌표 형식을 읽지 못했다: {text!r}")
    deg = int(m.group(1)) + float(m.group(2)) / 60.0
    if m.group(3) in ("S", "W"):
        deg = -deg
    return deg


def fmt(v, nd=None):
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        if nd is not None:
            return f"{v:.{nd}f}"
        return f"{v:g}"
    return str(v)


def read_sites(wb):
    """'Read me' 시트에서 지점 좌표와 경관 단위를 읽는다."""
    ws = wb["Read me"]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    hdr_i = None
    for i, r in enumerate(rows):
        if r and r[0] == "Location" and r[1] == "Latitude" and r[2] == "Longitude":
            hdr_i = i
            break
    if hdr_i is None:
        return {}
    sites = {}
    for r in rows[hdr_i + 1:]:
        if not r or r[0] is None:
            break
        code = str(r[0]).strip()            # 예: 'T5/S5/SX'
        tid = code.split("/")[0]
        sites[tid] = dict(site_id=tid, file_code=code,
                          lat_raw=str(r[1]).strip(), lon_raw=str(r[2]).strip(),
                          lat=dm_to_deg(r[1]), lon=dm_to_deg(r[2]),
                          landscape_unit=str(r[3]).strip())
    return sites


def read_thaw(wb):
    """융해 깊이 시트를 방문 단위 기록으로 읽는다."""
    name = [n for n in wb.sheetnames if n.lower().startswith("thaw depth")][0]
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = rows[0]
    cols = {}
    marks = {}
    for j, h in enumerate(hdr):
        if h is None or j < 2:
            continue
        m = RE_HDR.match(str(h).replace("-late", "- late"))
        if not m:
            raise ValueError(f"머리글을 읽지 못했다: {h!r}")
        if m.group(2) != "cm":
            raise ValueError(f"단위가 cm 가 아니다: {h!r}")
        cols[j] = m.group(1)
        if m.group(3):
            marks[m.group(3)] = m.group(1)
    recs = []
    foots = []
    for r in rows[1:]:
        if r[0] is None:
            continue
        if isinstance(r[0], str):
            m = RE_FOOT.match(r[0].strip())
            if m:
                foots.append(dict(mark=m.group(1), site_id=marks[m.group(1)],
                                  date=dt.date.fromisoformat(m.group(2)),
                                  value=float(m.group(3)), raw=r[0].strip()))
            continue
        year = int(r[0])
        date = r[1].date() if isinstance(r[1], dt.datetime) else None
        for j, tid in cols.items():
            raw = "" if r[j] is None else str(r[j]).strip()
            rec = dict(site_id=tid, year=year, date=date, raw=raw, kind="table",
                       value=None, censored=0, reps=[], reps_censored=[], status="")
            m = RE_VAL.match(raw)
            if m:
                rec["censored"] = 1 if m.group(1) == ">" else 0
                rec["value"] = float(m.group(2))
                for tok in m.group(3).split(","):
                    tok = tok.strip()
                    rec["reps_censored"].append(1 if tok.startswith(">") else 0)
                    rec["reps"].append(float(tok.lstrip(">").strip()))
                rec["status"] = "value"
            elif raw.lower().startswith("no data"):
                rec["status"] = "no_data"
            else:
                rec["status"] = "unparsed"
            recs.append(rec)
    for f in foots:
        recs.append(dict(site_id=f["site_id"], year=f["date"].year, date=f["date"], raw=f["raw"],
                         kind="footnote_installation", value=f["value"], censored=0,
                         reps=[], reps_censored=[], status="value"))
    return name, recs, foots


def read_ground_temp(path):
    """지점·센서 깊이별 지온 요약(연 최저, 평균, 최고, 0 °C 초과 시각 수)을 만든다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith("read"):
            continue
        it = ws.iter_rows(values_only=True)
        hdr = next(it)
        depth_cols = {}
        for j, h in enumerate(hdr):
            if h is None or j == 0:
                continue
            m = RE_DEPTH.match(str(h))
            if m:
                depth_cols[j] = (m.group(1), int(m.group(2)))
        vals = {j: [] for j in depth_cols}
        file_magt = {}
        t0 = t1 = None
        for r in it:
            if isinstance(r[0], dt.datetime):
                t0 = r[0] if t0 is None else t0
                t1 = r[0]
                for j in depth_cols:
                    if isinstance(r[j], (int, float)):
                        vals[j].append(float(r[j]))
            elif isinstance(r[0], str) and r[0].strip().lower().startswith("mean annual"):
                for j in depth_cols:
                    if isinstance(r[j], (int, float)):
                        file_magt[j] = float(r[j])
        for j, (tid, dcm) in depth_cols.items():
            v = vals[j]
            out.append(dict(site_id=tid, sheet=ws.title, obs_year=t0.year, t_first=t0.isoformat(sep=" "),
                            t_last=t1.isoformat(sep=" "), depth_cm=dcm, n=len(v),
                            t_min=min(v), t_mean=sum(v) / len(v), t_max=max(v),
                            n_above0=sum(1 for x in v if x > 0.0),
                            file_magt=file_magt.get(j)))
    wb.close()
    return out


def f2(x):
    """소수 2자리 표기. '-0.00' 은 '0.00' 으로 적는다."""
    t = f"{x:.2f}"
    return "0.00" if t == "-0.00" else t


def permafrost_evidence(tid, gt, recs):
    """영구동토 지점 판정. 반환: (판정, 근거 목록, 설명)."""
    g = sorted([x for x in gt if x["site_id"] == tid], key=lambda x: x["depth_cm"])
    crit = []
    desc = []
    frozen = [x for x in g if x["t_max"] <= 0.0]
    if frozen:
        crit.append("a")
        desc.append("연중 최고 지온 0 °C 이하 센서 깊이 " +
                    ", ".join(f"{x['depth_cm']} cm" for x in frozen) + f"({g[0]['obs_year']}년)")
    if g:
        d = g[-1]
        if d["t_mean"] < 0.0:
            crit.append("b")
        desc.append(f"최심 센서 {d['depth_cm']} cm 의 연평균 {f2(d['t_mean'])} °C, "
                    f"최저 {f2(d['t_min'])}, 최고 {f2(d['t_max'])}({d['obs_year']}년)")
    hit = [r for r in recs if r["site_id"] == tid and r["kind"] == "table"
           and r["status"] == "value" and not r["censored"]]
    if hit:
        crit.append("c")
        desc.append(f"2010-2012년 계절 말 측정에서 동결면 도달 {len(hit)}회")
    return (len(crit) > 0), crit, "; ".join(desc)


def cell_index(lat, lon):
    ky = math.floor(lat / CELL_DEG)
    phi = math.radians((ky + 0.5) * CELL_DEG)
    kx = math.floor(lon * math.cos(phi) / CELL_DEG)
    return ky, kx


def haversine_km(lat1, lon1, lat2, lon2):
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = (math.sin((p2 - p1) / 2) ** 2 +
         math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2)
    return 2 * 6371.0088 * math.asin(math.sqrt(a))


def nearest_v3(lat, lon):
    """v3 표(읽기 전용)에서 가장 가까운 F4_direct 셀을 찾는다. 표가 없으면 None."""
    if not os.path.exists(F_V3):
        return None
    natl = {"CALM_Scandinavia", "CALM_Svalbard", "CALM_Greenland"}
    best_all = best_other = None
    min_cheb = None
    with open(F_V3, newline="", encoding="utf-8") as f:
        for d in csv.DictReader(f):
            if d["source_id"] != "F4_direct":
                continue
            la, lo = float(d["lat"]), float(d["lon"])
            km = haversine_km(lat, lon, la, lo)
            cheb = max(abs(la - lat), abs(lo - lon))
            item = dict(loc_id=d["loc_id"], region=d["region"], lat=la, lon=lo, km=round(km, 1))
            if best_all is None or km < best_all["km_raw"]:
                best_all = dict(item, km_raw=km)
            if d["region"] not in natl and (best_other is None or km < best_other["km_raw"]):
                best_other = dict(item, km_raw=km)
            if min_cheb is None or cheb < min_cheb:
                min_cheb = cheb
    for b in (best_all, best_other):
        if b:
            b.pop("km_raw")
    return dict(nearest_f4_direct=best_all,
                nearest_f4_direct_excluding_natlantic_regions=best_other,
                natlantic_region_names_used=sorted(natl),
                min_chebyshev_deg_to_f4_direct=round(min_cheb, 4))


def main():
    for p in (F_THAW, F_TEMP):
        if not os.path.exists(p):
            sys.exit(f"원자료가 없다: {p}")
    os.makedirs(OUT_DIR, exist_ok=True)

    wb = openpyxl.load_workbook(F_THAW, data_only=True)
    sites = read_sites(wb)
    sheet_name, recs, foots = read_thaw(wb)
    gt = read_ground_temp(F_TEMP)

    unverified = []
    if not sites:
        # 좌표가 파일에 없으면 점 자료를 만들지 않는다(자료원 규칙).
        meta = dict(src_id=SRC_ID, url=URL, doi=DOI, accessed=ACCESSED, n_rows_points=0,
                    unverified_items=["좌표를 파일에서 찾지 못해 점 자료를 만들지 않았다"])
        with open(F_META, "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=1)
        sys.exit("좌표가 없어 점 자료를 만들지 않았다")

    # 좌표 점검: 북반구, 동경, 지점 개요 좌표(68°28'N 20°54'E)에서 0.02도 이내
    for s in sites.values():
        if not (68.44 <= s["lat"] <= 68.49 and 20.88 <= s["lon"] <= 20.92):
            raise ValueError(f"좌표가 지점 개요와 맞지 않는다: {s}")

    # 지온 요약의 연평균이 파일의 연평균 행과 같은지 점검
    gt_dev = max(abs(x["t_mean"] - x["file_magt"]) for x in gt if x["file_magt"] is not None)

    ev = {}
    for tid in sites:
        ok, crit, desc = permafrost_evidence(tid, gt, recs)
        ev[tid] = dict(permafrost_site=ok, criteria=crit, evidence=desc)

    # 중복 기록 점검: (지점, 연도, 날짜) 중복
    keys = [(r["site_id"], r["year"], r["date"]) for r in recs]
    n_dup = len(keys) - len(set(keys))

    excluded = {}
    points = []
    visits = []
    for r in sorted(recs, key=lambda x: (int(x["site_id"][1:]), x["year"])):
        s = sites[r["site_id"]]
        e = ev[r["site_id"]]
        reason = ""
        if r["status"] == "no_data":
            reason = "no_data_in_water"
        elif r["status"] == "unparsed":
            reason = "unparsed_text"
        elif not e["permafrost_site"]:
            reason = "no_permafrost_site_value_above_probe_limit" if r["censored"] else "no_permafrost_site"
        rep_mean = statistics.mean(r["reps"]) if r["reps"] else None
        rep_sd = statistics.stdev(r["reps"]) if len(r["reps"]) >= 2 else None
        visits.append(dict(site_id=r["site_id"], landscape_unit=s["landscape_unit"], year=r["year"],
                           date=r["date"].isoformat() if r["date"] else "", record_kind=r["kind"],
                           raw_text=r["raw"], stated_value_cm=fmt(r["value"]),
                           stated_censored=r["censored"],
                           replicates_cm=";".join((">" if c else "") + fmt(v)
                                                  for v, c in zip(r["reps"], r["reps_censored"])),
                           replicate_mean_cm=fmt(rep_mean, 2) if rep_mean is not None else "",
                           replicate_sd_cm=fmt(rep_sd, 2) if rep_sd is not None else "",
                           permafrost_site=int(e["permafrost_site"]),
                           permafrost_criteria="".join(e["criteria"]),
                           in_points=0 if reason else 1, exclude_reason=reason))
        if reason:
            excluded[reason] = excluded.get(reason, 0) + 1
            continue

        month = r["date"].month if r["date"] else None
        if r["date"] and r["date"].year != r["year"]:
            raise ValueError(f"연도 열과 날짜가 다르다: {r}")
        qc = []
        if r["censored"]:
            alt = PROBE_LIMIT_CM
            if abs(r["value"] - PROBE_LIMIT_CM) > 1e-9:
                raise ValueError(f"측정 한계 표기가 100 이 아니다: {r}")
        else:
            alt = r["value"]
        if not (ALT_MIN < alt <= ALT_MAX):
            qc.append("range")
        if r["year"] < 1990:
            qc.append("year_pre1990")
        if month is None:
            qc.append("date")
        label_def = "direct_eos" if month in EOS_MONTHS else "direct_dated"
        eos_basis = "record_date" if month in EOS_MONTHS else "none"

        notes = [f"경관 단위 {s['landscape_unit']}", f"영구동토 근거 {e['evidence']}"]
        if r["kind"] == "footnote_installation":
            notes.append("융해 깊이 시트의 각주에 적힌 설치 시 측정값이다. 반복 수 기재가 없어 n_obs 는 1 이다")
            later = [x for x in recs if x["site_id"] == r["site_id"] and x["kind"] == "table"
                     and x["status"] == "value"]
            if later and all(x["censored"] for x in later):
                notes.append("같은 지점의 2010-2012년 값은 모두 측정 한계 100 cm 초과다")
            n_obs = 1
        else:
            n_obs = len(r["reps"])
            if r["censored"]:
                notes.append(f"원자료 표기 '{r['raw']}'. 측정 한계 100 cm 초과이고 alt_cm 은 하한이다")
            else:
                notes.append(f"원자료 표기 '{r['raw']}'. 반복값 평균 {rep_mean:.2f} cm")
                if abs(rep_mean - r["value"]) > 0.5:
                    notes.append("파일의 지점 값과 반복값 평균의 차가 0.5 cm 를 넘는다")
                if any(abs(v - PROBE_LIMIT_CM) < 1e-9 for v in r["reps"]):
                    notes.append("반복값 가운데 100 cm 가 있다. 측정 한계에 닿은 값인지 확인하지 못했다")
        points.append(dict(
            src_id=SRC_ID, site_id=r["site_id"],
            site_name=f"Tavvavuoma {r['site_id']} ({s['landscape_unit']})",
            lat=fmt(s["lat"], 6), lon=fmt(s["lon"], 6), year=r["year"],
            month=month if month else "", alt_cm=fmt(alt), method="other", label_def=label_def,
            n_obs=n_obs, country=COUNTRY, macro=MACRO, citation=CITATION, license=LICENSE,
            subunit=SUBUNIT, date=r["date"].isoformat() if r["date"] else "", eos_basis=eos_basis,
            value_kind="single_visit", year_min="", year_max="",
            alt_sd_cm=fmt(rep_sd, 2) if (rep_sd is not None and not r["censored"]) else "",
            coord_prec_deg=f"{COORD_PREC_DEG:.7f}", disturbed=0, disturb_type="",
            right_censored=r["censored"], orig_source=ORIG_SOURCE, dup_of="",
            qc_flag=";".join(qc) if qc else "ok", notes=". ".join(notes)))

    with open(F_POINTS, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, lineterminator="\n")
        w.writeheader()
        w.writerows(points)
    with open(F_VISITS, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(visits[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(visits)
    with open(F_GT, "w", newline="", encoding="utf-8") as f:
        cols = ["site_id", "sheet", "obs_year", "t_first", "t_last", "depth_cm", "n", "t_min",
                "t_mean", "t_max", "n_above0", "file_magt"]
        w = csv.DictWriter(f, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        for x in gt:
            w.writerow({k: (f"{x[k]:.4f}" if isinstance(x[k], float) else x[k]) for k in cols})

    # ---------- 품질 요약 ----------
    def count(key, rows=points):
        c = {}
        for p in rows:
            c[str(p[key])] = c.get(str(p[key]), 0) + 1
        return dict(sorted(c.items()))

    main_rows = [p for p in points if p["label_def"] == "direct_eos" and p["right_censored"] == 0
                 and p["qc_flag"] == "ok"]
    vals = sorted(float(p["alt_cm"]) for p in main_rows)
    site_means = {}
    for p in main_rows:
        site_means.setdefault(p["site_id"], []).append(float(p["alt_cm"]))
    site_means = {k: round(statistics.mean(v), 2) for k, v in site_means.items()}
    table_rows = [p for p in main_rows if "각주" not in p["notes"]]
    site_means_table = {}
    for p in table_rows:
        site_means_table.setdefault(p["site_id"], []).append(float(p["alt_cm"]))
    site_means_table = {k: round(statistics.mean(v), 2) for k, v in site_means_table.items()}

    cells = {}
    for tid, s in sites.items():
        cells.setdefault("%d_%d" % cell_index(s["lat"], s["lon"]), []).append(tid)
    pt_sites = sorted({p["site_id"] for p in points}, key=lambda x: int(x[1:]))
    lat_c = statistics.mean(sites[t]["lat"] for t in pt_sites)
    lon_c = statistics.mean(sites[t]["lon"] for t in pt_sites)

    try:
        head = subprocess.run(["git", "-C", ROOT, "rev-parse", "HEAD"], capture_output=True,
                              text=True, check=True).stdout.strip()
    except Exception:
        head = ""
    n_table = sum(1 for r in recs if r["kind"] == "table")
    n_foot = sum(1 for r in recs if r["kind"] != "table")

    unverified += [
        "측정 기기와 방법은 파일에 기재가 없다. '>100' 표기로 보아 측정 한계가 100 cm 인 도구로 보이나 "
        "탐침인지는 확인하지 못했다. 원 논문(doi:10.1002/ppp.2045)의 본문은 출판사 누리집이 403 을 돌려주어 "
        "읽지 못했다. method 는 other 로 두었다",
        "T7 의 2011년 값은 파일이 84 cm 로 적었으나 반복값(74, 76, 100)의 평균은 83.33 cm 다. "
        "파일의 지점 값을 alt_cm 으로 썼다. 반복값 100 이 측정 한계에 닿은 값인지는 확인하지 못했다",
        "T1 은 지온 센서의 최대 깊이가 91 cm 다. 2007년 91 cm 의 연평균은 0 °C 미만이나 연중 최고는 0 °C 를 "
        "넘는다. 2010-2012년에 T1 아래에 영구동토가 남아 있었는지는 파일로 확인하지 못했다(세 해 모두 >100)",
        "T7 의 96 cm 지온은 연중 -0.12 에서 0.06 °C 범위다. 센서 정확도를 파일에서 확인하지 못했다. "
        "T7 은 계절 말 측정에서 동결면에 닿은 값이 있어 영구동토 지점으로 두었다",
        "지온 관측 연도(2007, 2008)와 융해 깊이 관측 연도(2010-2012)가 다르다. 영구동토 판정은 관측 연도가 "
        "다른 자료에 근거한다",
        "설치 시 융해 깊이 2건(T1 2005-09-23, T4 2006-09-04)의 측정 방법과 반복 수는 파일에 없다",
        "가장자리 지점(T1, T7, 열카르스트 호수 인접)을 교란 지점으로 볼지는 파일에 기준이 없다. disturbed 는 0 으로 두고 "
        "경관 단위를 notes 에 적었다",
        "같은 저자의 다른 자료(Bolin Centre sannel-2018-temperature-1, Tavvavuoma 2006-2013)는 이 작업에서 "
        "내려받지 않았다. 같은 1 km 셀의 연도를 늘릴 수 있는지는 확인하지 못했다",
    ]

    meta = dict(
        src_id=SRC_ID,
        name="Sannel 2020, Tavvavuoma 이탄지 지온, 융해 깊이, 적설 깊이(Bolin Centre Database)",
        url=URL, download_url=URL_FILE, doi=DOI, accessed=ACCESSED,
        files=[dict(name=os.path.relpath(p, RAW_DIR), size_bytes=os.path.getsize(p), sha256=sha256(p))
               for p in (F_ZIP, F_THAW, F_TEMP) if os.path.exists(p)],
        license=LICENSE,
        license_basis="자료 안내 쪽의 License 항목(Creative Commons Attribution 4.0 International License)",
        citation=CITATION, orig_source=ORIG_SOURCE,
        parser="scripts/1_data_prep/parse_ext_swe_tavvavuoma_sannel.py",
        parser_sha256=sha256(os.path.abspath(__file__)),
        parser_git_commit=head,
        parser_git_note="파서는 커밋하지 않은 작업 트리 파일이다. parser_git_commit 은 실행 시점의 HEAD 다",
        schema_version="1.0",
        thaw_sheet=sheet_name,
        n_rows_raw=len(recs),
        n_rows_raw_detail=dict(table_cells=n_table, footnote_installation_values=n_foot),
        n_rows_points=len(points),
        n_sites_raw=len(sites),
        n_sites_points=len(pt_sites),
        n_by_label_def=count("label_def"),
        n_by_method=count("method"),
        n_by_macro=count("macro"),
        n_by_subunit=count("subunit"),
        n_by_eos_basis=count("eos_basis"),
        n_by_right_censored=count("right_censored"),
        n_by_qc_flag=count("qc_flag"),
        n_by_month=count("month"),
        n_by_site=count("site_id"),
        n_excluded_by_reason=dict(sorted(excluded.items())),
        n_duplicate_site_year_date=n_dup,
        year_min=min(p["year"] for p in points), year_max=max(p["year"] for p in points),
        main_set=dict(
            definition="label_def = direct_eos, right_censored = 0, qc_flag = ok",
            n_rows=len(main_rows), n_sites=len(site_means),
            alt_cm_min=vals[0], alt_cm_median=statistics.median(vals), alt_cm_max=vals[-1],
            site_multiyear_mean_cm=site_means,
            n_rows_table_2010_2012=len(table_rows),
            site_multiyear_mean_cm_table_2010_2012=site_means_table),
        sites=[dict(site_id=t, file_code=s["file_code"], lat_raw=s["lat_raw"], lon_raw=s["lon_raw"],
                    lat=round(s["lat"], 6), lon=round(s["lon"], 6),
                    landscape_unit=s["landscape_unit"], in_points=t in pt_sites, **ev[t])
               for t, s in sites.items()],
        permafrost_rule=("(a) 연중 최고 지온이 0 °C 이하인 센서가 있다, (b) 최심 센서의 연평균 지온이 0 °C 미만이다, "
                         "(c) 2010-2012년 계절 말 측정에서 동결면에 닿은 값이 있다. 하나라도 만족하면 영구동토 지점이다"),
        ground_temp_check=dict(max_abs_diff_vs_file_magt=round(gt_dev, 6),
                               summary_file="data/raw/swe_tavvavuoma_sannel/derived_ground_temp_summary.csv"),
        visits_file="data/raw/swe_tavvavuoma_sannel/derived_thaw_depth_visits.csv",
        coordinate_conversion=("원자료는 도와 소수 분(예: 68°27.720’N, 20°54.118’E)이다. 십진 도 = 도 + 분/60. "
                               "북위와 동경이라 부호는 양수다. 소수 6자리로 적었다. 원자료 정밀도는 0.001분"
                               "(약 0.0000167도)이다. 측지계 기재는 파일에 없고 WGS84 로 가정했다"),
        units="원자료 단위는 cm 다(머리글의 '(cm)'). 변환하지 않았다",
        missing_codes="'No data (in water)' 는 결측이다. '>100' 은 측정 한계 초과다",
        censoring=dict(probe_limit_cm=PROBE_LIMIT_CM,
                       rule="'>100' 은 alt_cm = 100, right_censored = 1. 값은 하한이다"),
        cell_preview=dict(
            rule="ky = floor(lat/0.009), kx = floor(lon*cos(phi)/0.009), phi = (ky+0.5)*0.009 도",
            cells_all_sites=cells,
            n_cells_points_sites=len({"%d_%d" % cell_index(sites[t]["lat"], sites[t]["lon"])
                                      for t in pt_sites}),
            centroid_points_sites=dict(lat=round(lat_c, 6), lon=round(lon_c, 6)),
            cell_value_options_cm=dict(
                note=("셀 값은 v4 조립 단계에서 정한다. 아래는 주 집합 행으로 계산한 지점별 다년 평균의 평균이다. "
                      "T1 과 T4 는 절단되지 않은 값이 설치 시 측정값 1건뿐이고 2010-2012년 값은 모두 절단이다"),
                all_uncensored_rows=round(statistics.mean(site_means.values()), 2),
                table_2010_2012_rows_only=round(statistics.mean(site_means_table.values()), 2),
                n_sites_all_uncensored=len(site_means),
                n_sites_table_2010_2012=len(site_means_table))),
        v3_proximity=nearest_v3(lat_c, lon_c),
        unverified_items=unverified,
    )
    with open(F_META, "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=1)
        f.write("\n")

    print(f"원자료 기록 {len(recs)}건(표 {n_table}, 각주 {n_foot}), 점 자료 {len(points)}행, "
          f"지점 {len(pt_sites)}곳")
    print("제외:", meta["n_excluded_by_reason"])
    print("label_def:", meta["n_by_label_def"], "right_censored:", meta["n_by_right_censored"])
    print("주 집합:", meta["main_set"])
    print("셀:", meta["cell_preview"])
    print("v3 근접:", meta["v3_proximity"])
    print("지온 연평균 점검 최대 차:", meta["ground_temp_check"]["max_abs_diff_vs_file_magt"])


if __name__ == "__main__":
    main()
