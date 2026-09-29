"""ext_labels 파서: qtp_liux_compile (Liu X. 등 2022, figshare 21444879, 칭하이 고원 ALT 관측 문헌 집계표).

계획서: docs/EXPERIMENT_PLAN_LG_2026-09-29.md 6B 절(LGD, 개정 10).
형식 정의: data/processed/ext_labels/_schema.json (버전 1.0).

입력
- data/raw/qtp_liux_compile/ds01.xlsx (figshare 기록 21444879 의 유일한 파일, 31,302 바이트).
  없으면 탐색 단계 사본(data/raw/qtp_alt_open_2026-09-29/figshare21444879/ds01.xlsx)을 복사한다.
  둘 다 없으면 내려받기 명령을 출력하고 끝낸다(이 스크립트는 네트워크를 쓰지 않는다).
  시트 ALT_Data. 1행은 표 제목, 2행이 머리글, 3행부터 자료다.
  열: ID, ID_paper, Site name, Site Code, Latitude(°N), Longitude(°E), Observation period, ALT(m), Soure.
- data/processed/ext_labels/qtp_du_gpr_points.csv (읽기 전용). dup_of 판정의 기준 자료원이다.
  없으면 끝낸다(parse_ext_qtp_du_gpr.py 를 먼저 실행한다).
- data/processed/fidelity_base_v3.csv (읽기 전용, loc_id·lat·lon·region·source_id 열만).
  v3 셀과의 거리 기록과 새 셀 수의 예비 추정에만 쓴다. v3 는 고치지 않는다.
- data/raw/calm/PANGAEA_972777_CALM_ALT_NH.tab (읽기 전용, 있을 때만).
  출처가 'CALM network' 인 기록의 값을 CALM 연 값의 기간 평균과 대조한다.

라벨 규칙(6B.2, 6B.3). 출처 열(Soure)로 나눈다.
- 'Cao, 2020'(치롄 GPR): method = gpr, label_def = direct_dated, eos_basis = none.
  조사 월을 확인하지 못했다. 계절 말로 확인되면 direct_eos(dataset_statement)로 올린다.
  올리는 일은 새 지역 실행의 제출 전에만 할 수 있고 이 파서의 CAO_EOS_CONFIRMED 값을 바꿔서 한다.
- 'CALM network': method = borehole_temp, label_def = temp_derived, eos_basis = protocol.
- 그 밖의 출처: method = unknown, label_def = unknown, eos_basis = none.
- alt_cm = ALT(m) × 100 (십진 연산). 범위 QC 는 0 < alt_cm <= 600.
- 관측 기간 열의 시작과 끝을 year_min, year_max 에 넣는다. 한 해만 적힌 기록은 year 도 채운다.
  월과 날짜는 원자료에 없다. month 와 date 는 빈 칸이다.
- qtp_du_gpr 의 점과 체비쇼프 0.01° 이내이고 값이 같은(차이 0.5 cm 미만) 기록은 dup_of 에 적는다.
- 우측 절단 표기('>')는 원자료에 없다. right_censored = 0.
- disturbed = 0 은 '원자료에 교란 표기가 없다'는 뜻이다. 현장 교란 부재를 확인한 값이 아니다.
- 기록은 지우지 않는다. 규칙에 맞지 않는 기록은 label_def, qc_flag, dup_of 로 표시한다.

산출
- data/processed/ext_labels/qtp_liux_compile_points.csv  표준 점 자료(30열)
- data/processed/ext_labels/qtp_liux_compile_meta.json   메타와 품질 요약
- data/raw/qtp_liux_compile/ds01_all_rows.csv            시트 전체 행과 판정 결과

실행(ROOT, 스레드 1개):
  OMP_NUM_THREADS=1 OPENBLAS_NUM_THREADS=1 MKL_NUM_THREADS=1 python3 scripts/1_data_prep/parse_ext_qtp_liux_compile.py
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
SRC_ID = "qtp_liux_compile"
RAW = ROOT / "data" / "raw" / SRC_ID
XLSX_NAME = "ds01.xlsx"
XLSX = RAW / XLSX_NAME
LOCAL_COPY = ROOT / "data" / "raw" / "qtp_alt_open_2026-09-29" / "figshare21444879" / XLSX_NAME
FIGSHARE_JSON = RAW / "figshare_article_21444879.json"
RAW_DUMP = RAW / "ds01_all_rows.csv"
OUT_DIR = ROOT / "data" / "processed" / "ext_labels"
POINTS = OUT_DIR / f"{SRC_ID}_points.csv"
META = OUT_DIR / f"{SRC_ID}_meta.json"
SCHEMA = OUT_DIR / "_schema.json"
DU_SRC = "qtp_du_gpr"
DU_POINTS = OUT_DIR / f"{DU_SRC}_points.csv"
V3 = ROOT / "data" / "processed" / "fidelity_base_v3.csv"
CALM_TAB = ROOT / "data" / "raw" / "calm" / "PANGAEA_972777_CALM_ALT_NH.tab"

NAME = "Liu X. 등 2022, 칭하이 고원 ALT 관측 문헌 집계표(figshare 21444879)"
URL = "https://doi.org/10.6084/m9.figshare.21444879.v1"
LANDING = "https://figshare.com/articles/dataset/Active_layer_thickness_observations/21444879"
FILE_URL = "https://ndownloader.figshare.com/files/38063142"
API_URL = "https://api.figshare.com/v2/articles/21444879"
DOI = "10.6084/m9.figshare.21444879.v1"
FIGSHARE_MD5 = "b7eed03faa1d6fc45c78f21016574d2f"   # figshare API 의 computed_md5(2026-09-29 조회)
FIGSHARE_SIZE = 31302
ACCESSED = "2026-09-29"

LICENSE = "CC BY 4.0"
CITATION = (
    "Liu, X., Zhou, T., Zhang, Y., Luo, H., Yu, P., Zhang, J., Zeng, J., Xu, Y. (2022). "
    "Active layer thickness observations. figshare. Dataset. "
    "https://doi.org/10.6084/m9.figshare.21444879.v1"
)

SHEET = "ALT_Data"
HEADER = ["ID", "ID_paper", "Site name", "Site Code", "Latitude(°N)", "Longitude(°E)",
          "Observation period", "ALT(m)", "Soure"]
KEYS = ["table_id", "id_paper", "site_name", "site_code", "lat", "lon", "period", "alt_m", "source"]

# 출처 열의 값과 탐색 단계에서 센 기록 수. 수가 다르면 meta 의 source_count_check 에 적는다.
EXPECTED_SOURCES = {
    "Cao, 2020": 150, "Zhao et al., 2017": 66, "Wu et al., 2012": 24, "Qin et al., 2017": 18,
    "Mu et al., 2015": 9, "Luo et al., 2012": 9, "Cao et al., 2018": 8, "Zhang et al., 2012": 7,
    "Wu and Zhang, 2010": 7, "CALM network": 5, "Yue et al., 2013": 2, "Cai et al., 2016": 1,
}
CAO = "Cao, 2020"
CALM = "CALM network"
# 치롄 GPR 의 조사 시기가 계절 말로 확인되면 True 로 바꾼다(새 지역 실행의 제출 전에만).
CAO_EOS_CONFIRMED = False
SOURCE_RULE = {
    CAO: dict(method="gpr", label_def="direct_eos" if CAO_EOS_CONFIRMED else "direct_dated",
              eos_basis="dataset_statement" if CAO_EOS_CONFIRMED else "none", value_kind=""),
    CALM: dict(method="borehole_temp", label_def="temp_derived", eos_basis="protocol",
               value_kind="multiyear_mean"),
}
DEFAULT_RULE = dict(method="unknown", label_def="unknown", eos_basis="none", value_kind="")

# 좌표가 의심스러운 기록(값은 고치지 않고 qc_flag 에 coord 를 적는다). 키는 Site Code.
COORD_SUSPECT = {
    "S237": "위도 의심. 같은 계열 WQ01–WQ18 의 다른 17개 기록은 위도 35.27–35.59 인데 이 기록만 34.42 다. "
            "원 문헌과 대조하지 못했다",
}

# Tibet 의 지리 범위(6B.4)와 하위 단위 규칙. parse_ext_qtp_du_gpr.py 와 같다. 라벨 값을 쓰지 않는다.
TIBET_BOX = dict(lat_min=26.0, lat_max=40.0, lon_min=73.0, lon_max=105.0)
SUBUNIT_RULE = [
    ("Qilian", "치롄", dict(lat_min=37.0, lat_max=40.0, lon_min=96.0, lon_max=103.0, lon_max_incl=True)),
    ("YellowSrc_East", "황하원·동부", dict(lat_min=33.5, lat_max=36.5, lon_min=95.0, lon_max=101.0, lon_max_incl=True)),
    ("QTEC", "회랑", dict(lat_min=31.0, lat_max=36.2, lon_min=90.0, lon_max=95.0, lon_max_incl=False)),
    ("West_QTP", "서부", dict(lat_min=30.0, lat_max=38.0, lon_min=77.0, lon_max=90.0, lon_max_incl=False)),
]
DUP_DEG = 0.01             # 체비쇼프 거리 기준(6B.3)
DUP_EPS = 1e-9             # 부동소수 오차 허용
SAME_VALUE_CM = 0.5        # 값이 같다고 보는 차이의 상한(cm)
CELL = 0.009               # 약 1 km 셀(6B.3)
ALT_MAX_CM = 600.0
YEAR_MAIN = 1990
RE_YEAR = re.compile(r"^\d{4}$")
RE_RANGE = re.compile(r"^(\d{4})\s*[–—\-−~]\s*(\d{4})$")
RE_LIST = re.compile(r"^\d{4}(\s*[,，、;]\s*\d{4})+$")


def subunit_of(lat: float, lon: float) -> str:
    for name, _, b in SUBUNIT_RULE:
        lon_ok = (b["lon_min"] <= lon <= b["lon_max"]) if b["lon_max_incl"] else (b["lon_min"] <= lon < b["lon_max"])
        if b["lat_min"] <= lat <= b["lat_max"] and lon_ok:
            return name
    return "QTP_other"


def in_tibet(lat: float, lon: float) -> bool:
    b = TIBET_BOX
    return b["lat_min"] <= lat <= b["lat_max"] and b["lon_min"] <= lon <= b["lon_max"]


def file_hash(path: Path, algo: str) -> str:
    h = hashlib.new(algo)
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def n_decimals(x: float) -> int:
    s = repr(float(x))
    if "e" in s or "E" in s:
        s = format(Decimal(s), "f")
    return len(s.split(".")[1].rstrip("0")) if "." in s else 0


def fmt_num(x) -> str:
    """원자료 값을 반올림 없이 적는다(float 의 최단 표현)."""
    return repr(float(x))


def dec_str(d: Decimal) -> str:
    s = format(d.normalize(), "f")
    return s


def coord_prec(lat: float, lon: float) -> tuple[float, str]:
    """좌표 정밀도(도)와 판정 근거. 소수 6자리 이상은 도분초 환산값으로 보고 분, 초 단위를 가린다."""
    d = max(n_decimals(lat), n_decimals(lon))
    if d >= 6:
        if all(abs(v * 60 - round(v * 60)) < 1e-3 for v in (lat, lon)):
            return 0.0167, "dms_minute"
        if all(abs(v * 3600 - round(v * 3600)) < 0.02 for v in (lat, lon)):
            return 0.00028, "dms_second"
        return 1e-05, "decimal_6plus"
    return float(Decimal(1).scaleb(-d)), f"decimal_{d}"


def parse_period(v):
    """관측 기간 열. 반환: (year_min, year_max, years, kind, raw_str). 읽지 못하면 kind = 'unparsed'."""
    if v is None:
        return None, None, [], "missing", ""
    if isinstance(v, bool):
        return None, None, [], "unparsed", str(v)
    if isinstance(v, (int, float)):
        if float(v).is_integer():
            y = int(v)
            return y, y, [y], "single", str(y)
        return None, None, [], "unparsed", repr(v)
    s = str(v).strip()
    if RE_YEAR.match(s):
        y = int(s)
        return y, y, [y], "single", s
    m = RE_RANGE.match(s)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        return a, b, [a, b], "range", s
    if RE_LIST.match(s):
        ys = [int(t) for t in re.findall(r"\d{4}", s)]
        return min(ys), max(ys), ys, "list", s
    return None, None, [], "unparsed", s


def parse_alt(v):
    """ALT(m). 반환: (Decimal cm 또는 None, right_censored, 사유). 단위는 m 이고 100 을 곱한다."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None, 0, "alt_missing"
    if isinstance(v, bool):
        return None, 0, "alt_unparsed"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None, 0, "alt_missing"
        return Decimal(repr(float(v))) * Decimal(100), 0, ""
    s = str(v).strip()
    cens = 0
    if s[:1] in (">", "＞", "≥"):
        cens = 1
        s = s[1:].strip()
    try:
        return Decimal(s) * Decimal(100), cens, ""
    except Exception:
        return None, 0, "alt_unparsed"


def cell_key(lat: float, lon: float):
    ky = math.floor(lat / CELL)
    phi = math.radians((ky + 0.5) * CELL)
    kx = math.floor(lon * math.cos(phi) / CELL)
    return ky, kx


def block_of(lat: float, lon: float) -> int:
    return int(math.floor(lat / 0.5) * 100000 + math.floor(lon / 0.5))


def git_commit() -> str:
    try:
        return subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short", "HEAD"], capture_output=True,
                              text=True, check=True).stdout.strip()
    except Exception:
        return "unknown"


def dist(s: pd.Series) -> dict | None:
    if len(s) == 0:
        return None
    return dict(n=int(len(s)), min=float(s.min()), p25=float(s.quantile(0.25)), median=float(s.median()),
                mean=round(float(s.mean()), 2), p75=float(s.quantile(0.75)), max=float(s.max()))


def read_calm(path: Path) -> pd.DataFrame | None:
    """PANGAEA 972777 의 자료 부분. 머리말(/* ... */) 뒤의 탭 구분 표를 읽는다."""
    if not path.exists():
        return None
    txt = path.read_text(encoding="utf-8", errors="replace")
    i = txt.find("*/")
    if i < 0:
        return None
    df = pd.read_csv(io.StringIO(txt[i + 3:]), sep="\t", dtype=str, keep_default_na=False)
    need = {"Event", "Date/Time", "ALD [cm]"}
    if not need.issubset(df.columns):
        return None
    return df


def main() -> int:
    # ---------- 1. 원자료 확보와 검증 ----------
    RAW.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    origin = "data/raw/qtp_liux_compile/ds01.xlsx (이미 있음)"
    if not XLSX.exists():
        if LOCAL_COPY.exists():
            shutil.copy2(LOCAL_COPY, XLSX)
            origin = f"탐색 단계 사본 복사: {LOCAL_COPY.relative_to(ROOT)}"
        else:
            print("원자료가 없다. 다음 명령으로 받는다:")
            print(f"  curl -L -o {XLSX} {FILE_URL}")
            return 2
    if not DU_POINTS.exists():
        print(f"{DU_POINTS.relative_to(ROOT)} 가 없다. parse_ext_qtp_du_gpr.py 를 먼저 실행한다.")
        return 2
    size = XLSX.stat().st_size
    md5 = file_hash(XLSX, "md5")
    sha = file_hash(XLSX, "sha256")
    checksum_match = (md5 == FIGSHARE_MD5 and size == FIGSHARE_SIZE)
    if LOCAL_COPY.exists() and file_hash(LOCAL_COPY, "sha256") == sha:
        origin = f"탐색 단계 사본({LOCAL_COPY.relative_to(ROOT)})을 복사한 파일이다(sha256 같음)"
    if not checksum_match:
        print(f"경고: 파일의 크기 또는 md5 가 figshare 기록과 다르다(size={size}, md5={md5}).")

    schema = json.loads(SCHEMA.read_text(encoding="utf-8"))
    columns = [c["name"] for c in schema["columns"]]
    vocab = {c["name"]: set(c["vocab"]) for c in schema["columns"] if "vocab" in c}

    # ---------- 2. 시트 읽기 ----------
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    sheetnames = list(wb.sheetnames)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    title = ws.cell(row=1, column=1).value
    header = [ws.cell(row=2, column=j).value for j in range(1, len(HEADER) + 1)]
    header = [h.strip() if isinstance(h, str) else h for h in header]
    if header != HEADER:
        print(f"머리글이 예상과 다르다: {header}")
        return 3
    merged = [str(r) for r in ws.merged_cells.ranges]
    extra_cols = 0
    formula_cells = 0
    raw_rows = []
    for r in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=len(HEADER) + 8):
        vals = [c.value for c in r]
        main_vals = vals[:len(HEADER)]
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in main_vals):
            continue
        if any(v is not None for v in vals[len(HEADER):]):
            extra_cols += 1
        formula_cells += sum(1 for c in r[:len(HEADER)] if c.data_type == "f")
        row = dict(zip(KEYS, main_vals))
        row["excel_row"] = r[0].row
        raw_rows.append(row)
    n_raw = len(raw_rows)

    # ---------- 3. 기준 자료 읽기(Du 점, v3, CALM) ----------
    du = pd.read_csv(DU_POINTS, dtype=str, keep_default_na=False)
    du["lat_f"] = du.lat.astype(float)
    du["lon_f"] = du.lon.astype(float)
    du["alt_f"] = du.alt_cm.astype(float)
    du = du[du.qc_flag.isin(["ok", ""])].reset_index(drop=True)
    du_sha = file_hash(DU_POINTS, "sha256")

    v3 = pd.read_csv(V3, usecols=["loc_id", "lat", "lon", "region", "source_id"], low_memory=False)
    box = v3[v3.lat.between(TIBET_BOX["lat_min"] - 1, TIBET_BOX["lat_max"] + 1)
             & v3.lon.between(TIBET_BOX["lon_min"] - 1, TIBET_BOX["lon_max"] + 1)].reset_index(drop=True)

    calm = read_calm(CALM_TAB)

    # ---------- 4. 기록 단위 파싱 ----------
    coord_groups = defaultdict(list)
    for row in raw_rows:
        if isinstance(row["lat"], (int, float)) and isinstance(row["lon"], (int, float)):
            coord_groups[(float(row["lat"]), float(row["lon"]))].append(str(row["site_code"]))

    out_rows, dump_rows = [], []
    excluded = Counter()
    unexpected_sources = Counter()
    period_kinds = Counter()
    prec_kinds = Counter()
    du_same, du_near_diff, v3_near_rows, calm_check = [], [], [], []
    seen_ids = Counter()

    for row in raw_rows:
        src = str(row["source"]).strip() if row["source"] is not None else ""
        code = str(row["site_code"]).strip() if row["site_code"] is not None else ""
        sname = str(row["site_name"]).strip() if row["site_name"] is not None else ""
        flags, notes = [], []
        reason = ""

        lat, lon = row["lat"], row["lon"]
        coord_ok = isinstance(lat, (int, float)) and isinstance(lon, (int, float)) and not isinstance(lat, bool)
        if coord_ok:
            lat, lon = float(lat), float(lon)
            coord_ok = (-90.0 <= lat <= 90.0) and (-180.0 <= lon <= 180.0)
        if not coord_ok:
            reason = "coord_missing"

        alt, cens, alt_reason = parse_alt(row["alt_m"])
        if alt is None and not reason:
            reason = alt_reason

        y0, y1, years, pkind, praw = parse_period(row["period"])
        period_kinds[pkind] += 1

        dump = dict(excel_row=row["excel_row"], **{k: ("" if row[k] is None else
                    (fmt_num(row[k]) if isinstance(row[k], float) else str(row[k]))) for k in KEYS})
        if reason:
            excluded[reason] += 1
            dump.update(used=0, reason=reason)
            dump_rows.append(dump)
            continue

        if src not in EXPECTED_SOURCES:
            unexpected_sources[src] += 1
        rule = SOURCE_RULE.get(src, DEFAULT_RULE)

        site_id = code if code else f"{SRC_ID}_{int(row['table_id']):04d}"
        seen_ids[site_id] += 1
        if not sname:
            sname = site_id

        # QC
        alt_f = float(alt)
        if not (0.0 < alt_f <= ALT_MAX_CM):
            flags.append("range")
        if code in COORD_SUSPECT:
            flags.append("coord")
            notes.append(COORD_SUSPECT[code])
        tib = in_tibet(lat, lon)
        if not tib:
            flags.append("coord")
            notes.append("좌표가 Tibet 범위(26–40°N, 73–105°E) 밖이다")
        if pkind in ("missing", "unparsed"):
            flags.append("date")
            notes.append(f"관측 기간을 읽지 못했다(원문 '{praw}')")
        elif y0 < 1900 or y1 > 2026 or y0 > y1:
            flags.append("date")
        elif y0 < YEAR_MAIN:
            flags.append("year_pre1990")

        # 출처별 설명
        notes.append(f"period_raw={praw}")
        notes.append(f"table_id={row['table_id']}")
        notes.append(f"id_paper={row['id_paper']}")
        if src == CAO:
            notes.append("치롄 GPR(출처 표기 'Cao, 2020'). 조사 월을 확인하지 못했다. 기간 안의 조사 연도도 표에 없다")
        elif src == CALM:
            notes.append("CALM 지점. 지온에서 유도한 값의 기간 평균으로 보인다")
        else:
            notes.append("측정 방법이 표에 없다")
        if pkind == "list":
            notes.append("관측 연도 " + ", ".join(str(y) for y in years) + "(불연속 표기)")
        if pkind in ("range", "list"):
            if rule["value_kind"] == "":
                notes.append("기간 값의 집계 방식(평균, 최댓값, 단일 조사)을 확인하지 못했다")

        # 같은 좌표의 다른 기록
        same = [c for c in coord_groups[(lat, lon)] if c != code]
        if same:
            notes.append("같은 좌표의 기록: " + ", ".join(same))

        # Du 점과의 대조
        dup_of = ""
        cheb = np.maximum(np.abs(du.lat_f.values - lat), np.abs(du.lon_f.values - lon))
        near = np.where(cheb <= DUP_DEG + DUP_EPS)[0]
        if len(near):
            dv = np.abs(du.alt_f.values[near] - alt_f)
            same_idx = near[dv < SAME_VALUE_CM]
            if len(same_idx):
                cand = sorted(same_idx, key=lambda j: (du.dup_of.values[j] != "", cheb[j], du.site_id.values[j]))
                j = cand[0]
                dup_of = f"{DU_SRC}:{du.site_id.values[j]}"
                others = [f"{du.site_id.values[k]}({du.year.values[k]})" for k in cand[1:]]
                in_period = (y0 is not None and du.year.values[j] != "" and y0 <= int(du.year.values[j]) <= y1)
                notes.append(f"{DU_SRC} 의 {du.site_id.values[j]}({du.year.values[j]}, {du.method.values[j]})와 "
                             f"{cheb[j]:.4f}° 이내이고 값이 같다"
                             + ("" if in_period else ". 연도는 관측 기간과 다르다")
                             + (". 같은 조건의 다른 점: " + ", ".join(others) if others else ""))
                du_same.append(dict(site_id=site_id, site_name=sname, orig_source=src, lat=lat, lon=lon,
                                    period=praw, alt_cm=alt_f, du_site_id=str(du.site_id.values[j]),
                                    du_year=str(du.year.values[j]), du_method=str(du.method.values[j]),
                                    du_lat=float(du.lat_f.values[j]), du_lon=float(du.lon_f.values[j]),
                                    du_alt_cm=float(du.alt_f.values[j]), chebyshev_deg=round(float(cheb[j]), 5),
                                    year_in_period=bool(in_period), other_candidates=others))
            else:
                k = near[np.argmin(dv)]
                notes.append(f"{DU_SRC} 점 {len(near)}개가 0.01° 이내(값 다름, 가장 가까운 값 차이 {float(dv.min()):.0f} cm)")
                du_near_diff.append(dict(site_id=site_id, orig_source=src, n_du_points=int(len(near)),
                                         min_abs_diff_cm=round(float(dv.min()), 1),
                                         du_site_id_min_diff=str(du.site_id.values[k])))

        # v3 셀과의 거리
        if len(box):
            cv = np.maximum(np.abs(box.lat.values - lat), np.abs(box.lon.values - lon))
            for j in np.where(cv <= DUP_DEG + DUP_EPS)[0]:
                notes.append(f"v3 loc_id={int(box.loc_id.values[j])}({box.region.values[j]}, "
                             f"{box.source_id.values[j]}) 0.01° 이내")
                v3_near_rows.append(dict(site_id=site_id, site_name=sname, orig_source=src,
                                         v3_loc_id=int(box.loc_id.values[j]), v3_region=str(box.region.values[j]),
                                         v3_source_id=str(box.source_id.values[j]),
                                         chebyshev_deg=round(float(cv[j]), 5)))

        # CALM 원자료와의 대조
        if src == CALM and calm is not None and y0 is not None:
            ev = f"CALM_{sname}"
            sub = calm[calm["Event"] == ev]
            yrs, vals, n_cens, cens_vals = [], [], 0, []
            for yy, vv in zip(sub["Date/Time"], sub["ALD [cm]"]):
                try:
                    yi = int(str(yy)[:4])
                except ValueError:
                    continue
                if not (y0 <= yi <= y1):
                    continue
                s = str(vv).strip()
                if s == "":
                    continue
                if s.startswith(">"):
                    n_cens += 1
                    try:
                        cens_vals.append(float(s[1:]))
                    except ValueError:
                        pass
                    continue
                try:
                    vals.append(float(s))
                    yrs.append(yi)
                except ValueError:
                    continue
            if vals:
                mean_unc = float(np.mean(vals))
                mean_all = float(np.mean(vals + cens_vals)) if cens_vals else mean_unc
                notes.append(f"PANGAEA 972777 {ev} 의 같은 기간 평균 {mean_unc:.2f} cm({len(vals)}개 연도"
                             + (f", '>' 표기 {n_cens}개 연도 제외" if n_cens else "") + ")")
                calm_check.append(dict(site_id=site_id, site_name=sname, event=ev, period=praw, table_alt_cm=alt_f,
                                       calm_mean_uncensored_cm=round(mean_unc, 2), n_years_uncensored=len(vals),
                                       n_years_censored=n_cens,
                                       calm_mean_with_censored_lower_bounds_cm=round(mean_all, 2),
                                       diff_table_minus_calm_cm=round(alt_f - mean_unc, 2)))
            else:
                calm_check.append(dict(site_id=site_id, site_name=sname, event=ev, period=praw, table_alt_cm=alt_f,
                                       calm_mean_uncensored_cm=None, n_years_uncensored=0,
                                       n_years_censored=n_cens))

        prec, pk = coord_prec(lat, lon)
        prec_kinds[pk] += 1

        out = dict.fromkeys(columns, "")
        out.update(
            src_id=SRC_ID, site_id=site_id, site_name=sname,
            lat=fmt_num(lat), lon=fmt_num(lon),
            year=(str(y0) if pkind == "single" else ""), month="",
            alt_cm=dec_str(alt), method=rule["method"], label_def=rule["label_def"],
            n_obs="1", country="China", macro=("Tibet" if tib else "other"),
            citation=CITATION, license=LICENSE,
            subunit=(subunit_of(lat, lon) if tib else ""), date="",
            eos_basis=rule["eos_basis"], value_kind=rule["value_kind"],
            year_min=("" if y0 is None else str(y0)), year_max=("" if y1 is None else str(y1)),
            alt_sd_cm="", coord_prec_deg=repr(prec), disturbed="0", disturb_type="",
            right_censored=str(cens), orig_source=src, dup_of=dup_of,
            qc_flag=(";".join(dict.fromkeys(flags)) if flags else "ok"),
            notes="; ".join(notes),
        )
        for k, allowed in vocab.items():
            v = out[k]
            if v == "":
                continue
            if k in ("disturbed", "right_censored"):
                ok = int(v) in allowed
            else:
                ok = v in allowed
            if not ok:
                print(f"어휘 밖의 값: {k}={v} (site_id={site_id})")
                return 4
        out_rows.append(out)
        dump.update(used=1, reason="", site_id=site_id, method=out["method"], label_def=out["label_def"],
                    dup_of=dup_of, qc_flag=out["qc_flag"])
        dump_rows.append(dump)

    dup_ids = {k: v for k, v in seen_ids.items() if v > 1}
    if dup_ids:
        print(f"site_id 가 겹친다: {dup_ids}")
        return 5

    # ---------- 5. 쓰기 ----------
    with open(POINTS, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columns, lineterminator="\n")
        w.writeheader()
        w.writerows(out_rows)
    dump_cols = ["excel_row"] + KEYS + ["used", "reason", "site_id", "method", "label_def", "dup_of", "qc_flag"]
    with open(RAW_DUMP, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=dump_cols, lineterminator="\n", restval="")
        w.writeheader()
        w.writerows(dump_rows)

    # ---------- 6. 품질 요약 ----------
    df = pd.DataFrame(out_rows)
    df["lat_f"] = df.lat.astype(float)
    df["lon_f"] = df.lon.astype(float)
    df["alt_f"] = df.alt_cm.astype(float)
    df["y0"] = pd.to_numeric(df.year_min, errors="coerce")
    df["y1"] = pd.to_numeric(df.year_max, errors="coerce")

    du_cells = {cell_key(a, b) for a, b, d in zip(du.lat_f, du.lon_f, du.dup_of) if d == ""}

    def cell_estimate(sub: pd.DataFrame, label: str) -> dict:
        """6B.3 의 셀 규칙으로 센 예비 추정. qc_flag 가 ok 이고 dup_of 가 빈 행만 쓴다."""
        m = sub[(sub.qc_flag == "ok") & (sub.dup_of == "")].copy()
        if len(m) == 0:
            return dict(subset=label, n_rows_used=0, n_cells=0)
        m["cell"] = [cell_key(a, b) for a, b in zip(m.lat_f, m.lon_f)]
        cells = m.groupby("cell", as_index=False).agg(lat=("lat_f", "mean"), lon=("lon_f", "mean"),
                                                      alt_cm=("alt_f", "mean"), n_loc=("site_id", "size"))
        cells["block"] = [block_of(a, b) for a, b in zip(cells.lat, cells.lon)]
        cells["subunit"] = [subunit_of(a, b) for a, b in zip(cells.lat, cells.lon)]

        def near(ref: pd.DataFrame) -> np.ndarray:
            if len(ref) == 0:
                return np.zeros(len(cells), dtype=bool)
            dl = np.abs(cells.lat.values[:, None] - ref.lat.values[None, :])
            dn = np.abs(cells.lon.values[:, None] - ref.lon.values[None, :])
            return (np.maximum(dl, dn) <= DUP_DEG + DUP_EPS).any(axis=1)

        cells["dup_v3_direct"] = near(box[box.source_id == "F4_direct"])
        cells["near_v3_other"] = near(box[box.source_id != "F4_direct"])
        cells["shared_du"] = [c in du_cells for c in cells.cell]
        new = cells[~cells.dup_v3_direct]
        new_nodu = new[~new.shared_du]
        return dict(
            subset=label, n_rows_used=int(len(m)), n_cells=int(len(cells)), n_blocks=int(cells.block.nunique()),
            n_cells_dup_v3_F4_direct=int(cells.dup_v3_direct.sum()),
            n_cells_near_v3_other_source=int(cells.near_v3_other.sum()),
            n_new_cells=int(len(new)), n_new_blocks=int(new.block.nunique()),
            n_new_cells_sharing_cell_with_qtp_du_gpr=int(new.shared_du.sum()),
            n_new_cells_not_in_qtp_du_gpr=int(len(new_nodu)),
            n_new_blocks_not_in_qtp_du_gpr=int(new_nodu.block.nunique()),
            new_cells_by_subunit={k: int(v) for k, v in new.subunit.value_counts().items()},
            new_blocks_by_subunit={k: int(v) for k, v in new.groupby("subunit").block.nunique().items()},
            cell_alt_cm=dist(new.alt_cm),
        )

    def exploration_rule(sub: pd.DataFrame, label: str) -> dict:
        """탐색 단계 규칙(newcells.py)의 재현. 0.01° 반올림 격자, v3 전체 행과 0.01° 미만이면 중복."""
        if len(sub) == 0:
            return dict(subset=label, n_rows=0)
        d = pd.DataFrame(dict(lat=sub.lat_f.values, lon=sub.lon_f.values))
        d["clat"] = (d.lat / 0.01).round() * 0.01
        d["clon"] = (d.lon / 0.01).round() * 0.01
        c = d.groupby(["clat", "clon"], as_index=False).agg(lat=("lat", "mean"), lon=("lon", "mean"))
        if len(v3):
            dl = np.abs(c.lat.values[:, None] - box.lat.values[None, :])
            dn = np.abs(c.lon.values[:, None] - box.lon.values[None, :])
            dupv = (np.maximum(dl, dn) < 0.01).any(axis=1) if len(box) else np.zeros(len(c), dtype=bool)
        else:
            dupv = np.zeros(len(c), dtype=bool)
        c["block"] = [block_of(a, b) for a, b in zip(c.lat, c.lon)]
        new = c[~dupv]
        return dict(subset=label, n_rows=int(len(d)), n_cells_0p01deg=int(len(c)), n_dup_v3=int(dupv.sum()),
                    n_new_cells=int(len(new)), n_new_blocks=int(new.block.nunique()))

    cao_df = df[df.orig_source == CAO]
    cell_est = dict(
        rule="6B.3 의 약 1 km 셀(ky = floor(lat/0.009), kx = floor(lon·cos φ/0.009)). 위치는 site_id 단위. "
             "qc_flag 가 ok 이고 dup_of 가 빈 행만 썼다",
        status="예비 추정이다. 확정 값은 v4 조립 단계에서 정한다. label_def 가 direct_eos 가 아닌 행은 "
               "대상 셀이 아니다(6B.4)",
        subsets=[
            cell_estimate(cao_df, "orig_source = 'Cao, 2020' (치롄 GPR)"),
            cell_estimate(df[df.label_def == "unknown"], "label_def = unknown"),
            cell_estimate(df[df.label_def == "temp_derived"], "label_def = temp_derived"),
            cell_estimate(df, "표 전체"),
        ],
        exploration_rule_check=[
            exploration_rule(cao_df, "orig_source = 'Cao, 2020'"),
            exploration_rule(df, "표 전체"),
        ],
        exploration_estimate="치롄 GPR 새 셀 63개(블록 2개), 표 전체 새 셀 172개(블록 29개). 탐색 단계 값이고 "
                             "격자 정의가 다르다(0.01° 반올림 격자)",
        v3_rows_in_tibet_box=[dict(region=a, source_id=b, n=int(c)) for (a, b), c in
                              box.groupby(["region", "source_id"]).size().items()],
    )

    # 표 안의 같은 좌표 묶음
    same_coord = []
    for (la, lo), codes in coord_groups.items():
        if len(codes) > 1:
            g = df[df.site_id.isin(codes)]
            same_coord.append(dict(lat=la, lon=lo, site_ids=list(g.site_id), site_names=list(g.site_name),
                                   orig_sources=sorted(set(g.orig_source)), alt_cm=[float(a) for a in g.alt_f],
                                   same_value=bool(g.alt_f.nunique() == 1)))
    dup_names = {k: list(g.site_id) for k, g in df.groupby("site_name") if len(g) > 1}

    # 치롄 GPR 의 요약 통계(Cao 등 2017 초록의 Yeniu Gou 값과 대조)
    cao_stat = None
    if len(cao_df):
        a = cao_df.alt_f / 100.0
        cao_stat = dict(n=int(len(a)), mean_m=round(float(a.mean()), 3), sd_m=round(float(a.std(ddof=1)), 3),
                        min_m=float(a.min()), max_m=float(a.max()),
                        abstract_yeniugou="mean 2.72 ± 0.88 m, range 1.07 m to 4.86 m",
                        abstract_eboling="mean 1.32 ± 0.29 m, range 0.81 to 2.1 m",
                        abstract_source="Cao 등 2017, JGR Earth Surface 122, 574–591, doi:10.1002/2016JF004018 "
                                        "(Crossref 의 초록, 2026-09-29 조회)",
                        site_name_range=f"{cao_df.site_name.iloc[0]}–{cao_df.site_name.iloc[-1]}",
                        lat_range=[float(cao_df.lat_f.min()), float(cao_df.lat_f.max())],
                        lon_range=[float(cao_df.lon_f.min()), float(cao_df.lon_f.max())])

    by_source = []
    for s, g in df.groupby("orig_source"):
        by_source.append(dict(orig_source=s, n=int(len(g)), method=g.method.iloc[0], label_def=g.label_def.iloc[0],
                              year_min=int(g.y0.min()), year_max=int(g.y1.max()),
                              n_single_year=int((g.year != "").sum()),
                              n_dup_of=int((g.dup_of != "").sum()),
                              n_qc_not_ok=int((g.qc_flag != "ok").sum()),
                              subunits={k: int(v) for k, v in g.subunit.value_counts().items()},
                              alt_cm=dist(g.alt_f)))
    by_source.sort(key=lambda r: -r["n"])

    source_count_check = dict(
        expected=EXPECTED_SOURCES,
        found={k: int(v) for k, v in Counter(r["source"] if r["source"] is None else str(r["source"]).strip()
                                              for r in raw_rows).items()},
        unexpected_sources=dict(unexpected_sources),
    )
    source_count_check["match"] = (source_count_check["found"] == EXPECTED_SOURCES)

    other = df[~df.orig_source.isin([CAO, CALM])]
    n_other_src, n_other_rec = int(other.orig_source.nunique()), int(len(other))
    zhao = df[df.orig_source == "Zhao et al., 2017"]
    n_zhao, n_zhao_dup = int(len(zhao)), int((zhao.dup_of != "").sum())

    meta = dict(
        src_id=SRC_ID, name=NAME, url=URL, landing_page=LANDING, doi=DOI, accessed=ACCESSED,
        figshare=dict(article_id=21444879, title="Active layer thickness observations", version=1,
                      published_date="2022-11-02", defined_type="dataset", file_id=38063142,
                      file_md5=FIGSHARE_MD5, checksum_match=bool(checksum_match), api=API_URL,
                      description="Dataset of active layer thickness (ALT) observations on the Qinghai Plateau.",
                      linked_paper="기록에 연결 논문(resource_doi, references)이 없다",
                      linked_paper_candidates_checked=[
                          dict(doi="10.1186/s13021-022-00203-z", journal="Carbon Balance and Management (2022)",
                               result="저자 목록이 figshare 기록과 다르다. 연결을 확인하지 못했다"),
                          dict(doi="10.1016/j.geoderma.2023.116488", journal="Geoderma 434 (2023)",
                               result="저자 목록이 figshare 기록과 다르다. 본문을 읽지 못했다. 연결을 확인하지 못했다"),
                      ]),
        files=[dict(name=XLSX_NAME, size=size, sha256=sha, md5=md5, url=FILE_URL, origin=origin)]
              + ([dict(name=FIGSHARE_JSON.name, size=FIGSHARE_JSON.stat().st_size,
                       sha256=file_hash(FIGSHARE_JSON, "sha256"), url=API_URL)] if FIGSHARE_JSON.exists() else []),
        license=LICENSE,
        license_detail="figshare 기록의 약관 표기는 CC BY 4.0 이다(https://creativecommons.org/licenses/by/4.0/). "
                       "재배포를 허용하므로 점 자료를 저장소에 둘 수 있다. 집계표가 옮겨 실은 원 문헌 값의 "
                       "개별 약관은 확인하지 못했다",
        citation=CITATION,
        parser="scripts/1_data_prep/parse_ext_qtp_liux_compile.py",
        parser_git_commit=f"{git_commit()} (HEAD, 파서는 아직 커밋되지 않았다)",
        schema_version=schema.get("version"),
        sheet=dict(file=XLSX_NAME, sheet_names=sheetnames, sheet_used=ws.title, title_row=title, header_row=2,
                   header=HEADER, merged_cells=merged, n_rows_with_values_outside_9_columns=extra_cols,
                   n_formula_cells=formula_cells),
        n_rows_raw=n_raw, n_rows_points=len(out_rows), n_sites=int(df.site_id.nunique()),
        n_unique_locations=int(len(coord_groups)),
        n_by_label_def={k: int(v) for k, v in df.label_def.value_counts().items()},
        n_by_method={k: int(v) for k, v in df.method.value_counts().items()},
        n_by_eos_basis={k: int(v) for k, v in df.eos_basis.value_counts().items()},
        n_by_value_kind={(k if k else "(빈 칸)"): int(v) for k, v in df.value_kind.value_counts().items()},
        n_by_macro={k: int(v) for k, v in df.macro.value_counts().items()},
        n_by_subunit={k: int(v) for k, v in df.subunit.value_counts().items()},
        n_by_orig_source={k: int(v) for k, v in df.orig_source.value_counts().items()},
        n_by_qc_flag={k: int(v) for k, v in df.qc_flag.value_counts().items()},
        n_by_period_kind=dict(period_kinds),
        n_with_year=int((df.year != "").sum()),
        n_by_year={str(k): int(v) for k, v in df[df.year != ""].year.value_counts().sort_index().items()},
        n_right_censored=int((df.right_censored == "1").sum()),
        n_with_dup_of=int((df.dup_of != "").sum()),
        year_range=[int(df.y0.min()), int(df.y1.max())],
        n_year_min_pre1990=int((df.y0 < YEAR_MAIN).sum()),
        alt_cm=dist(df.alt_f),
        alt_cm_by_label_def={k: dist(g.alt_f) for k, g in df.groupby("label_def")},
        alt_cm_by_subunit={k: dist(g.alt_f) for k, g in df.groupby("subunit")},
        by_orig_source=by_source,
        lat_range=[float(df.lat_f.min()), float(df.lat_f.max())],
        lon_range=[float(df.lon_f.min()), float(df.lon_f.max())],
        n_excluded_by_reason=dict(excluded),
        exclusion_reason_desc=dict(coord_missing="좌표가 없거나 범위 밖이다", alt_missing="ALT 값이 없다",
                                   alt_unparsed="ALT 값을 숫자로 읽지 못했다"),
        label_rule=dict(
            by_source={CAO: SOURCE_RULE[CAO], CALM: SOURCE_RULE[CALM], "그 밖의 출처": DEFAULT_RULE},
            cao_eos_confirmed=CAO_EOS_CONFIRMED,
            note="치롄 GPR 의 조사 시기가 계절 말로 확인되면 direct_eos(dataset_statement)로 올린다. "
                 "새 지역 실행의 제출 전에만 바꿀 수 있다",
        ),
        source_count_check=source_count_check,
        duplicates=dict(
            rule=f"{DU_SRC} 의 점과 체비쇼프 {DUP_DEG}° 이내이고 값의 차이가 {SAME_VALUE_CM} cm 미만이면 "
                 "dup_of 에 적는다. 후보가 여럿이면 dup_of 가 빈 Du 점, 가까운 점의 순서로 고른다",
            reference_file=str(DU_POINTS.relative_to(ROOT)), reference_sha256=du_sha,
            reference_rows_used=int(len(du)),
            same_value_within_0p01deg=du_same,
            n_same_value=len(du_same),
            n_records_with_du_point_within_0p01deg_different_value=len(du_near_diff),
            different_value_by_source={k: int(v) for k, v in
                                       Counter(r["orig_source"] for r in du_near_diff).items()},
            same_coordinate_groups_within_table=same_coord,
            n_same_coordinate_groups=len(same_coord),
            n_records_in_same_coordinate_groups=int(sum(len(g["site_ids"]) for g in same_coord)),
            repeated_site_names=dup_names,
            handling="기록은 지우지 않았다. Du 점과 겹치는 기록은 dup_of 에, 표 안의 같은 좌표 기록과 "
                     "v3 셀 근접은 notes 에 적었다",
        ),
        v3_proximity=dict(rule=f"v3 의 행과 체비쇼프 {DUP_DEG}° 이내인 기록", rows=v3_near_rows,
                          n_records=len({r["site_id"] for r in v3_near_rows})),
        calm_crosscheck=dict(
            reference="data/raw/calm/PANGAEA_972777_CALM_ALT_NH.tab" if calm is not None else None,
            method="집계표의 관측 기간 안에 있는 CALM 연 값('>' 표기 연도 제외)의 평균과 집계표 값을 견준다",
            rows=calm_check,
        ),
        cao2020_identification=cao_stat,
        coordinate_conversion="변환하지 않았다. 원자료가 십진 도(Latitude(°N), Longitude(°E) 열)다. 측지계 표기는 "
                              "원자료에 없다(WGS84 로 가정). 북위, 동경이므로 부호는 양수다. 값은 float 의 최단 "
                              "표현으로 적었다(반올림하지 않았다). 소수 8자리 값은 원 문헌의 도분초를 집계표 "
                              "작성자가 환산한 것으로 보인다",
        coord_prec_rule="d 는 위도와 경도의 소수 자릿수 가운데 큰 값이다. d 가 5 이하이면 10^-d. d 가 6 이상이면 "
                        "두 좌표가 모두 분의 정수배일 때 0.0167, 초의 정수배일 때 0.00028, 그 밖에는 0.00001 이다. "
                        "측위 정확도는 확인하지 못했다",
        n_by_coord_prec_kind=dict(prec_kinds),
        subunit_rule=[dict(subunit=a, label=b, box=c) for a, b, c in SUBUNIT_RULE]
                     + [dict(subunit="QTP_other", label="그 밖", box="위 네 상자 밖의 Tibet 범위")],
        macro_rule=dict(macro="Tibet", box=TIBET_BOX),
        unit_conversion="alt_cm = ALT(m) × 100 (십진 연산). 표 제목에 단위가 m 로 적혀 있다",
        missing_value_code="빈 셀(None). 숫자 결측 부호(-999 등)는 없다. 9개 열에 빈 셀이 없다",
        right_censor_code="'>' 표기는 원자료에 없다(ALT 열은 모두 숫자다)",
        disturbed_convention="disturbed = 0 은 원자료에 교란 표기가 없다는 뜻이다",
        eos_evidence=[
            "점별 관측 날짜와 월이 표에 없다",
            "치롄 GPR(Cao, 2020)의 조사 월을 확인하지 못했다. Cao 등 2017 의 본문은 이 서버에서 403 이고 "
            "초록에는 조사 시기가 없다. TPDC 자료 설명 쪽은 스크립트 렌더링 방식이라 본문을 읽지 못했다",
        ],
        preliminary_cell_estimate=cell_est,
        unverified_items=[
            "치롄 GPR(출처 표기 'Cao, 2020', 150기록)의 조사 월과 조사 연도. 표에는 기간 2011–2014 만 있다",
            "출처 표기 'Cao, 2020' 이 가리키는 문헌. TPDC 자료(doi:10.11888/Geocry.tpdc.270324)로 보이나 "
            "집계표에 참고문헌 목록이 없다",
            "연결 논문. figshare 기록에 논문 DOI 가 없다. 표 제목의 'Data Set S1' 은 논문 보충 자료임을 뜻한다",
            f"CALM 과 치롄 GPR 을 뺀 {n_other_src}개 출처({n_other_rec}기록)의 측정 방법(시추공 지온, 탐침, 탐갱, GPR "
            "가운데 어느 것인지)",
            "관측 기간이 여러 해인 기록의 값이 평균인지, 최댓값인지, 단일 조사값인지",
            "관측 기간이 한 해인 기록의 조사 월",
            "S237(WQ06)의 위도 34.42 가 오기인지 여부",
            "표 안에서 좌표가 같은 기록(소수 2자리 좌표의 Zhao et al., 2017 기록 등)이 서로 다른 지점인지 여부",
            f"Zhao et al., 2017 의 {n_zhao}개 기록 가운데 Du 점과 값이 같은 {n_zhao_dup}개(BQ 계열, 2011년)를 뺀 "
            f"{n_zhao - n_zhao_dup}개 기록이 GPR 조사값인지 여부. 겹치는 Du 점은 모두 GPR 이다",
            "측지계(WGS84 로 가정)",
            "교란 지점 여부. 원자료에 표기가 없어 disturbed 는 0 으로 두었다",
            "행마다의 원관측 수. n_obs 는 1 로 두었다",
        ],
    )
    META.write_text(json.dumps(meta, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    # ---------- 7. 화면 요약 ----------
    print(f"원자료 {n_raw}행 → 점 자료 {len(out_rows)}행, 제외 {sum(excluded.values())}행 {dict(excluded)}")
    print(f"파일 검증: size={size}, md5 일치={checksum_match}")
    print("label_def:", meta["n_by_label_def"])
    print("method:", meta["n_by_method"])
    print("subunit:", meta["n_by_subunit"])
    print("qc_flag:", meta["n_by_qc_flag"])
    print("연도 범위:", meta["year_range"], "한 해 기록:", meta["n_with_year"])
    print("ALT(cm):", meta["alt_cm"])
    print("dup_of:", meta["n_with_dup_of"], [r["site_id"] + "→" + r["du_site_id"] for r in du_same])
    print("CALM 대조:", [(r["site_name"], r["table_alt_cm"], r.get("calm_mean_uncensored_cm")) for r in calm_check])
    for s in cell_est["subsets"]:
        print("셀 예비 추정:", {k: s.get(k) for k in ("subset", "n_rows_used", "n_cells", "n_blocks", "n_new_cells",
                                                   "n_new_blocks", "n_new_cells_not_in_qtp_du_gpr")})
    for s in cell_est["exploration_rule_check"]:
        print("탐색 규칙 재현:", s)
    print(f"쓰기: {POINTS.relative_to(ROOT)}, {META.relative_to(ROOT)}, {RAW_DUMP.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
